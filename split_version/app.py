import sys
import os
import json
import base64
import math
import re
import threading
import logging
import traceback
import subprocess
import argparse
import shutil
import time
from queue import Empty, Queue
from datetime import datetime, timezone

# split_version 안에서 실행해도 원본 공통 모듈을 그대로 import한다.
_SPLIT_DIR = os.path.dirname(os.path.abspath(__file__))
_PROJECT_ROOT = os.path.dirname(_SPLIT_DIR)
for _path in (_SPLIT_DIR, _PROJECT_ROOT):
    if _path not in sys.path:
        sys.path.insert(0, _path)


_EARLY_AGENT_LOCK = None


def _early_env_truthy(name):
    value = os.environ.get(name, "")
    return value.strip().lower() in {"1", "true", "yes", "y", "on"}


def _early_agent_mode_requested(argv=None):
    argv = list(sys.argv[1:] if argv is None else argv)
    if "--diagnostics-probe" in argv or "--repair-local-state" in argv:
        return False
    if "--legacy-ui" in argv or _early_env_truthy("AIMAX_LEGACY_UI"):
        return False
    if any(flag in argv for flag in ("--connect", "--status", "--open-settings", "--open_settings")):
        return True
    if "--agent" in argv or _early_env_truthy("AIMAX_AGENT_MODE"):
        return True
    return bool(getattr(sys, "frozen", False)) or "__compiled__" in globals()


def _early_agent_request_kind(argv=None):
    joined = " ".join(list(sys.argv[1:] if argv is None else argv)).lower()
    if "open_settings" in joined or "open-settings" in joined or "settings" in joined:
        return "open_settings"
    if "status" in joined:
        return "status"
    return "connect"


def _preacquire_agent_lock():
    global _EARLY_AGENT_LOCK
    if not _early_agent_mode_requested():
        return
    try:
        from local_agent.single_instance import SingleInstanceError, acquire_single_instance_lock

        _EARLY_AGENT_LOCK = acquire_single_instance_lock()
    except SingleInstanceError:
        try:
            from local_agent.single_instance import signal_existing_instance

            signal_existing_instance(_early_agent_request_kind())
        except Exception:
            pass
        sys.exit(0)
    except Exception:
        pass


_preacquire_agent_lock()


def _hidden_subprocess_kwargs():
    if os.name != "nt":
        return {}
    kwargs = {}
    creationflags = getattr(subprocess, "CREATE_NO_WINDOW", 0)
    if creationflags:
        kwargs["creationflags"] = creationflags
    try:
        startupinfo = subprocess.STARTUPINFO()
        startupinfo.dwFlags |= subprocess.STARTF_USESHOWWINDOW
        startupinfo.wShowWindow = subprocess.SW_HIDE
        kwargs["startupinfo"] = startupinfo
    except Exception:
        pass
    return kwargs


# macOS PyInstaller 번들 SSL 인증서 경로 설정
try:
    import certifi
    os.environ.setdefault("SSL_CERT_FILE", certifi.where())
    os.environ.setdefault("REQUESTS_CA_BUNDLE", certifi.where())
except ImportError:
    pass

import ttkbootstrap as ttk
from ttkbootstrap.constants import *
from tkinter import filedialog, END
import tkinter as tk
import aimax_compliance as aimax
from local_agent.runtime import HeadlessAgentMixin, agent_mode_requested, env_truthy

# ── 플랫폼별 폰트 ──
if sys.platform == "darwin":
    FONT_UI   = "Apple SD Gothic Neo"
    FONT_MONO = "Monaco"
elif sys.platform.startswith("win"):
    FONT_UI   = "Malgun Gothic"
    FONT_MONO = "Consolas"
else:
    FONT_UI   = "DejaVu Sans"
    FONT_MONO = "DejaVu Sans Mono"



# ──────────────────────────────────────────────
# 설정 파일 경로 (플랫폼별 사용자 데이터 디렉토리 사용)
# ──────────────────────────────────────────────
from paths import SETTINGS_PATH as _SETTINGS_PATH, APP_DATA_DIR as _APP_DATA_DIR

SETTINGS_PATH = str(_SETTINGS_PATH)
SECRET_FALLBACK_PATH = str(_APP_DATA_DIR / ".settings_secrets.json")
EXPORTS_DIR = str(_APP_DATA_DIR / "exports")
CREW_RESULTS_DIR = str(_APP_DATA_DIR / "crew_results")


# ──────────────────────────────────────────────
# 설정 저장 / 로드
# ──────────────────────────────────────────────
# 새 저장값은 AIMAX 이름으로 관리하고, 기존 설치 사용자는 NaverBlogAuto에서 1회 폴백/마이그레이션한다.
_KR_SERVICE = os.environ.get("AIMAX_KEYRING_SERVICE", "AIMAX")
_KR_LEGACY_SERVICE = "NaverBlogAuto"
_SECRET_ENV_KEYS = {
    "naver_pw": "AIMAX_NAVER_PASSWORD",
    "gemini_api_key": "AIMAX_GEMINI_API_KEY",
    "claude_api_key": "AIMAX_CLAUDE_API_KEY",
    "openai_api_key": ("AIMAX_OPENAI_API_KEY", "OPENAI_API_KEY", "openai_api_key"),
    "apify_api_token": ("AIMAX_APIFY_API_TOKEN", "APIFY_API_TOKEN", "apify_api_token"),
}
_DEFAULT_AI_MODEL = "gemini-2.5-pro"
API_KEY_GUIDE_URL = os.environ.get(
    "AIMAX_API_KEY_GUIDE_URL",
    "https://www.notion.so/367b31f1da5581ed9b11f23757476cd2",
) or "https://www.notion.so/367b31f1da5581ed9b11f23757476cd2"
_KEYCHAIN_READ_UNAVAILABLE = False
_USD_KRW_RATE = 1476
_USD_KRW_RATE_LABEL = "2026-05-06 Wise/Investing.com spot"
_GEMINI_IMAGE_PRICE_USD = 0.039
_OPENAI_IMAGE_PRICE_USD = 0.042
_AI_TEXT_PRICE_USD_PER_1M = {
    "gemini-2.5-pro": {"input": 1.25, "output": 10.00, "label": "Gemini 2.5 Pro"},
    "gemini-2.5-flash": {"input": 0.30, "output": 2.50, "label": "Gemini 2.5 Flash"},
    "claude": {"input": 3.00, "output": 15.00, "label": "Claude Sonnet"},
    "gpt-5.4-mini": {"input": 0.75, "output": 4.50, "label": "GPT-5.4 mini"},
    "gpt-5-mini": {"input": 0.25, "output": 2.00, "label": "GPT-5 mini"},
}
_LEGACY_AI_MODEL_MAP = {
    "gemini": _DEFAULT_AI_MODEL,
    "gemini-pro": "gemini-2.5-pro",
    "gemini-flash": "gemini-2.5-flash",
    "gemini-3.1-pro-preview": _DEFAULT_AI_MODEL,
}
_PLACEHOLDER_SECRET_VALUES = {
    "",
    "your_naver_id",
    "your_naver_password",
    "your_gemini_api_key",
    "your_claude_api_key",
    "your_openai_api_key",
    "your_apify_api_token",
    "your_api_key_here",
    "your_api_key",
    "changeme",
    "change_me",
}
_PROVIDER_SECRET_KEYS = {
    "gemini_api_key",
    "claude_api_key",
    "openai_api_key",
    "apify_api_token",
}


def _normalize_ai_model(value):
    value = (value or "").strip()
    if value in ("claude", "gemini-2.5-flash", "gemini-2.5-pro", "gpt-5.4-mini", "gpt-5-mini"):
        return value
    return _LEGACY_AI_MODEL_MAP.get(value, _DEFAULT_AI_MODEL)


def _is_openai_model(value):
    return str(value or "").startswith("gpt-")


def _normalize_image_count(value):
    try:
        count = int(value)
    except (TypeError, ValueError):
        count = 3
    return max(0, min(8, count))


def _normalize_target_char_count(value):
    raw = str(value or "").split(" - ")[0].strip()
    raw = raw.replace("자", "").replace("단어", "").strip()
    try:
        count = int(raw)
    except (TypeError, ValueError):
        count = 1500
    return max(300, min(6000, count))


def _is_browser_session_error(error):
    text = str(error or "").lower()
    return any(
        marker in text
        for marker in (
            "invalid session",
            "invalidsessionid",
            "not connected to devtools",
            "disconnected",
            "browser has closed",
            "chrome not reachable",
            "target window already closed",
            "target frame detached",
            "frame detached",
            "no such window",
        )
    )


def _limit_image_blocks(content_list, image_count):
    limit = _normalize_image_count(image_count)
    image_seen = 0
    filtered = []
    for item in content_list:
        if item and item[0] == "image":
            if image_seen >= limit:
                continue
            image_seen += 1
        filtered.append(item)
    return filtered


def _is_empty_image_prompt(prompt):
    text = str(prompt or "").strip().strip("\"'`")
    text = re.sub(r"\s+", " ", text)
    if not text:
        return True
    normalized = text.lower().strip(" .,:;!()[]{}")
    return normalized in {
        "image",
        "image prompt",
        "prompt",
        "photo",
        "picture",
        "이미지",
        "이미지 프롬프트",
        "프롬프트",
        "사진",
        "그림",
    }


def _fallback_image_prompt(title, source, index):
    seed = str(title or source or "블로그 글").strip()
    seed = re.sub(r"\s+", " ", seed)
    seed = seed[:80] or "블로그 글"
    return (
        f"{seed} 주제를 시각적으로 설명하는 네이버 블로그용 이미지 {index}, "
        "자연스러운 사진 스타일, 밝고 선명한 분위기, 텍스트나 로고 없이"
    )


def _repair_empty_image_prompts(content_list, title="", source=""):
    repaired = []
    image_index = 0
    changed = 0
    for item in content_list:
        if item and item[0] == "image":
            image_index += 1
            prompt = item[1] if len(item) > 1 else ""
            if _is_empty_image_prompt(prompt):
                prompt = _fallback_image_prompt(title, source, image_index)
                changed += 1
            repaired.append(("image", prompt))
            continue
        repaired.append(item)
    return repaired, changed


def _usage_number(value):
    try:
        if value is None:
            return 0
        return int(value)
    except (TypeError, ValueError):
        return 0


def _merge_usage_totals(totals, usage):
    for key in ("input_tokens", "output_tokens", "thinking_tokens", "billable_output_tokens", "total_tokens"):
        value = _usage_number((usage or {}).get(key))
        if value:
            totals[key] = _usage_number(totals.get(key)) + value
    return totals


def _safe_image_failures(raw, limit=20):
    if not isinstance(raw, list):
        return []
    failures = []
    for item in raw[:limit]:
        if not isinstance(item, dict):
            continue
        failures.append({
            "index": max(0, _usage_number(item.get("index"))),
            "stage": str(item.get("stage") or "image_completion")[:80],
            "error_code": str(item.get("error_code") or "image_not_inserted")[:80],
            "provider": str(item.get("provider") or "")[:40],
            "method": str(item.get("method") or "")[:40],
            "message": str(item.get("message") or "")[:180],
        })
    return failures


def _image_failure_stage(failures):
    stages = [str(item.get("stage") or "").strip() for item in failures if isinstance(item, dict)]
    for stage in ("image_generation", "image_upload", "image_insert_exception", "image_prompt_empty"):
        if stage in stages:
            return stage
    return "image_completion"


def _image_failure_message(requested, inserted, failures):
    stage = _image_failure_stage(failures)
    stage_labels = {
        "image_generation": "이미지 생성 실패",
        "image_upload": "네이버 이미지 업로드 실패",
        "image_insert_exception": "이미지 삽입 중 예외 발생",
        "image_prompt_empty": "이미지 프롬프트 누락",
        "image_completion": "이미지 첨부 완료 확인 실패",
    }
    codes = sorted({str(item.get("error_code") or "").strip() for item in failures if item.get("error_code")})
    code_text = f" 원인 코드: {', '.join(codes[:3])}." if codes else ""
    return (
        f"{stage_labels.get(stage, '이미지 첨부 실패')}: 요청 {requested}장 중 {inserted}장만 첨부되었습니다."
        f"{code_text} 글이 이미지 없이 저장되지 않도록 작업을 중단했습니다."
    )


def _won_from_usd(usd):
    return int(math.ceil(max(0.0, float(usd or 0)) * _USD_KRW_RATE))


def _calculate_generation_cost(ai_model, usage_totals, image_generated, image_provider_counts=None):
    model = _normalize_ai_model(ai_model)
    price = _AI_TEXT_PRICE_USD_PER_1M.get(model)
    input_tokens = _usage_number((usage_totals or {}).get("input_tokens"))
    output_tokens = _usage_number((usage_totals or {}).get("billable_output_tokens")) or _usage_number((usage_totals or {}).get("output_tokens"))
    image_generated = max(0, int(image_generated or 0))
    provider_counts = image_provider_counts if isinstance(image_provider_counts, dict) else {}
    gemini_images = max(0, _usage_number(provider_counts.get("gemini")))
    openai_images = max(0, _usage_number(provider_counts.get("openai")))
    unattributed_images = max(0, image_generated - gemini_images - openai_images)
    if unattributed_images:
        gemini_images += unattributed_images

    text_usd = 0.0
    if price:
        text_usd = (
            input_tokens * price["input"] / 1_000_000
            + output_tokens * price["output"] / 1_000_000
        )
    image_usd = gemini_images * _GEMINI_IMAGE_PRICE_USD + openai_images * _OPENAI_IMAGE_PRICE_USD
    total_usd = text_usd + image_usd
    return {
        "currency": "KRW",
        "exchange_rate": _USD_KRW_RATE,
        "exchange_rate_label": _USD_KRW_RATE_LABEL,
        "price_available": bool(price),
        "model": model,
        "model_label": (price or {}).get("label", model),
        "input_tokens": input_tokens,
        "output_tokens": _usage_number((usage_totals or {}).get("output_tokens")),
        "thinking_tokens": _usage_number((usage_totals or {}).get("thinking_tokens")),
        "billable_output_tokens": output_tokens,
        "image_generated": image_generated,
        "image_provider_counts": {
            "gemini": gemini_images,
            "openai": openai_images,
        },
        "text_won": _won_from_usd(text_usd),
        "image_won": _won_from_usd(image_usd),
        "total_won": _won_from_usd(total_usd),
    }


def _build_write_result(
    ok,
    success,
    total,
    mode,
    ai_model,
    usage_totals,
    image_totals,
    posts=None,
    failed_posts=None,
    error=None,
    stage=None,
    failed_keyword=None,
):
    image_totals = image_totals or {}
    result = {
        "ok": bool(ok),
        "success": int(success or 0),
        "total": int(total or 0),
        "mode": mode,
        "usage": {
            "input_tokens": _usage_number((usage_totals or {}).get("input_tokens")),
            "output_tokens": _usage_number((usage_totals or {}).get("output_tokens")),
            "thinking_tokens": _usage_number((usage_totals or {}).get("thinking_tokens")),
            "billable_output_tokens": _usage_number((usage_totals or {}).get("billable_output_tokens")) or _usage_number((usage_totals or {}).get("output_tokens")),
            "total_tokens": _usage_number((usage_totals or {}).get("total_tokens")),
        },
        "images": {
            "attempted": _usage_number(image_totals.get("attempted")),
            "generated": _usage_number(image_totals.get("generated")),
            "inserted": _usage_number(image_totals.get("inserted")),
            "provider_counts": image_totals.get("providers") if isinstance(image_totals.get("providers"), dict) else {},
            "failure_count": len(_safe_image_failures(image_totals.get("failures"))),
            "failures": _safe_image_failures(image_totals.get("failures")),
        },
        "posts": (posts or [])[-20:],
    }
    failed_posts = failed_posts or []
    if failed_posts:
        result["failed_posts"] = failed_posts[-20:]
    if stage:
        result["stage"] = str(stage)[:80]
    if failed_keyword:
        result["failed_keyword"] = str(failed_keyword)[:160]
    result["cost"] = _calculate_generation_cost(
        ai_model,
        result["usage"],
        result["images"]["generated"],
        result["images"].get("provider_counts"),
    )
    if error:
        result["error"] = str(error)[:500]
    return result


def _is_real_secret(value):
    value = (value or "").strip()
    return bool(value) and value.lower() not in _PLACEHOLDER_SECRET_VALUES


def _load_settings_data():
    if not os.path.exists(SETTINGS_PATH):
        return {}
    try:
        with open(SETTINGS_PATH, "r", encoding="utf-8") as f:
            data = json.load(f)
        return data if isinstance(data, dict) else {}
    except (json.JSONDecodeError, OSError):
        return {}


def _save_settings_data(data):
    os.makedirs(os.path.dirname(SETTINGS_PATH), exist_ok=True)
    clean_data = data if isinstance(data, dict) else {}
    with open(SETTINGS_PATH, "w", encoding="utf-8") as f:
        json.dump(clean_data, f, ensure_ascii=False, indent=2)


def _load_secret_fallback_data():
    if not os.path.exists(SECRET_FALLBACK_PATH):
        return {}
    try:
        with open(SECRET_FALLBACK_PATH, "r", encoding="utf-8") as f:
            data = json.load(f)
        return data if isinstance(data, dict) else {}
    except (json.JSONDecodeError, OSError):
        return {}


def _save_secret_fallback_data(data):
    try:
        os.makedirs(os.path.dirname(SECRET_FALLBACK_PATH), exist_ok=True)
        tmp_path = SECRET_FALLBACK_PATH + ".tmp"
        fd = os.open(tmp_path, os.O_WRONLY | os.O_CREAT | os.O_TRUNC, 0o600)
        with os.fdopen(fd, "w", encoding="utf-8") as f:
            json.dump(data if isinstance(data, dict) else {}, f, ensure_ascii=False, indent=2)
        os.replace(tmp_path, SECRET_FALLBACK_PATH)
        try:
            os.chmod(SECRET_FALLBACK_PATH, 0o600)
        except OSError:
            pass
    except OSError:
        pass


def _secret_fallback_get(key):
    value = _load_secret_fallback_data().get(key, "")
    return value if _is_real_secret(value) else ""


def _secret_fallback_set(key, value):
    if not _is_real_secret(value):
        return
    data = _load_secret_fallback_data()
    data[key] = value
    _save_secret_fallback_data(data)


def _secret_fallback_delete(key):
    data = _load_secret_fallback_data()
    if key in data:
        data.pop(key, None)
        _save_secret_fallback_data(data)


def _cleared_secret_keys():
    raw = _load_settings_data().get("cleared_secret_keys", [])
    if not isinstance(raw, list):
        return set()
    return {str(item) for item in raw if str(item)}


def _secret_auto_recovery_blocked(key):
    return key in _cleared_secret_keys()


def _mark_secret_cleared(key):
    data = _load_settings_data()
    cleared = _cleared_secret_keys()
    cleared.add(key)
    data["cleared_secret_keys"] = sorted(cleared)
    _save_settings_data(data)


def _unmark_secret_cleared(key):
    data = _load_settings_data()
    cleared = _cleared_secret_keys()
    if key not in cleared:
        return
    cleared.discard(key)
    data["cleared_secret_keys"] = sorted(cleared)
    _save_settings_data(data)


def _mark_keychain_unavailable():
    global _KEYCHAIN_READ_UNAVAILABLE
    _KEYCHAIN_READ_UNAVAILABLE = True
    data = _load_settings_data()
    data["keychain_unavailable"] = True
    _save_settings_data(data)


def _keychain_available():
    if env_truthy("AIMAX_DISABLE_KEYCHAIN") or _KEYCHAIN_READ_UNAVAILABLE:
        return False
    data = _load_settings_data()
    return not bool(data.get("keychain_unavailable"))


def _update_popup_key(version_info):
    info = version_info if isinstance(version_info, dict) else {}
    platform = str(info.get("platform") or sys.platform or "unknown").strip().lower()
    version = str(info.get("latest_version") or info.get("min_version") or "").strip()
    if not version:
        return ""
    return f"{platform}:{version}"


def _is_update_popup_dismissed(version_info):
    key = _update_popup_key(version_info)
    if not key:
        return False
    dismissed = _load_settings_data().get("dismissed_update_popups")
    if isinstance(dismissed, dict):
        return bool(dismissed.get(key))
    if isinstance(dismissed, list):
        return key in {str(item) for item in dismissed}
    return False


def dismiss_update_popup(version_info):
    key = _update_popup_key(version_info)
    if not key:
        return
    data = _load_settings_data()
    dismissed = data.get("dismissed_update_popups")
    if isinstance(dismissed, dict):
        dismissed_list = [str(k) for k, v in dismissed.items() if v]
    elif isinstance(dismissed, list):
        dismissed_list = [str(item) for item in dismissed if str(item)]
    else:
        dismissed_list = []
    if key not in dismissed_list:
        dismissed_list.append(key)
    data["dismissed_update_popups"] = dismissed_list[-30:]
    _save_settings_data(data)


def _keyring_write(action, key, value=None):
    if not _keychain_available():
        return False
    timeout = float(os.environ.get("AIMAX_KEYCHAIN_WRITE_TIMEOUT_SECONDS", "2") or "2")
    result = {"ok": False, "error": ""}

    def _worker():
        try:
            import keyring

            if action == "delete":
                keyring.delete_password(_KR_SERVICE, key)
            else:
                keyring.set_password(_KR_SERVICE, key, value or "")
            result["ok"] = True
        except Exception as error:
            result["error"] = str(error)

    thread = threading.Thread(target=_worker, daemon=True)
    thread.start()
    thread.join(timeout)
    if thread.is_alive():
        _mark_keychain_unavailable()
        return False
    if result.get("error"):
        return False
    return bool(result.get("ok"))


def _keyring_set(key, value):
    _keyring_write("set", key, value)


def _keyring_delete(key):
    _keyring_write("delete", key)


def _keyring_get_password(service, key):
    if not _keychain_available():
        return ""
    timeout = float(os.environ.get("AIMAX_KEYCHAIN_TIMEOUT_SECONDS", "2") or "2")
    result = {"value": "", "error": ""}

    def _worker():
        try:
            import keyring

            result["value"] = keyring.get_password(service, key) or ""
        except Exception as error:
            result["value"] = ""
            result["error"] = str(error)

    thread = threading.Thread(target=_worker, daemon=True)
    thread.start()
    thread.join(timeout)
    if thread.is_alive():
        _mark_keychain_unavailable()
        return ""
    if result.get("error"):
        _mark_keychain_unavailable()
        return ""
    return result.get("value", "") or ""


def _env_secret_value(key):
    names = _SECRET_ENV_KEYS.get(key)
    if not names:
        return ""
    if isinstance(names, str):
        names = (names,)
    for name in names:
        value = os.environ.get(name)
        if value:
            return value
    return ""


def _keyring_get(key):
    env_value = _env_secret_value(key)
    if env_value:
        return env_value
    if _secret_auto_recovery_blocked(key):
        return ""
    fallback_value = _secret_fallback_get(key)
    if fallback_value:
        return fallback_value
    if not _keychain_available():
        return ""
    try:
        value = _keyring_get_password(_KR_SERVICE, key)
        if value:
            _secret_fallback_set(key, value)
            return value
        if env_truthy("AIMAX_DISABLE_LEGACY_KEYCHAIN") or not _keychain_available():
            return ""
        legacy_value = _keyring_get_password(_KR_LEGACY_SERVICE, key)
        if legacy_value:
            _secret_fallback_set(key, legacy_value)
            try:
                threading.Thread(target=_keyring_set, args=(key, legacy_value), daemon=True).start()
            except Exception:
                pass
        return legacy_value
    except Exception:
        return ""


def _keyring_get_password_direct(service, key):
    if env_truthy("AIMAX_DISABLE_KEYCHAIN"):
        return ""
    timeout = float(os.environ.get("AIMAX_KEYCHAIN_RECOVERY_TIMEOUT_SECONDS", "2") or "2")
    result = {"value": "", "error": ""}

    def _worker():
        try:
            import keyring

            result["value"] = keyring.get_password(service, key) or ""
        except Exception as error:
            result["value"] = ""
            result["error"] = str(error)

    thread = threading.Thread(target=_worker, daemon=True)
    thread.start()
    thread.join(timeout)
    if thread.is_alive() or result.get("error"):
        return ""
    return result.get("value", "") or ""


def recover_missing_settings_secrets(existing=None):
    existing = existing if isinstance(existing, dict) else {}
    recovered = {}
    keys = (
        "naver_pw",
        "gemini_api_key",
        "claude_api_key",
        "openai_api_key",
        "apify_api_token",
    )
    for key in keys:
        if _is_real_secret(existing.get(key)):
            continue
        value = _env_secret_value(key)
        if not value and _secret_auto_recovery_blocked(key):
            continue
        value = value or _secret_fallback_get(key)
        if not value:
            value = _keyring_get_password_direct(_KR_SERVICE, key)
        if not value and not env_truthy("AIMAX_DISABLE_LEGACY_KEYCHAIN"):
            value = _keyring_get_password_direct(_KR_LEGACY_SERVICE, key)
        if _is_real_secret(value):
            _secret_fallback_set(key, value)
            recovered[key] = value
    return recovered


def _profile_key(naver_id):
    """내 블로그 프로필 keyring 키 — 네이버 계정별 분리"""
    nid = (naver_id or "").strip().lower() or "_default"
    return f"blog_profile::{nid}"


def save_blog_profile(naver_id, profile_text):
    """내 블로그 프로필을 OS 키체인에 안전 저장 (계정별 분리)"""
    _secret_fallback_set(_profile_key(naver_id), profile_text or "")
    _keyring_set(_profile_key(naver_id), profile_text or "")


def load_blog_profile(naver_id):
    """내 블로그 프로필을 OS 키체인에서 불러옴. 없으면 빈 문자열."""
    return _keyring_get(_profile_key(naver_id)) or ""


def _save_secret_setting(key, value, *, delete_when_blank=False):
    if _is_real_secret(value):
        _unmark_secret_cleared(key)
        _secret_fallback_set(key, value)
        _keyring_set(key, value)
        return
    if not delete_when_blank:
        return
    _secret_fallback_delete(key)
    _keyring_delete(key)
    _mark_secret_cleared(key)


def save_settings(naver_id, naver_pw, api_key, ai_model=_DEFAULT_AI_MODEL, claude_key="", openai_key="", apify_key=""):
    # 민감 정보는 로컬 보안 저장소에만 저장한다. 일반 저장에서 빈 값은 기존 키 보존으로 처리한다.
    _save_secret_setting("naver_pw", naver_pw)
    _save_secret_setting("gemini_api_key", api_key)
    _save_secret_setting("claude_api_key", claude_key)
    _save_secret_setting("openai_api_key", openai_key)
    _save_secret_setting("apify_api_token", apify_key)
    data = _load_settings_data()
    data["naver_id"] = naver_id
    data["ai_model"] = _normalize_ai_model(ai_model)
    _save_settings_data(data)


def save_local_security_settings(naver_id, naver_pw, ai_model=None):
    """Save Naver/local settings only, without touching provider API keys."""
    _save_secret_setting("naver_pw", naver_pw)
    data = _load_settings_data()
    data["naver_id"] = naver_id
    if ai_model is not None:
        data["ai_model"] = _normalize_ai_model(ai_model)
    _save_settings_data(data)


def repair_accidental_provider_clear_markers():
    data = _load_settings_data()
    cleared = _cleared_secret_keys()
    repaired = sorted(cleared.intersection(_PROVIDER_SECRET_KEYS))
    if not repaired:
        return []
    cleared = cleared.difference(_PROVIDER_SECRET_KEYS)
    data["cleared_secret_keys"] = sorted(cleared)
    data["provider_clear_marker_repaired_at"] = datetime.now(timezone.utc).isoformat()
    _save_settings_data(data)
    return repaired


def _safe_legacy_b64_secret(value):
    if not value:
        return ""
    try:
        decoded = base64.b64decode(value).decode()
    except Exception:
        return ""
    return decoded if _is_real_secret(decoded) else ""


def _legacy_plain_secret(data, storage_key, data_key):
    if _secret_auto_recovery_blocked(storage_key):
        return ""
    value = data.get(data_key, "")
    return value if _is_real_secret(value) else ""


def load_settings():
    data = _load_settings_data()
    if not data:
        return "", "", "", _DEFAULT_AI_MODEL, "", "", ""
    repair_accidental_provider_clear_markers()
    naver_id = data.get("naver_id", "")
    ai_model = _normalize_ai_model(data.get("ai_model", _DEFAULT_AI_MODEL))
    # 민감 정보는 키체인에서 복원, 없으면 구버전 base64 필드 폴백
    naver_pw = _keyring_get("naver_pw") or _safe_legacy_b64_secret(data.get("naver_pw", ""))
    api_key = _keyring_get("gemini_api_key") or _legacy_plain_secret(data, "gemini_api_key", "api_key")
    claude_key = _keyring_get("claude_api_key") or _legacy_plain_secret(data, "claude_api_key", "claude_key")
    openai_key = _keyring_get("openai_api_key") or _legacy_plain_secret(data, "openai_api_key", "openai_key")
    apify_key = _keyring_get("apify_api_token") or _legacy_plain_secret(data, "apify_api_token", "apify_key")
    naver_pw = naver_pw if _is_real_secret(naver_pw) else ""
    api_key = api_key if _is_real_secret(api_key) else ""
    claude_key = claude_key if _is_real_secret(claude_key) else ""
    openai_key = openai_key if _is_real_secret(openai_key) else ""
    apify_key = apify_key if _is_real_secret(apify_key) else ""
    return naver_id, naver_pw, api_key, ai_model, claude_key, openai_key, apify_key


def _neighbor_messages_key(naver_id):
    nid = (naver_id or "").strip().lower() or "_default"
    return nid


def save_neighbor_messages(naver_id, messages):
    clean = [m.strip() for m in (messages or []) if m and m.strip()]
    data = _load_settings_data()
    by_account = data.get("neighbor_messages_by_account")
    if not isinstance(by_account, dict):
        by_account = {}
    by_account[_neighbor_messages_key(naver_id)] = clean
    data["neighbor_messages_by_account"] = by_account
    _save_settings_data(data)


def load_neighbor_messages(naver_id):
    data = _load_settings_data()
    by_account = data.get("neighbor_messages_by_account")
    if not isinstance(by_account, dict):
        return []
    messages = by_account.get(_neighbor_messages_key(naver_id), [])
    if not isinstance(messages, list):
        return []
    return [str(m).strip() for m in messages if str(m).strip()]


# ──────────────────────────────────────────────
# 로그를 GUI로 보내는 핸들러
# ──────────────────────────────────────────────
class QueueHandler(logging.Handler):
    def __init__(self, queue):
        super().__init__()
        self.queue = queue

    def emit(self, record):
        msg = self.format(record)
        self.queue.put(("log", msg))


# ──────────────────────────────────────────────
# 색상 정의
# ──────────────────────────────────────────────
COLORS = {
    # ── 라이트 모드 팔레트 (2026-04-24) ──
    "sidebar_bg": "#F1F4F8",          # 사이드바 — 아주 연한 그레이
    "sidebar_hover": "#E3E8EF",
    "sidebar_active": "#FF6B9D",       # 핑크 (active indicator/accent bar)
    "content_bg": "#FFFFFF",          # 메인 영역 — 흰색
    "card_bg": "#FFFFFF",             # 카드 — 흰색
    "card_border": "#D9DDE3",         # 카드 테두리 — 살짝 진하게 해서 3D 느낌
    "card_shadow": "#E8ECEF",         # 카드 밑그림자 (가짜 그림자용)
    "text_primary": "#1F2937",        # 거의 검정
    "text_secondary": "#4B5563",      # 중간 회색
    "text_muted": "#8B95A5",          # 옅은 회색
    "accent": "#FF6B9D",              # 브랜드 핑크
    "accent_hover": "#E85A8A",
    # 네비 버튼 (3D 효과) 전용
    "nav_btn_bg": "#FFFFFF",
    "nav_btn_border": "#DEE3EA",
    "nav_btn_hover_bg": "#F7F9FC",
    "nav_btn_active_bg": "#FFF0F5",    # 아주 연한 핑크
    "nav_btn_active_border": "#FF6B9D",
    "success": "#198754",
    "danger": "#dc3545",
    "terminal_bg": "#0d1117",
    "terminal_fg": "#7ee787",
    "input_bg": "#2b3035",
    "input_fg": "#e9ecef",
    "input_border": "#495057",
}


# ──────────────────────────────────────────────
# 메인 앱
# ──────────────────────────────────────────────
class NaverBlogApp:
    MODE_CONFIG = {
        "all": {
            "title": aimax.APP_VERSION_LABEL,
            "brand": "예리씨",
            "subtitle": "소통·글쓰기 전담",
            "employee_name": "예리씨",
            "employee_initial": "예",
            "avatar_file": "avatar_yeri_circle.png",
            "brand_subtitle": "소통·글쓰기 전담",
            "startup_name": "소통·글쓰기 예리씨",
            "default_panel": "settings",
            "nav_items": [
                ("find_keyword", "고객을 찾아볼게요", "잠재 고객 찾아 서로이웃 신청"),
                ("engage", "고객과 친해질게요", "공감·댓글로 소통"),
                ("write", "고객을 설득할게요", "글을 써서 마음을 움직여요"),
            ],
        },
        "find": {
            "title": "AIMAX-현주씨-영업사원",
            "brand": "현주씨",
            "subtitle": "메이크패밀리 영업사원",
            "employee_name": "현주씨",
            "employee_initial": "현",
            "avatar_file": "avatar_hyunju_circle.png",
            "brand_subtitle": "메이크패밀리 영업사원",
            "startup_name": "영업사원 현주씨",
            "default_panel": "find_keyword",
            "nav_items": [
                ("find_keyword", "고객을 찾아볼게요", "키워드·블로거 링크로 찾기"),
            ],
        },
        "engage_write": {
            "title": "AIMAX-예리씨-블로그글쓰기",
            "brand": "예리씨",
            "subtitle": "소통·글쓰기 전담",
            "employee_name": "예리씨",
            "employee_initial": "예",
            "avatar_file": "avatar_yeri_circle.png",
            "brand_subtitle": "소통·글쓰기 전담",
            "startup_name": "소통·글쓰기 예리씨",
            "default_panel": "engage",
            "nav_items": [
                ("engage", "고객과 친해질게요", "공감·댓글로 소통"),
                ("write", "고객을 설득할게요", "글을 써서 마음을 움직여요"),
            ],
        },
        "engage": {
            "title": "AIMAX-예리씨-블로그글쓰기",
            "brand": "예리씨",
            "subtitle": "소통·글쓰기 전담",
            "employee_name": "예리씨",
            "employee_initial": "예",
            "avatar_file": "avatar_yeri_circle.png",
            "brand_subtitle": "소통·글쓰기 전담",
            "startup_name": "소통·글쓰기 예리씨",
            "default_panel": "engage",
            "nav_items": [
                ("engage", "고객과 친해질게요", "공감·댓글로 소통"),
            ],
        },
        "write": {
            "title": "AIMAX-예리씨-블로그글쓰기",
            "brand": "예리씨",
            "subtitle": "소통·글쓰기 전담",
            "employee_name": "예리씨",
            "employee_initial": "예",
            "avatar_file": "avatar_yeri_circle.png",
            "brand_subtitle": "소통·글쓰기 전담",
            "startup_name": "소통·글쓰기 예리씨",
            "default_panel": "write",
            "nav_items": [
                ("write", "고객을 설득할게요", "글을 써서 마음을 움직여요"),
            ],
        },
    }

    def __init__(self, app_mode=None):
        self.app_mode = self._normalize_app_mode(app_mode or os.environ.get("APP_MODE", "all"))
        self.mode_config = self.MODE_CONFIG[self.app_mode]
        self.root = self._create_root_window()
        self.root.protocol("WM_DELETE_WINDOW", self._on_close)
        if not self._ensure_aimax_consent():
            self.root.destroy()
            raise SystemExit(0)

        self.queue = Queue()
        self.running = False
        self.worker_thread = None
        self.driver = None
        self.stop_event = threading.Event()

        # StringVars
        self.naver_id_var = ttk.StringVar()
        self.naver_pw_var = ttk.StringVar()
        self.api_key_var = ttk.StringVar()
        self.ai_model_var = ttk.StringVar(value=_DEFAULT_AI_MODEL)
        self.claude_key_var = ttk.StringVar()
        self.openai_key_var = ttk.StringVar()
        self.apify_key_var = ttk.StringVar()
        self.web_email_var = ttk.StringVar()
        self.web_password_var = ttk.StringVar()
        self.web_status_var = ttk.StringVar(value="웹앱 연결 안 됨")
        self.last_scraper_csv_path = ""
        self.web_agent_client = None
        self.web_agent_thread = None
        self.web_agent_stop_event = threading.Event()
        self.web_agent_active_job_id = None
        self._shown_update_popup_keys = set()

        # 패널 관리
        self.panels = {}
        self.nav_buttons = {}
        self.current_panel = None
        self.panel_canvases = {}  # panel_key → Canvas (스크롤 라우팅용)
        self._wheel_remainder = 0.0
        # 전역 마우스 휠: 현재 패널의 canvas로 1회만 라우팅
        self.root.bind_all("<MouseWheel>", self._on_mousewheel_global)
        self.root.bind_all("<Button-4>", self._on_mousewheel_global)
        self.root.bind_all("<Button-5>", self._on_mousewheel_global)
        self._bind_touchpad_scroll()

        self._load_saved_settings()
        self._load_web_agent_state()
        self._build_ui()
        self._setup_logging()
        self._poll_queue()

        # 초기 패널 표시 (INIT_PANEL 환경변수로 오버라이드 가능)
        _init_panel = os.environ.get("INIT_PANEL", self.mode_config["default_panel"])
        if not self.naver_id_var.get() or not self.naver_pw_var.get():
            _init_panel = "settings"
        self._show_panel(_init_panel)
        self.root.after(700, self._recover_missing_settings_for_gui)

        # macOS Tk 9.0: bind_all 이벤트 수신에 창 포커스 필수
        if sys.platform == "darwin":
            self.root.after(200, self._focus_root_for_macos_scroll)

        # 시작 환영 메시지
        self._log(f"{aimax.APP_VERSION_LABEL} {self.mode_config.get('startup_name', self.mode_config['brand'])} 직원이 출근했습니다 🟢")
        self._log("좌측 [직원 설정]에서 계정을 등록한 후 작업을 시켜보세요.")
        self._restore_web_agent_session()

        # 첫 실행 또는 계정 미설정 시 온보딩 가이드 표시
        nid, npw, _, _, _, _, _ = load_settings()
        if not nid or not npw:
            self.root.after(300, self._show_onboarding)

    def _normalize_app_mode(self, value):
        value = (value or "all").strip().lower()
        if value in self.MODE_CONFIG:
            return value
        return "all"

    def _mode_allows_panel(self, panel_key):
        if panel_key == "settings" or self.app_mode == "all":
            return True
        if self.app_mode == "find":
            return panel_key in {"find_keyword", "find_link"}
        if self.app_mode == "engage":
            return panel_key == "engage"
        if self.app_mode == "engage_write":
            return panel_key in {"engage", "write"}
        if self.app_mode == "write":
            return panel_key == "write"
        return False

    def _create_root_window(self):
        title = self.mode_config["title"]
        if sys.platform == "darwin":
            root = tk.Tk()
            root.title(title)
            root.geometry("1080x760")
            root.resizable(True, True)
            root.minsize(960, 680)
            self._style = ttk.Style("flatly")
            return root

        root = ttk.Window(
            title=title,
            themename="flatly",
            size=(1080, 760),
            resizable=(True, True),
            iconphoto=None,
        )
        root.minsize(960, 680)
        return root

    def _focus_root_for_macos_scroll(self):
        try:
            focused = self.root.focus_get()
            if focused is not None:
                cls = focused.winfo_class()
                if cls in {"Entry", "TEntry", "Text", "TCombobox", "TSpinbox"}:
                    return
            self.root.lift()
            self.root.focus_force()
        except Exception:
            pass

    def _bind_touchpad_scroll(self):
        try:
            self.root.bind_all("<TouchpadScroll>", self._on_touchpad_scroll_global)
        except tk.TclError:
            pass

    # ── AIMAX 약관/라이선스 ──
    def _ensure_aimax_consent(self):
        if aimax.has_current_consent():
            return True
        return self._show_aimax_consent_dialog()

    def _center_child(self, win, width, height):
        self.root.update_idletasks()
        rx = self.root.winfo_rootx()
        ry = self.root.winfo_rooty()
        rw = max(self.root.winfo_width(), width)
        rh = max(self.root.winfo_height(), height)
        x = rx + (rw - width) // 2
        y = ry + (rh - height) // 2
        win.geometry(f"{width}x{height}+{x}+{y}")

    def _make_readonly_text(self, parent, text, height=18):
        frame = tk.Frame(parent, bg=COLORS["card_bg"])
        frame.pack(fill=BOTH, expand=YES, padx=12, pady=12)
        scrollbar = tk.Scrollbar(frame)
        scrollbar.pack(side=RIGHT, fill=Y)
        text_widget = tk.Text(
            frame, height=height, wrap=WORD, font=(FONT_UI, 9),
            bg="white", fg=COLORS["text_primary"],
            relief="flat", padx=10, pady=8,
            yscrollcommand=scrollbar.set,
        )
        text_widget.insert("1.0", text)
        text_widget.configure(state=DISABLED)
        text_widget.pack(side=LEFT, fill=BOTH, expand=YES)
        scrollbar.configure(command=text_widget.yview)
        return text_widget

    def _show_aimax_consent_dialog(self):
        dlg = tk.Toplevel(self.root)
        dlg.title(f"{aimax.APP_VERSION_LABEL} 이용 동의")
        dlg.configure(bg=COLORS["content_bg"])
        dlg.resizable(True, True)
        dlg.transient(self.root)
        self._center_child(dlg, 820, 680)
        dlg.minsize(760, 620)

        accepted = {"value": False}

        tk.Frame(dlg, bg=COLORS["accent"], height=6).pack(fill=X)
        header = tk.Frame(dlg, bg=COLORS["content_bg"])
        header.pack(fill=X, padx=24, pady=(18, 10))
        tk.Label(
            header, text=aimax.APP_VERSION_LABEL,
            font=(FONT_UI, 18, "bold"),
            bg=COLORS["content_bg"], fg=COLORS["text_primary"], anchor="w",
        ).pack(fill=X)
        tk.Label(
            header,
            text=f"{aimax.SELLER} · {aimax.LICENSE_UNIT} / {aimax.LICENSE_TERM}",
            font=(FONT_UI, 10),
            bg=COLORS["content_bg"], fg=COLORS["text_secondary"], anchor="w",
        ).pack(fill=X, pady=(4, 0))

        notebook = ttk.Notebook(dlg)
        notebook.pack(fill=BOTH, expand=YES, padx=24, pady=(0, 10))
        for title, body in [
            ("이용약관", aimax.TERMS_TEXT),
            ("면책 조항", aimax.DISCLAIMER_TEXT),
            ("개인정보 처리방침", aimax.PRIVACY_TEXT),
        ]:
            tab = tk.Frame(notebook, bg=COLORS["card_bg"])
            notebook.add(tab, text=title)
            self._make_readonly_text(tab, body)

        agree_frame = tk.Frame(dlg, bg=COLORS["content_bg"])
        agree_frame.pack(fill=X, padx=24, pady=(0, 12))

        terms_var = ttk.BooleanVar(value=False)
        disclaimer_var = ttk.BooleanVar(value=False)
        privacy_var = ttk.BooleanVar(value=False)
        start_button = ttk.Button(agree_frame, text="동의하고 시작", bootstyle="primary", state=DISABLED)

        def _refresh_button(*_):
            state = NORMAL if (terms_var.get() and disclaimer_var.get() and privacy_var.get()) else DISABLED
            start_button.configure(state=state)

        for label, var in [
            ("이용약관에 동의합니다", terms_var),
            ("면책 조항을 확인하였으며 동의합니다", disclaimer_var),
            ("개인정보 처리방침에 동의합니다", privacy_var),
        ]:
            ttk.Checkbutton(
                agree_frame, text=label, variable=var,
                command=_refresh_button,
            ).pack(anchor=W, pady=2)

        button_row = tk.Frame(dlg, bg=COLORS["content_bg"])
        button_row.pack(fill=X, padx=24, pady=(0, 18))

        def _decline():
            accepted["value"] = False
            try:
                dlg.grab_release()
            except tk.TclError:
                pass
            dlg.destroy()

        def _accept():
            if not (terms_var.get() and disclaimer_var.get() and privacy_var.get()):
                return
            aimax.accept_current_terms()
            accepted["value"] = True
            try:
                dlg.grab_release()
            except tk.TclError:
                pass
            dlg.destroy()

        start_button.configure(command=_accept)
        ttk.Button(
            button_row, text="동의하지 않음", bootstyle="secondary-outline",
            command=_decline,
        ).pack(side=RIGHT, padx=(8, 0))
        start_button.pack(side=RIGHT)

        dlg.protocol("WM_DELETE_WINDOW", _decline)
        dlg.grab_set()
        self.root.wait_window(dlg)
        return accepted["value"]

    # ── 첫 실행 온보딩 ──
    def _show_onboarding(self):
        # 사용자가 다른 화면에서 입력 중일 때 안내창이 앱 입력을 가로막지 않도록
        # 설정 화면에서만 띄운다. 비밀번호가 비어 있으면 설득/친해질게요 입력칸도
        # 막히던 원인이 여기의 modal grab이었다.
        if self.current_panel != "settings":
            return

        dlg = tk.Toplevel(self.root)
        dlg.title("시작 가이드")
        dlg.configure(bg=COLORS["content_bg"])
        dlg.resizable(False, False)
        dlg.transient(self.root)

        W = 560
        H = 540
        rx = self.root.winfo_x() + (self.root.winfo_width() - W) // 2
        ry = self.root.winfo_y() + (self.root.winfo_height() - H) // 2
        dlg.geometry(f"{W}x{H}+{rx}+{ry}")

        # ── 헤더 ──
        hdr = tk.Frame(dlg, bg=COLORS["accent"], height=6)
        hdr.pack(fill=X)

        tk.Label(
            dlg, text="👋  처음 오셨군요!",
            font=(FONT_UI, 18, "bold"),
            bg=COLORS["content_bg"], fg=COLORS["text_primary"],
        ).pack(pady=(28, 4))
        tk.Label(
            dlg, text="아래 3단계만 따라하면 바로 사용할 수 있어요.",
            font=(FONT_UI, 10),
            bg=COLORS["content_bg"], fg=COLORS["text_secondary"],
        ).pack(pady=(0, 24))

        # ── 단계 카드 ──
        steps = [
            (
                "1단계  |  계정 설정",
                "지금 열려있는 Settings 탭에\n네이버 아이디와 비밀번호를 입력하고\n[설정 저장] 버튼을 눌러주세요.",
                COLORS["accent"],
            ),
            (
                "2단계  |  AI API 키 입력",
                "글 자동 작성에 AI API 키가 필요합니다.\n• Claude: console.anthropic.com → API Keys\n• Gemini: aistudio.google.com → Get API key\n발급한 키를 Settings 탭에 붙여넣고 저장하세요.",
                "#198754",
            ),
            (
                "3단계  |  기능 사용",
                "설정 완료 후 왼쪽 메뉴에서 원하는 기능을 선택하세요.\n• Write — 키워드로 블로그 글 자동 작성·발행\n• Bulk Post — 엑셀로 여러 글 한꺼번에 발행\n• Like / Comment / Neighbor — 이웃 활동 자동화",
                "#fd7e14",
            ),
        ]

        cards_frame = tk.Frame(dlg, bg=COLORS["content_bg"])
        cards_frame.pack(fill=X, padx=28)

        for title, body, accent in steps:
            card = tk.Frame(
                cards_frame, bg=COLORS["card_bg"],
                highlightbackground=accent, highlightthickness=1,
            )
            card.pack(fill=X, pady=5, ipady=10)

            accent_bar = tk.Frame(card, bg=accent, width=4)
            accent_bar.pack(side=LEFT, fill=Y)

            text_frame = tk.Frame(card, bg=COLORS["card_bg"])
            text_frame.pack(side=LEFT, fill=BOTH, expand=True, padx=14, pady=4)

            tk.Label(
                text_frame, text=title,
                font=(FONT_UI, 9, "bold"),
                bg=COLORS["card_bg"], fg=accent, anchor="w",
            ).pack(anchor="w", pady=(4, 2))
            tk.Label(
                text_frame, text=body,
                font=(FONT_UI, 9),
                bg=COLORS["card_bg"], fg=COLORS["text_secondary"],
                anchor="w", justify="left",
            ).pack(anchor="w")

        # ── 버튼 ──
        btn_frame = tk.Frame(dlg, bg=COLORS["content_bg"])
        btn_frame.pack(pady=24)

        def _close_onboarding():
            try:
                dlg.grab_release()
            except tk.TclError:
                pass
            dlg.destroy()
            self._show_panel("settings")

        dlg.protocol("WM_DELETE_WINDOW", _close_onboarding)

        def on_start():
            _close_onboarding()

        ttk.Button(
            btn_frame, text="✓  시작하기 — Settings 열기",
            bootstyle="primary", width=28,
            command=on_start,
        ).pack()

    # ── 설정 로드 ──
    def _load_saved_settings(self):
        nid, npw, akey, ai_model, claude_key, openai_key, apify_key = load_settings()
        self.naver_id_var.set(nid)
        self.naver_pw_var.set(npw)
        self.api_key_var.set(akey)
        self.ai_model_var.set(ai_model)
        self.claude_key_var.set(claude_key)
        self.openai_key_var.set(openai_key)
        self.apify_key_var.set(apify_key)

    def _recover_missing_settings_for_gui(self):
        if getattr(self, "_settings_recovery_started", False):
            return
        snapshot = {
            "naver_pw": self.naver_pw_var.get(),
            "gemini_api_key": self.api_key_var.get(),
            "claude_api_key": self.claude_key_var.get(),
            "openai_api_key": self.openai_key_var.get(),
            "apify_api_token": self.apify_key_var.get(),
        }
        if all(str(value or "").strip() for value in snapshot.values()):
            return
        self._settings_recovery_started = True
        self._log("[설정] 이전 로컬 저장소에 저장된 항목이 있는지 확인합니다.")

        def _worker():
            try:
                recovered = recover_missing_settings_secrets(snapshot) or {}
            except Exception:
                recovered = {}

            def _apply():
                mapping = {
                    "naver_pw": self.naver_pw_var,
                    "gemini_api_key": self.api_key_var,
                    "claude_api_key": self.claude_key_var,
                    "openai_api_key": self.openai_key_var,
                    "apify_api_token": self.apify_key_var,
                }
                restored = 0
                for key, var in mapping.items():
                    value = recovered.get(key)
                    if value and not str(var.get() or "").strip():
                        var.set(value)
                        restored += 1
                if restored:
                    self._log(f"[설정] 이전 로컬 저장소에서 저장된 항목 {restored}개를 복원했습니다.")
                    self._send_immediate_web_agent_heartbeat("settings_recovered")
                else:
                    self._log("[설정] 추가로 복원할 저장 항목을 찾지 못했습니다.")

            try:
                self.root.after(0, _apply)
            except Exception:
                pass

        threading.Thread(target=_worker, daemon=True).start()

    # ── 웹앱 Agent 연결 ──
    def _load_web_agent_state(self):
        try:
            from web_agent.client import load_state
            state = load_state()
            self.web_email_var.set(state.get("email", ""))
        except Exception:
            self.web_email_var.set("")

    def _make_web_agent_client(self, session_token=None):
        from web_agent.client import AimaxWebAgentClient, load_state, load_session_token
        state = load_state()
        token = session_token if session_token is not None else load_session_token()
        return AimaxWebAgentClient(base_url=state.get("base_url"), session_token=token)

    def _set_web_agent_status(self, text, color=None):
        self.web_status_var.set(text)
        if hasattr(self, "web_status_label"):
            self.web_status_label.configure(fg=color or COLORS["text_muted"])

    def _restore_web_agent_session(self):
        try:
            from web_agent.client import load_session_token
            token = load_session_token()
        except Exception:
            token = ""
        if not token:
            return
        try:
            client = self._make_web_agent_client(session_token=token)
            self._start_web_agent_polling(client)
            self._set_web_agent_status("저장된 웹앱 세션 확인 중...", COLORS["text_muted"])
        except Exception as e:
            self._log(f"[웹앱 연결] 저장된 세션 복원 실패: {e}")

    def _connect_web_agent(self):
        email = self.web_email_var.get().strip().lower()
        raw_password = self.web_password_var.get()
        password = raw_password.strip()
        if not email:
            self._log("[웹앱 연결] 이메일과 비밀번호를 입력해주세요.")
            self._set_web_agent_status("이메일과 비밀번호가 필요합니다.", "#C0392B")
            return
        try:
            from web_agent.client import password_input_error

            password_error = password_input_error(raw_password)
        except Exception:
            password_error = ""
        if password_error:
            self.web_password_var.set("")
            self._log(f"[웹앱 연결] 로그인 전 입력 확인: {password_error}")
            self._set_web_agent_status(password_error, "#C0392B")
            return

        if hasattr(self, "web_connect_btn"):
            self.web_connect_btn.configure(state=DISABLED)
        self._set_web_agent_status("웹앱 로그인 중...", COLORS["text_muted"])

        def _worker():
            try:
                from web_agent.client import (
                    current_platform_label,
                    default_device_label,
                    friendly_error_message,
                    save_session_token,
                    save_state,
                )
                client = self._make_web_agent_client(session_token="")
                result = client.login(email, password, device_label=default_device_label())
                token = result.get("session_token", "")
                if not token:
                    raise RuntimeError("로그인은 성공했지만 세션 토큰이 없습니다.")
                token_stored = save_session_token(token)
                save_state(email=email, base_url=client.base_url, device_label=default_device_label())
                self.queue.put(("web_agent_clear_password", None))
                if not token_stored:
                    self.queue.put((
                        "log",
                        "[웹앱 연결] 로그인은 성공했지만 이 PC의 안전 저장소에 세션을 저장하지 못했습니다. "
                        "재시작 후 다시 로그인해야 할 수 있습니다.",
                    ))

                if result.get("requires_password_change") or not result.get("can_execute"):
                    self.queue.put((
                        "web_agent_status",
                        ("첫 로그인 비밀번호 변경이 필요합니다. 웹앱에서 변경 후 다시 연결해주세요.", "#C0392B"),
                    ))
                    self.queue.put(("log", "[웹앱 연결] 로그인 성공. 비밀번호 변경 전에는 작업 실행이 제한됩니다."))
                    return

                self.queue.put(("log", f"[웹앱 연결] 로그인 성공: {email} / {current_platform_label()}"))
                self._start_web_agent_polling(client)
                if not token_stored:
                    self.queue.put((
                        "web_agent_status",
                        (
                            "웹앱 로그인 성공. 다만 안전 저장소에 세션을 저장하지 못해 재시작 후 다시 로그인해야 할 수 있습니다.",
                            "#C77C02",
                        ),
                    ))
            except Exception as e:
                try:
                    message = friendly_error_message(e)
                except Exception:
                    message = str(e)
                self.queue.put(("log", f"[웹앱 연결] 로그인 실패: {message}"))
                self.queue.put(("web_agent_status", (f"로그인 실패. {message}", "#C0392B")))
            finally:
                self.queue.put(("web_agent_controls", None))

        threading.Thread(target=_worker, daemon=True).start()

    def _disconnect_web_agent(self):
        self.web_agent_stop_event.set()
        client = self.web_agent_client
        self.web_agent_client = None
        self.web_agent_active_job_id = None
        try:
            from web_agent.client import clear_session_token
            clear_session_token()
        except Exception:
            pass
        if client:
            def _logout():
                try:
                    client.logout()
                except Exception:
                    pass
            threading.Thread(target=_logout, daemon=True).start()
        self._set_web_agent_status("웹앱 연결 해제됨", COLORS["text_muted"])
        self._log("[웹앱 연결] 연결을 해제했습니다.")

    def _start_web_agent_polling(self, client):
        if self.web_agent_thread and self.web_agent_thread.is_alive():
            self.web_agent_client = client
            self.queue.put(("web_agent_status", ("웹앱 연결됨. 작업 대기 중입니다.", "#198754")))
            return
        self.web_agent_client = client
        self.web_agent_stop_event.clear()
        self.web_agent_thread = threading.Thread(target=self._web_agent_loop, args=(client,), daemon=True)
        self.web_agent_thread.start()

    def _send_immediate_web_agent_heartbeat(self, reason="settings_saved"):
        client = getattr(self, "web_agent_client", None)
        if not client:
            return

        def _worker():
            try:
                from web_agent.client import current_platform_label, default_device_label

                local_status = "busy" if self.running or self.web_agent_active_job_id else "connected"
                client.heartbeat(
                    status=local_status,
                    version=aimax.APP_VERSION,
                    platform_label=current_platform_label(),
                    device_label=default_device_label(),
                    readiness=self._collect_web_agent_readiness(),
                )
                self.queue.put(("log", "[웹앱 연결] 설정 변경 상태를 대시보드에 즉시 반영했습니다."))
            except Exception as error:
                self.queue.put(("log", f"[웹앱 연결] 설정 변경 상태 즉시 반영 실패: {error}"))

        threading.Thread(target=_worker, daemon=True).start()

    def _web_agent_status_value(self, ready):
        return "ready" if ready else "missing"

    def _web_agent_tool_candidates(self, command, env_path=""):
        candidates = []
        configured = (os.environ.get(env_path) or "").strip() if env_path else ""
        if configured:
            candidates.append(configured)
        found = shutil.which(command)
        if found:
            candidates.append(found)
        if command == "yt-dlp" and os.name == "nt":
            local_app_data = (os.environ.get("LOCALAPPDATA") or "").strip()
            if local_app_data:
                candidates.append(os.path.join(
                    local_app_data,
                    "AIMAX",
                    "media-tools",
                    "win32",
                    "x64",
                    "yt-dlp.exe",
                ))
            app_dir = os.path.dirname(os.path.abspath(__file__))
            parent_dir = os.path.dirname(app_dir)
            bundle_dir = getattr(sys, "_MEIPASS", "")
            executable_dir = os.path.dirname(sys.executable) if getattr(sys, "frozen", False) else ""
            for base_dir in (app_dir, parent_dir, bundle_dir, executable_dir):
                if base_dir:
                    candidates.append(os.path.join(
                        base_dir,
                        "oracle",
                        "aimax-reports-api",
                        "vendor",
                        "media-tools",
                        "win32",
                        "x64",
                        "yt-dlp.exe",
                    ))
                    candidates.append(os.path.join(
                        base_dir,
                        "vendor",
                        "media-tools",
                        "win32",
                        "x64",
                        "yt-dlp.exe",
                    ))
        unique_candidates = []
        seen = set()
        for candidate in candidates:
            value = os.path.normpath(str(candidate or "").strip())
            if not value or value in seen:
                continue
            seen.add(value)
            unique_candidates.append(value)
        return unique_candidates

    def _web_agent_resolve_tool(self, command, env_path=""):
        for candidate in self._web_agent_tool_candidates(command, env_path):
            resolved = shutil.which(candidate) or candidate
            if os.path.isfile(resolved):
                return resolved
        return ""

    def _web_agent_tool_status(self, command, version_args=None, env_path=""):
        executable = self._web_agent_resolve_tool(command, env_path)
        if not executable:
            if command == "yt-dlp":
                try:
                    import yt_dlp.version
                    return {"status": "ready", "version": str(getattr(yt_dlp.version, "__version__", "") or "python-module")[:80]}
                except Exception:
                    pass
            return {"status": "missing", "version": ""}
        args = tuple(version_args or ["--version"])
        cache_key = (command, executable, args)
        cache = getattr(self, "_web_agent_tool_status_cache", {})
        now = time.monotonic()
        cached = cache.get(cache_key)
        if cached and now < float(cached.get("expires_at", 0.0) or 0.0):
            return dict(cached.get("value") or {"status": "needs_attention", "version": ""})
        try:
            completed = subprocess.run(
                [executable, *args],
                capture_output=True,
                text=True,
                timeout=8,
                **_hidden_subprocess_kwargs(),
            )
            output = (completed.stdout or completed.stderr or "").splitlines()
            version = output[0].strip() if output else ""
            result = {"status": "ready" if completed.returncode == 0 else "needs_attention", "version": version[:80]}
        except Exception:
            result = {"status": "needs_attention", "version": ""}
        cache[cache_key] = {
            "expires_at": now + (600 if result["status"] == "ready" else 60),
            "value": dict(result),
        }
        self._web_agent_tool_status_cache = cache
        return result

    def _collect_web_agent_diagnostics(self):
        local_state = {}
        polling_state = getattr(self, "_web_agent_polling_diagnostics", {})
        try:
            from local_agent.state_repair import collect_local_state_diagnostics

            data = collect_local_state_diagnostics()
            request_files = []
            for item in data.get("request_files") or []:
                request_files.append({
                    "name": str(item.get("name") or "")[:80],
                    "exists": bool(item.get("exists")),
                    "stale": bool(item.get("stale")),
                    "repair_action": str(item.get("repair_action") or "")[:80],
                    "age_seconds": item.get("age_seconds"),
                })
            lock_file = data.get("lock_file") if isinstance(data.get("lock_file"), dict) else {}
            local_state = {
                "available": True,
                "repair_available": bool(data.get("repair_available")),
                "repair_strategy": str(data.get("repair_strategy") or "")[:80],
                "legacy_candidate_count": int(data.get("legacy_candidate_count") or 0),
                "stale_request_count": int(data.get("stale_request_count") or 0),
                "lock_file_action": str(lock_file.get("repair_action") or "")[:80],
                "request_files": request_files[:5],
            }
        except Exception as error:
            local_state = {
                "available": False,
                "error": str(error)[:200],
            }
        return {
            "app_mode": getattr(self, "app_mode", "all"),
            "version": aimax.APP_VERSION,
            "local_state": local_state,
            "web_agent": {
                "heartbeat_only": env_truthy("AIMAX_AGENT_HEARTBEAT_ONLY"),
                "disable_commands": env_truthy("AIMAX_AGENT_DISABLE_COMMANDS"),
                "disable_jobs": env_truthy("AIMAX_AGENT_DISABLE_JOBS"),
                "running": bool(getattr(self, "running", False)),
                "active_job_id": str(getattr(self, "web_agent_active_job_id", "") or "")[:80],
                "last_next_job_at": str(getattr(self, "web_agent_last_next_job_at", "") or "")[:40],
                "last_next_job_result": str(getattr(self, "web_agent_last_next_job_result", "") or "")[:120],
                "last_next_job_error": str(getattr(self, "web_agent_last_next_job_error", "") or "")[:200],
                "last_command_error": str(getattr(self, "web_agent_last_command_error", "") or "")[:200],
            },
            "polling": {
                "heartbeat_only": bool(polling_state.get("heartbeat_only")),
                "skip_commands": bool(polling_state.get("skip_commands")),
                "skip_jobs": bool(polling_state.get("skip_jobs")),
                "last_next_job_at": str(polling_state.get("last_next_job_at") or "")[:40],
                "last_next_job_status": str(polling_state.get("last_next_job_status") or "")[:40],
                "last_next_job_id": str(polling_state.get("last_next_job_id") or "")[:80],
                "last_next_job_job_status": str(polling_state.get("last_next_job_job_status") or "")[:40],
                "last_next_job_error": str(polling_state.get("last_next_job_error") or "")[:200],
                "active_job_id": str(getattr(self, "web_agent_active_job_id", "") or "")[:80],
            },
        }

    def _queue_update_popup_if_needed(self, version_info):
        info = version_info if isinstance(version_info, dict) else {}
        if not (info.get("update_required") or info.get("update_available")):
            return
        key = _update_popup_key(info)
        if not key:
            return
        if not hasattr(self, "_shown_update_popup_keys"):
            self._shown_update_popup_keys = set()
        if key in self._shown_update_popup_keys or _is_update_popup_dismissed(info):
            return
        self._shown_update_popup_keys.add(key)
        self.queue.put(("web_agent_update_popup", dict(info)))

    def _collect_web_agent_readiness(self):
        naver_id = (self.naver_id_var.get() or "").strip()
        naver_pw = (self.naver_pw_var.get() or "").strip()
        gemini_key = (self.api_key_var.get() or "").strip()
        claude_key = (self.claude_key_var.get() or "").strip()
        openai_key = (self.openai_key_var.get() or "").strip()
        apify_key = (self.apify_key_var.get() or "").strip()
        ai_model = _normalize_ai_model(self.ai_model_var.get() or _DEFAULT_AI_MODEL)
        naver_ready = bool(naver_id and naver_pw)
        if ai_model == "claude":
            selected_ai_ready = bool(claude_key)
        elif _is_openai_model(ai_model):
            selected_ai_ready = bool(openai_key)
        else:
            selected_ai_ready = bool(gemini_key)

        try:
            neighbor_messages = load_neighbor_messages(naver_id)
        except Exception:
            neighbor_messages = []

        yeri_ready = naver_ready and selected_ai_ready
        hyunju_ready = naver_ready and bool(neighbor_messages)
        browser_ready = bool(getattr(self, "driver", None))
        yt_dlp_status = self._web_agent_tool_status("yt-dlp", env_path="AIMAX_SONGI_YTDLP_PATH")

        def worker_status(kind, ready):
            if not self._supports_remote_job(kind):
                return "unavailable"
            return self._web_agent_status_value(ready)

        return {
            "web_login": True,
            "naver_account": {
                "status": self._web_agent_status_value(naver_ready),
                "has_id": bool(naver_id),
                "has_password": bool(naver_pw),
            },
            "ai_keys": {
                "gemini": self._web_agent_status_value(bool(gemini_key)),
                "claude": self._web_agent_status_value(bool(claude_key)),
                "openai": self._web_agent_status_value(bool(openai_key)),
                "apify": self._web_agent_status_value(bool(apify_key)),
                "selected_model": ai_model,
                "selected_model_ready": self._web_agent_status_value(selected_ai_ready),
            },
            "neighbor_messages": {
                "status": self._web_agent_status_value(bool(neighbor_messages)),
                "count": len(neighbor_messages),
            },
            "browser": {
                "status": "ready" if browser_ready else "unknown",
                "last_check_at": None,
            },
            "media_tools": {
                "yt_dlp": yt_dlp_status["status"],
                "yt_dlp_version": yt_dlp_status["version"],
            },
            "workers": {
                "yeri_write": worker_status("yeri_write", yeri_ready),
                "hyunju_find": worker_status("hyunju_find", hyunju_ready),
            },
            "diagnostics": self._collect_web_agent_diagnostics(),
        }

    def _web_agent_loop(self, client):
        from web_agent.client import AimaxApiError, current_platform_label, default_device_label
        import time

        heartbeat_seconds = max(5, int(os.environ.get("AIMAX_AGENT_HEARTBEAT_SECONDS", os.environ.get("AIMAX_AGENT_POLL_SECONDS", "20")) or 20))
        command_poll_seconds = max(2, int(os.environ.get("AIMAX_AGENT_COMMAND_POLL_SECONDS", "5") or 5))
        version_check_seconds = max(60, int(os.environ.get("AIMAX_AGENT_VERSION_CHECK_SECONDS", "600") or 600))
        device_label = default_device_label()
        platform_label = current_platform_label()
        version_status = ""
        version_notice_key = ""
        heartbeat_only = env_truthy("AIMAX_AGENT_HEARTBEAT_ONLY")
        skip_commands = heartbeat_only or env_truthy("AIMAX_AGENT_DISABLE_COMMANDS")
        skip_jobs = heartbeat_only or env_truthy("AIMAX_AGENT_DISABLE_JOBS")
        runner_start_timeout_seconds = max(10, int(os.environ.get("AIMAX_AGENT_RUNNER_START_TIMEOUT_SECONDS", "90") or 90))
        self._web_agent_polling_diagnostics = {
            "heartbeat_only": heartbeat_only,
            "skip_commands": skip_commands,
            "skip_jobs": skip_jobs,
            "last_next_job_at": "",
            "last_next_job_status": "not_polled",
            "last_next_job_id": "",
            "last_next_job_job_status": "",
            "last_next_job_error": "",
        }

        def _refresh_version_status():
            nonlocal version_notice_key
            version_data = client.version(aimax.APP_VERSION, platform_label).get("agent", {})
            next_status = ""
            notice_key = ""
            if version_data.get("update_required"):
                next_status = "필수 업데이트 필요"
                notice_key = f"required:{_update_popup_key(version_data)}"
                if notice_key != version_notice_key:
                    self.queue.put(("log", "[업데이트] 현재 실행기는 필수 업데이트 대상입니다. 웹앱에서 새 파일을 내려받아 설치해주세요."))
            elif version_data.get("update_available"):
                next_status = "업데이트 가능"
                notice_key = f"available:{_update_popup_key(version_data)}"
                if notice_key != version_notice_key:
                    self.queue.put(("log", "[업데이트] 새 실행기 버전이 있습니다. 웹앱에서 업데이트 파일을 내려받을 수 있습니다."))
            version_notice_key = notice_key
            self._queue_update_popup_if_needed(version_data)
            return next_status

        try:
            me = client.me()
            if me.get("requires_password_change") or not me.get("can_execute"):
                self.queue.put((
                    "web_agent_status",
                    ("첫 로그인 비밀번호 변경이 필요합니다. 웹앱에서 변경 후 다시 연결해주세요.", "#C0392B"),
                ))
                return
            version_status = _refresh_version_status()
        except AimaxApiError as e:
            if e.status_code in (401, 403):
                self.queue.put(("web_agent_status", ("웹앱 세션이 만료되었거나 권한 확인이 필요합니다.", "#C0392B")))
                try:
                    from web_agent.client import clear_session_token
                    clear_session_token()
                except Exception:
                    pass
                return
            self.queue.put(("log", f"[웹앱 연결] 초기 확인 실패: {e}"))
        except Exception as e:
            self.queue.put(("log", f"[웹앱 연결] 초기 확인 실패: {e}"))

        self.queue.put(("web_agent_status", ("웹앱 연결됨. 작업 대기 중입니다.", "#198754")))
        next_heartbeat_at = 0.0
        next_version_check_at = time.monotonic() + version_check_seconds
        while not self.web_agent_stop_event.is_set():
            try:
                now = time.monotonic()
                if self.web_agent_active_job_id and not self.running:
                    claimed_at = float(getattr(self, "web_agent_active_job_claimed_at", 0.0) or 0.0)
                    if claimed_at and now - claimed_at >= runner_start_timeout_seconds:
                        stuck_job_id = self.web_agent_active_job_id
                        try:
                            client.update_job(
                                stuck_job_id,
                                "failed",
                                "로컬 실행기가 작업을 받았지만 내부 실행 워커가 시작되지 않았습니다. 실행기를 재시작한 뒤 다시 시도해주세요.",
                                "error",
                                result={
                                    "ok": False,
                                    "stage": "runner_start_timeout",
                                    "error": "local_worker_not_started_after_claim",
                                },
                            )
                        except Exception as update_error:
                            self.queue.put(("log", f"[웹앱 작업] 시작 타임아웃 상태 전송 오류: {update_error}"))
                        self.queue.put(("log", f"[웹앱 작업] 시작 타임아웃으로 작업을 실패 처리했습니다: {stuck_job_id}"))
                        self.web_agent_active_job_id = None
                        self.web_agent_active_job_claimed_at = 0.0

                if now >= next_version_check_at:
                    next_version_check_at = now + version_check_seconds
                    version_status = _refresh_version_status()

                if now >= next_heartbeat_at:
                    next_heartbeat_at = now + heartbeat_seconds
                    local_status = "busy" if self.running or self.web_agent_active_job_id else "connected"
                    client.heartbeat(
                        status=local_status,
                        version=aimax.APP_VERSION,
                        platform_label=platform_label,
                        device_label=device_label,
                        readiness=self._collect_web_agent_readiness(),
                    )
                    if version_status:
                        self.queue.put(("web_agent_status", (f"웹앱 연결됨. {version_status}.", "#C0392B")))
                    elif self.running or self.web_agent_active_job_id:
                        self.queue.put(("web_agent_status", ("웹앱 연결됨. 로컬 작업 실행 중입니다.", "#1E5E8C")))
                    else:
                        self.queue.put(("web_agent_status", ("웹앱 연결됨. 작업 대기 중입니다.", "#198754")))

                if not skip_commands:
                    try:
                        command = client.next_command(platform_label).get("command")
                        self.web_agent_last_command_error = ""
                        if command:
                            self.queue.put(("web_agent_command", {"client": client, "command": command}))
                    except AimaxApiError as command_error:
                        self.web_agent_last_command_error = f"{command_error.status_code}:{command_error.error}"
                        self.queue.put(("log", f"[웹앱 명령] 폴링 오류: {command_error}"))
                    except Exception as command_error:
                        self.web_agent_last_command_error = str(command_error)
                        self.queue.put(("log", f"[웹앱 명령] 폴링 오류: {command_error}"))

                if skip_jobs:
                    self._web_agent_polling_diagnostics.update({
                        "last_next_job_status": "disabled",
                        "last_next_job_error": "",
                    })
                elif not self.running and not self.web_agent_active_job_id:
                    try:
                        self.web_agent_last_next_job_at = datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")
                        self._web_agent_polling_diagnostics.update({
                            "last_next_job_at": self.web_agent_last_next_job_at,
                            "last_next_job_status": "polling",
                            "last_next_job_error": "",
                        })
                        job = client.next_job(platform_label, device_label).get("job")
                        self.web_agent_last_next_job_error = ""
                        self.web_agent_last_next_job_result = f"job:{job.get('id')}" if job else "empty"
                        self._web_agent_polling_diagnostics.update({
                            "last_next_job_status": "job" if job else "empty",
                            "last_next_job_id": str((job or {}).get("id") or "")[:80],
                            "last_next_job_job_status": str((job or {}).get("status") or "")[:40],
                            "last_next_job_error": "",
                        })
                        if job:
                            job_id = job.get("id")
                            self.web_agent_active_job_id = job_id
                            self.web_agent_active_job_claimed_at = time.monotonic()
                            try:
                                client.update_job(
                                    job_id,
                                    "running",
                                    "로컬 실행기가 작업을 받았고 내부 실행 큐에 넣었습니다.",
                                    result={"runner_event": "claimed"},
                                )
                            except Exception as update_error:
                                self.queue.put(("log", f"[웹앱 작업] 수신 확인 전송 오류: {update_error}"))
                            self.queue.put(("remote_job", {"client": client, "job": job}))
                    except AimaxApiError as job_error:
                        self.web_agent_last_next_job_error = f"{job_error.status_code}:{job_error.error}"
                        self.web_agent_last_next_job_result = "error"
                        self._web_agent_polling_diagnostics.update({
                            "last_next_job_status": "error",
                            "last_next_job_error": self.web_agent_last_next_job_error[:200],
                        })
                        self.queue.put(("log", f"[웹앱 작업] 폴링 오류: {job_error}"))
                    except Exception as job_error:
                        self.web_agent_last_next_job_error = str(job_error)
                        self.web_agent_last_next_job_result = "error"
                        self._web_agent_polling_diagnostics.update({
                            "last_next_job_status": "error",
                            "last_next_job_error": self.web_agent_last_next_job_error[:200],
                        })
                        self.queue.put(("log", f"[웹앱 작업] 폴링 오류: {job_error}"))
            except AimaxApiError as e:
                if e.status_code in (401, 403):
                    self.queue.put(("web_agent_status", ("웹앱 세션이 만료되었거나 권한 확인이 필요합니다.", "#C0392B")))
                    try:
                        from web_agent.client import clear_session_token
                        clear_session_token()
                    except Exception:
                        pass
                    return
                self.queue.put(("log", f"[웹앱 연결] 폴링 오류: {e}"))
            except Exception as e:
                self.queue.put(("log", f"[웹앱 연결] 폴링 오류: {e}"))

            self.web_agent_stop_event.wait(command_poll_seconds)

    def _local_provider_secret_values(self):
        return {
            "gemini": (self.api_key_var.get() or "").strip(),
            "apify": (self.apify_key_var.get() or "").strip(),
            "openai": (self.openai_key_var.get() or "").strip(),
            "claude": (self.claude_key_var.get() or "").strip(),
        }

    def _import_local_provider_secrets(self, client, command):
        allowed = ("gemini", "apify", "openai", "claude")
        payload = command.get("payload") if isinstance(command.get("payload"), dict) else {}
        requested = payload.get("providers") if isinstance(payload.get("providers"), list) else allowed
        providers = []
        for provider in requested:
            provider = str(provider or "").strip().lower()
            if provider in allowed and provider not in providers:
                providers.append(provider)
        if not providers:
            providers = list(allowed)

        local_values = self._local_provider_secret_values()
        result = {"type": "local_provider_secret_import", "providers": {}}
        for provider in providers:
            value = local_values.get(provider, "")
            if not _is_real_secret(value):
                result["providers"][provider] = {"status": "missing"}
                continue
            try:
                client.put_user_secret(provider, value)
                result["providers"][provider] = {"status": "imported"}
            except Exception as error:
                result["providers"][provider] = {
                    "status": "failed",
                    "error": str(error)[:160],
                }

        statuses = [item.get("status") for item in result["providers"].values()]
        imported_count = statuses.count("imported")
        missing_count = statuses.count("missing")
        failed_count = statuses.count("failed")
        result["imported_count"] = imported_count
        result["missing_count"] = missing_count
        result["failed_count"] = failed_count
        result["requested_count"] = len(providers)
        log = f"AI/API 키 가져오기 완료: {imported_count}개 저장, {missing_count}개 없음, {failed_count}개 실패"
        return result, log

    def _songi_bounded_int(self, value, fallback, minimum, maximum):
        try:
            number = int(value)
        except Exception:
            number = fallback
        return max(minimum, min(maximum, number))

    def _songi_parse_number(self, value):
        if value in (None, "", "NA"):
            return 0
        try:
            return max(0, float(value))
        except Exception:
            return 0

    def _songi_upload_date_iso(self, value):
        text = str(value or "").strip()
        if not re.match(r"^\d{8}$", text):
            return ""
        try:
            parsed = datetime(int(text[:4]), int(text[4:6]), int(text[6:8]), tzinfo=timezone.utc)
            return parsed.isoformat().replace("+00:00", "Z")
        except Exception:
            return ""

    def _songi_run_ytdlp(self, args, timeout=45):
        ytdlp = self._web_agent_resolve_tool("yt-dlp", env_path="AIMAX_SONGI_YTDLP_PATH")
        if not ytdlp:
            raise RuntimeError("yt-dlp 실행 파일을 찾지 못했습니다.")
        completed = subprocess.run(
            [ytdlp, *args],
            capture_output=True,
            text=True,
            timeout=timeout,
            **_hidden_subprocess_kwargs(),
        )
        if completed.returncode != 0:
            stderr = (completed.stderr or completed.stdout or "").strip().splitlines()
            message = stderr[-1] if stderr else "yt-dlp 실행이 실패했습니다."
            raise RuntimeError(message[:500])
        return completed.stdout or ""

    def _songi_ytdlp_module_version(self):
        try:
            import yt_dlp.version
            return str(getattr(yt_dlp.version, "__version__", "") or "python-module")
        except Exception:
            return ""

    def _songi_youtube_detail_with_module(self, url):
        try:
            import yt_dlp
            with yt_dlp.YoutubeDL({
                "quiet": True,
                "no_warnings": True,
                "skip_download": True,
                "noplaylist": True,
            }) as ydl:
                info = ydl.extract_info(url, download=False) or {}
        except Exception:
            return {}
        return {
            "upload_date": info.get("upload_date"),
            "like_count": self._songi_parse_number(info.get("like_count")),
            "comment_count": self._songi_parse_number(info.get("comment_count")),
            "channel_follower_count": self._songi_parse_number(info.get("channel_follower_count") or info.get("uploader_follower_count")),
            "duration": self._songi_parse_number(info.get("duration")),
            "webpage_url": info.get("webpage_url") or "",
        }

    def _songi_youtube_search_with_module(self, query, max_results):
        try:
            import yt_dlp
        except Exception as error:
            raise RuntimeError("yt-dlp 실행 파일 또는 Python 모듈을 찾지 못했습니다.") from error
        try:
            with yt_dlp.YoutubeDL({
                "quiet": True,
                "no_warnings": True,
                "skip_download": True,
                "extract_flat": True,
                "playlistend": max_results,
            }) as ydl:
                info = ydl.extract_info(query, download=False) or {}
        except Exception as error:
            raise RuntimeError(str(error)[:500]) from error
        return [entry for entry in (info.get("entries") or []) if isinstance(entry, dict)]

    def _songi_youtube_shorts_queries(self, keyword, search_limit):
        clean_keyword = re.sub(r"\s+", " ", str(keyword or "")).strip()
        return [
            f"ytsearch{search_limit}:{clean_keyword} shorts",
            f"ytsearch{search_limit}:{clean_keyword} #shorts",
            f"ytsearch{search_limit}:{clean_keyword} 쇼츠",
            f"ytsearch{search_limit}:{clean_keyword} 숏폼",
        ]

    def _songi_youtube_has_short_hint(self, *values):
        text = " ".join(str(value or "") for value in values).lower()
        return bool(re.search(r"(/shorts/|#shorts|\bshorts\b|shortform|short form|쇼츠|숏폼)", text))

    def _songi_youtube_video_id(self, item):
        video_id = str(item.get("id") or "").strip()
        url = str(item.get("webpage_url") or item.get("url") or "").strip()
        url_match = re.search(r"(?:/shorts/|[?&]v=)([^&/?#]+)", url)
        if url_match:
            video_id = url_match.group(1)
        return video_id[:80]

    def _songi_youtube_candidate_url(self, video_id, url, is_short_form):
        if is_short_form and video_id:
            return f"https://www.youtube.com/shorts/{video_id}"
        if video_id and not str(url or "").startswith("http"):
            return f"https://www.youtube.com/watch?v={video_id}"
        return str(url or "").strip()

    def _songi_youtube_detail(self, url):
        if not self._web_agent_resolve_tool("yt-dlp", env_path="AIMAX_SONGI_YTDLP_PATH"):
            return self._songi_youtube_detail_with_module(url)
        separator = "\x1f"
        template = separator.join([
            "%(upload_date)s",
            "%(like_count)s",
            "%(comment_count)s",
            "%(channel_follower_count)s",
            "%(duration)s",
            "%(webpage_url)s",
        ])
        try:
            stdout = self._songi_run_ytdlp([
                "--skip-download",
                "--no-warnings",
                "--print",
                template,
                url,
            ], timeout=35)
        except Exception:
            return {}
        parts = (stdout.strip().splitlines()[-1] if stdout.strip() else "").split(separator)
        while len(parts) < 6:
            parts.append("")
        return {
            "upload_date": parts[0],
            "like_count": self._songi_parse_number(parts[1]),
            "comment_count": self._songi_parse_number(parts[2]),
            "channel_follower_count": self._songi_parse_number(parts[3]),
            "duration": self._songi_parse_number(parts[4]),
            "webpage_url": parts[5],
        }

    def _songi_youtube_discovery(self, command):
        payload = command.get("payload") if isinstance(command.get("payload"), dict) else {}
        keyword = str(payload.get("keyword") or "").strip()
        if not keyword:
            raise ValueError("찾고 싶은 키워드가 비어 있습니다.")
        max_results = self._songi_bounded_int(payload.get("max_results"), 12, 5, 20)
        days = self._songi_bounded_int(payload.get("date_range_days"), 30, 1, 90)
        search_limit = min(60, max(max_results * 4, max_results + 12))
        queries = self._songi_youtube_shorts_queries(keyword, search_limit)
        use_module = not self._web_agent_resolve_tool("yt-dlp", env_path="AIMAX_SONGI_YTDLP_PATH")
        rows = []
        seen_video_ids = set()
        for query in queries:
            stdout = ""
            search_items = []
            if use_module:
                search_items = self._songi_youtube_search_with_module(query, search_limit)
            else:
                stdout = self._songi_run_ytdlp([
                    "--skip-download",
                    "--flat-playlist",
                    "--no-warnings",
                    "--dump-json",
                    "--playlist-end",
                    str(search_limit),
                    query,
                ], timeout=max(45, search_limit * 4))
                for line in stdout.splitlines():
                    line = line.strip()
                    if not line:
                        continue
                    try:
                        search_items.append(json.loads(line))
                    except Exception:
                        continue
            for item in search_items:
                if not isinstance(item, dict):
                    continue
                video_id = self._songi_youtube_video_id(item)
                if not video_id or video_id in seen_video_ids:
                    continue
                seen_video_ids.add(video_id)
                url = str(item.get("webpage_url") or item.get("url") or "").strip()
                provisional_url = self._songi_youtube_candidate_url(video_id, url, self._songi_youtube_has_short_hint(url))
                if not provisional_url.startswith("http"):
                    continue
                detail = self._songi_youtube_detail(provisional_url)
                duration_seconds = round(self._songi_parse_number(item.get("duration")) or self._songi_parse_number(detail.get("duration")))
                detail_url = str(detail.get("webpage_url") or "").strip()
                is_short_form = (
                    self._songi_youtube_has_short_hint(url, detail_url, item.get("title"), item.get("description"))
                    or (0 < duration_seconds <= 180)
                )
                if duration_seconds > 180 and not self._songi_youtube_has_short_hint(url, detail_url):
                    continue
                if not is_short_form:
                    continue
                url = self._songi_youtube_candidate_url(video_id, detail_url or url, is_short_form)
                if "youtube.com/" not in url:
                    continue
                published_at = self._songi_upload_date_iso(detail.get("upload_date"))
                if published_at:
                    try:
                        published_dt = datetime.fromisoformat(published_at.replace("Z", "+00:00"))
                        age_hours = max(1, (datetime.now(timezone.utc) - published_dt).total_seconds() / 3600)
                    except Exception:
                        age_hours = max(1, days * 24)
                else:
                    age_hours = max(1, days * 24)
                view_count = self._songi_parse_number(item.get("view_count"))
                like_count = self._songi_parse_number(detail.get("like_count"))
                comment_count = self._songi_parse_number(detail.get("comment_count"))
                thumbnails = item.get("thumbnails") if isinstance(item.get("thumbnails"), list) else []
                thumbnail = ""
                if thumbnails:
                    sorted_thumbs = sorted(thumbnails, key=lambda row: int(row.get("width") or 0), reverse=True)
                    thumbnail = str(sorted_thumbs[0].get("url") or "")
                if not thumbnail:
                    thumbnail = f"https://i.ytimg.com/vi/{video_id}/hqdefault.jpg"
                rows.append({
                    "video_id": video_id,
                    "url": url,
                    "title": str(item.get("title") or "").strip()[:180],
                    "creator": str(item.get("channel") or item.get("uploader") or "").strip()[:120],
                    "description": str(item.get("description") or "").strip()[:900],
                    "thumbnail_url": thumbnail,
                    "published_at": published_at,
                    "metrics": {
                        "view_count": round(view_count),
                        "like_count": round(like_count),
                        "comment_count": round(comment_count),
                        "duration_seconds": duration_seconds,
                        "is_short_form": True,
                        "age_hours": round(age_hours, 1),
                        "views_per_hour": round(view_count / age_hours, 1) if age_hours else 0,
                        "engagement_rate": round((like_count + comment_count) / view_count, 4) if view_count else 0,
                        "channel_follower_count": round(self._songi_parse_number(detail.get("channel_follower_count"))),
                    },
                    "content_format": "youtube_shorts",
                    "measurement_badge": "쇼츠 공개 검색",
                })
                if len(rows) >= max_results:
                    break
            if len(rows) >= max_results:
                break
        version_info = self._web_agent_tool_status("yt-dlp", env_path="AIMAX_SONGI_YTDLP_PATH")
        if use_module and not version_info.get("version"):
            version_info["version"] = self._songi_ytdlp_module_version()
        return {
            "ok": True,
            "run_id": payload.get("run_id") or "",
            "project_id": payload.get("project_id") or "",
            "keyword": keyword,
            "platform": "youtube",
            "source_mode": "local_ytdlp",
            "source_version": version_info.get("version") or "",
            "candidates": rows[:max_results],
        }, f"YouTube 쇼츠 후보 {len(rows[:max_results])}개를 찾았습니다."

    def _open_web_requested_local_settings_dialog(self):
        result = {"saved": False}
        dlg = tk.Toplevel(self.root)
        dlg.title("AIMAX 로컬 보안 설정")
        dlg.configure(bg=COLORS["content_bg"])
        dlg.transient(self.root)
        dlg.resizable(True, True)
        self._center_child(dlg, 760, 430)
        dlg.minsize(560, 360)

        outer = tk.Frame(dlg, bg=COLORS["content_bg"])
        outer.pack(fill=BOTH, expand=YES, padx=24, pady=22)
        outer.columnconfigure(1, weight=1)

        tk.Label(
            outer, text="AIMAX 로컬 보안 설정",
            font=(FONT_UI, 16, "bold"),
            bg=COLORS["content_bg"], fg=COLORS["text_primary"], anchor="w",
        ).grid(row=0, column=0, columnspan=2, sticky=EW, pady=(0, 8))
        tk.Label(
            outer,
            text=(
                "네이버 계정과 브라우저 세션만 이 PC에 저장합니다. "
                "Gemini, Claude, OpenAI, Apify 키는 웹 설정 탭의 AI/API 연결에서 관리합니다."
            ),
            font=(FONT_UI, 9),
            bg=COLORS["content_bg"], fg=COLORS["text_secondary"],
            anchor="w", justify="left", wraplength=680,
        ).grid(row=1, column=0, columnspan=2, sticky=EW, pady=(0, 18))

        values = {
            "naver_id": tk.StringVar(value=self.naver_id_var.get() or ""),
            "naver_pw": tk.StringVar(value=self.naver_pw_var.get() or ""),
        }
        local_provider_count = sum(
            1
            for value in (
                self.api_key_var.get(),
                self.claude_key_var.get(),
                self.openai_key_var.get(),
                self.apify_key_var.get(),
            )
            if _is_real_secret(value)
        )
        if local_provider_count:
            initial_status = (
                "이 PC에 저장된 기존 AI/API 키는 삭제하지 않고 유지합니다. "
                "웹 설정의 '기존 실행기 키 가져오기'로 웹 보안 저장소에 옮길 수 있습니다."
            )
        else:
            initial_status = (
                "AI/API 키는 여기서 입력하지 않습니다. "
                "송이와 웹 기반 AI 작업은 웹 설정 탭의 AI/API 연결에서 저장해주세요."
            )
        status_var = tk.StringVar(value=initial_status)

        tk.Label(
            outer, text="네이버 ID", font=(FONT_UI, 9, "bold"),
            bg=COLORS["content_bg"], fg=COLORS["text_secondary"], anchor="w",
        ).grid(row=2, column=0, sticky=W, pady=6, padx=(0, 14))
        ttk.Entry(outer, textvariable=values["naver_id"]).grid(row=2, column=1, sticky=EW, pady=6)

        tk.Label(
            outer, text="네이버 비밀번호", font=(FONT_UI, 9, "bold"),
            bg=COLORS["content_bg"], fg=COLORS["text_secondary"], anchor="w",
        ).grid(row=3, column=0, sticky=W, pady=6, padx=(0, 14))
        pw_entry = ttk.Entry(outer, textvariable=values["naver_pw"], show="*")
        pw_entry.grid(row=3, column=1, sticky=EW, pady=6)
        pw_entry.bind(
            "<FocusIn>",
            lambda _event: status_var.set("비밀번호 입력 전 한/영 상태가 영어인지 확인해주세요. 한글로 입력되면 로그인에 실패할 수 있습니다."),
        )

        tk.Label(
            outer,
            textvariable=status_var,
            font=(FONT_UI, 8),
            bg=COLORS["content_bg"], fg=COLORS["text_muted"],
            anchor="w", justify="left", wraplength=680,
        ).grid(row=4, column=0, columnspan=2, sticky=EW, pady=(16, 10))

        buttons = tk.Frame(outer, bg=COLORS["content_bg"])
        buttons.grid(row=5, column=0, columnspan=2, sticky=EW, pady=(16, 0))

        def _open_api_guide():
            try:
                import webbrowser
                webbrowser.open("https://www.notion.so/makefriends/349b31f1da5581f4b87fc7cbe85ccdb7")
                status_var.set("AI/API 연결 안내를 브라우저로 열었습니다.")
            except Exception as error:
                status_var.set(f"AI/API 연결 안내를 열 수 없습니다: {error}")

        def _cancel():
            dlg.destroy()

        def _save():
            try:
                if str(save_button.cget("state")) == "disabled":
                    return
            except Exception:
                pass
            ai_model = _normalize_ai_model(self.ai_model_var.get() or _DEFAULT_AI_MODEL)
            naver_id = values["naver_id"].get().strip()
            naver_pw = values["naver_pw"].get().strip()
            status_var.set("로컬 보안 설정을 저장하는 중입니다...")
            save_button.configure(state=DISABLED)
            cancel_button.configure(state=DISABLED)
            finish_queue = Queue()

            def _worker():
                error = ""
                try:
                    save_local_security_settings(naver_id, naver_pw, ai_model)
                except Exception as exc:
                    error = str(exc)
                finish_queue.put(error)

            def _poll_finish():
                try:
                    error = finish_queue.get_nowait()
                except Empty:
                    try:
                        self.root.after(50, _poll_finish)
                    except Exception:
                        pass
                    return
                if error:
                    status_var.set(f"저장하지 못했습니다: {error}")
                    save_button.configure(state=NORMAL)
                    cancel_button.configure(state=NORMAL)
                    return
                self.naver_id_var.set(naver_id)
                self.naver_pw_var.set(naver_pw)
                self.ai_model_var.set(ai_model)
                try:
                    self._send_immediate_web_agent_heartbeat("settings_saved")
                except Exception:
                    pass
                result["saved"] = True
                dlg.destroy()

            threading.Thread(target=_worker, daemon=True).start()
            self.root.after(50, _poll_finish)

        ttk.Button(buttons, text="AI/API 연결 안내", bootstyle="secondary-outline", command=_open_api_guide).pack(side=LEFT)
        cancel_button = ttk.Button(buttons, text="취소", bootstyle="secondary-outline", command=_cancel)
        cancel_button.pack(side=RIGHT, padx=(8, 0))
        save_button = ttk.Button(buttons, text="저장", bootstyle="primary", command=_save)
        save_button.pack(side=RIGHT)
        dlg.bind("<Return>", lambda _event: _save())
        dlg.bind("<Escape>", lambda _event: _cancel())
        dlg.protocol("WM_DELETE_WINDOW", _cancel)
        dlg.grab_set()
        self.root.wait_window(dlg)
        return result["saved"]

    def _handle_web_agent_command(self, data):
        client = data.get("client")
        command = data.get("command") or {}
        command_id = command.get("id") or ""
        command_type = command.get("type") or ""

        def _send_command_update(status, log, result=None):
            if not client or not command_id:
                return
            def _worker():
                try:
                    client.update_command(command_id, status, log, result=result)
                except Exception as update_error:
                    self.queue.put(("log", f"[웹앱 연결] 명령 상태 전송 오류: {update_error}"))
            threading.Thread(target=_worker, daemon=True).start()

        try:
            if command_type == "import_local_provider_secrets":
                result, log = self._import_local_provider_secrets(client, command)
                self._log(f"[웹앱 연결] {log}")
                _send_command_update("done", log, result)
                return
            if command_type == "songi_youtube_discovery":
                self._log("[송이] YouTube 쇼츠 키워드 후보 찾기를 시작합니다.")

                def _discovery_worker():
                    try:
                        result, log = self._songi_youtube_discovery(command)
                        self.queue.put(("log", f"[송이] {log}"))
                        _send_command_update("done", log, result)
                    except Exception as error:
                        message = str(error) or "YouTube 쇼츠 후보 찾기에 실패했습니다."
                        payload = command.get("payload") if isinstance(command.get("payload"), dict) else {}
                        self.queue.put(("log", f"[송이] YouTube 쇼츠 후보 찾기 실패: {message}"))
                        _send_command_update("failed", message, {
                            "ok": False,
                            "run_id": payload.get("run_id") or "",
                            "error": "local_ytdlp_discovery_failed",
                            "message": message,
                        })

                threading.Thread(target=_discovery_worker, daemon=True).start()
                return
            if command_type != "open_settings":
                raise ValueError(f"지원하지 않는 웹앱 명령입니다: {command_type}")
            saved = self._open_web_requested_local_settings_dialog()
            self.root.lift()
            self.root.focus_force()
            if saved:
                self._log("[웹앱 연결] 로컬 보안 설정을 이 PC에 저장했습니다.")
                _send_command_update("done", "로컬 보안 설정을 이 PC에 저장했습니다.")
            else:
                self._log("[웹앱 연결] 로컬 보안 설정 저장을 취소했습니다.")
                _send_command_update("failed", "로컬 보안 설정 저장을 취소했습니다.")
        except Exception as e:
            self._log(f"[웹앱 연결] 명령 처리 실패: {e}")
            _send_command_update("failed", str(e))

    def _supports_remote_job(self, kind):
        if self.app_mode == "all":
            return kind in {"yeri_write", "hyunju_find"}
        if self.app_mode in {"engage", "write", "engage_write"}:
            return kind == "yeri_write"
        if self.app_mode == "find":
            return kind == "hyunju_find"
        return False

    def _remote_payload_keywords(self, payload):
        raw = payload.get("keywords")
        if raw is None:
            raw = payload.get("keyword") or payload.get("write_keyword") or payload.get("search_keyword")
        if isinstance(raw, list):
            return [str(item).strip() for item in raw if str(item).strip()]
        if isinstance(raw, str):
            return [item.strip() for item in raw.split(",") if item.strip()]
        return []

    def _remote_job_artifact(self, job):
        artifact = job.get("artifact") if isinstance(job.get("artifact"), dict) else None
        if not artifact:
            return None
        content = str(
            artifact.get("content_markdown")
            or artifact.get("markdown")
            or artifact.get("content")
            or ""
        ).strip()
        if not content:
            return None
        normalized = dict(artifact)
        normalized["content_markdown"] = content
        normalized["title"] = str(artifact.get("title") or "").strip()
        return normalized

    def _remote_artifact_image_count(self, artifact):
        content = str((artifact or {}).get("content_markdown") or "")
        return sum(1 for line in content.splitlines() if line.strip().startswith("[이미지]"))

    def _start_remote_job(self, data):
        client = data.get("client")
        job = data.get("job") or {}
        job_id = job.get("id") or "-"
        kind = job.get("kind") or "-"
        if self.running:
            try:
                client.update_job(job_id, "failed", "로컬 실행기가 이미 다른 작업을 처리 중입니다.", "warning")
            except Exception:
                pass
            self.web_agent_active_job_id = None
            self._log(f"[웹앱 작업] 이미 실행 중이라 작업을 건너뜁니다: {job_id}")
            return
        if not self._supports_remote_job(kind):
            try:
                client.update_job(job_id, "failed", "이 실행기에서 지원하지 않는 직원 작업입니다.", "warning")
            except Exception:
                pass
            self.web_agent_active_job_id = None
            self._log(f"[웹앱 작업] 이 실행기에서 지원하지 않는 작업입니다: {kind}")
            return
        self._log(f"[웹앱 작업] 수신: {kind} ({job_id})")
        self._start_worker(self._worker_remote_job, client=client, job=job)

    def _worker_remote_job(self, client, job):
        job_id = job.get("id") or ""
        kind = job.get("kind") or ""
        payload = job.get("payload") if isinstance(job.get("payload"), dict) else {}
        artifact = self._remote_job_artifact(job) if kind == "yeri_write" else None
        worker_started = False
        try:
            if not self._validate_credentials(need_api=(kind == "yeri_write" and not artifact)):
                raise ValueError("로컬 실행기의 네이버 계정 또는 AI API Key 설정이 필요합니다.")
            artifact_image_count = 0
            if artifact:
                requested_images = _normalize_image_count(payload.get("image_count") or payload.get("images") or 3)
                artifact_image_count = min(self._remote_artifact_image_count(artifact), requested_images)
            if artifact_image_count > 0 and not (self.api_key_var.get() or self.openai_key_var.get()):
                raise ValueError("서버에서 글은 생성되었지만 이미지 생성을 위한 Gemini 또는 OpenAI API Key가 로컬 실행기에 없습니다.")

            client.update_job(job_id, "running", "로컬 실행기가 작업을 시작했습니다.")
            if kind == "yeri_write":
                kwargs = self._remote_write_kwargs(payload, artifact=artifact)
                worker_started = True
                result = self._worker_write(**kwargs)
                status = "cancelled" if self.stop_event.is_set() else "done"
                if isinstance(result, dict) and not result.get("ok", True) and status != "cancelled":
                    client.update_job(job_id, "failed", result.get("error") or "글쓰기 작업이 완료되지 않았습니다.", "error", result=result)
                    return
                client.update_job(job_id, status, "로컬 실행기 작업이 종료되었습니다. 상세 결과는 앱 로그를 확인해주세요.", result=result if isinstance(result, dict) else None)
            elif kind == "hyunju_find":
                kwargs = self._remote_neighbor_kwargs(payload)
                worker_started = True
                self._worker_neighbor(**kwargs)
                status = "cancelled" if self.stop_event.is_set() else "done"
                client.update_job(job_id, status, "로컬 실행기 작업이 종료되었습니다. 상세 결과는 앱 로그를 확인해주세요.")
            else:
                raise ValueError(f"지원하지 않는 웹앱 작업입니다: {kind}")
        except Exception as e:
            self._log(f"[웹앱 작업] 실패: {e}")
            try:
                client.update_job(job_id, "failed", str(e), "error")
            except Exception as update_error:
                self._log(f"[웹앱 작업] 실패 상태 전송 오류: {update_error}")
        finally:
            self.web_agent_active_job_id = None
            if not worker_started:
                self.queue.put(("done", None))

    def _remote_write_kwargs(self, payload, artifact=None):
        keywords = self._remote_payload_keywords(payload)
        md_file = str(payload.get("md_file") or payload.get("markdown_file") or "").strip()
        if artifact and not keywords:
            artifact_title = str(artifact.get("title") or "").strip()
            keywords = [artifact_title or "서버 생성 글"]
        if not keywords and not md_file and not artifact:
            raise ValueError("웹앱 작업에 글쓰기 키워드 또는 MD 파일 경로가 없습니다.")

        mode = str(payload.get("mode") or "publish").strip()
        if mode == "draft":
            mode = "save"
        if mode not in {"publish", "save", "schedule"}:
            mode = "publish"
        word_count = _normalize_target_char_count(payload.get("word_count") or payload.get("wordcount") or 1500)
        return {
            "keywords": keywords,
            "md_file": md_file,
            "mode": mode,
            "style_id": str(payload.get("style_id") or payload.get("style") or "info").strip() or "info",
            "category": str(payload.get("category") or "").strip() or None,
            "cta_link": str(payload.get("cta_link") or "").strip() or None,
            "cta_text": str(payload.get("cta_text") or "").strip() or None,
            "schedule_date": str(payload.get("schedule_date") or "").strip() or None,
            "schedule_hour": str(payload.get("schedule_hour") or "").strip() or None,
            "schedule_interval": str(payload.get("schedule_interval") or "").strip() or None,
            "word_count": word_count,
            "image_count": _normalize_image_count(payload.get("image_count") or payload.get("images") or 3),
            "font_name": str(payload.get("font_name") or "").strip() or None,
            "ai_model": _normalize_ai_model(payload.get("ai_model") or payload.get("model") or self.ai_model_var.get()),
            "seo_brief": payload.get("seo_brief") if isinstance(payload.get("seo_brief"), dict) else None,
            "artifact": artifact if isinstance(artifact, dict) else None,
        }

    def _remote_neighbor_kwargs(self, payload):
        keywords = self._remote_payload_keywords(payload)
        if not keywords:
            raise ValueError("웹앱 작업에 검색 키워드가 없습니다.")
        raw_messages = payload.get("messages") or payload.get("neighbor_messages")
        if isinstance(raw_messages, list):
            messages = [str(item).strip() for item in raw_messages if str(item).strip()]
        elif isinstance(raw_messages, str):
            messages = [item.strip() for item in raw_messages.splitlines() if item.strip()]
        else:
            messages = load_neighbor_messages(self.naver_id_var.get())
        if not messages:
            raise ValueError("서로이웃 신청 멘트가 없습니다. 앱에서 멘트를 저장하거나 웹앱에서 전달해야 합니다.")
        max_per_keyword = payload.get("max_per_keyword") or payload.get("count") or 10
        cooldown_every = payload.get("cooldown_every") or 10
        daily_limit = payload.get("daily_limit") or aimax.RECOMMENDED_LIMITS["daily_neighbor_requests"]
        try:
            max_per_keyword = max(1, int(max_per_keyword))
        except (TypeError, ValueError):
            max_per_keyword = 10
        try:
            cooldown_every = max(0, int(cooldown_every))
        except (TypeError, ValueError):
            cooldown_every = 10
        try:
            daily_limit = max(0, int(daily_limit))
        except (TypeError, ValueError):
            daily_limit = aimax.RECOMMENDED_LIMITS["daily_neighbor_requests"]
        speed_mode = str(payload.get("speed_mode") or payload.get("speed") or "normal").strip()
        if speed_mode not in {"safe", "normal", "fast"}:
            speed_mode = "normal"
        return {
            "keywords": keywords,
            "max_per_keyword": max_per_keyword,
            "messages": messages,
            "speed_mode": speed_mode,
            "cooldown_every": cooldown_every,
            "daily_limit": daily_limit,
        }

    # ── 로깅 설정 ──
    def _setup_logging(self):
        handler = QueueHandler(self.queue)
        handler.setFormatter(logging.Formatter("[%(asctime)s] %(message)s", datefmt="%H:%M:%S"))
        root_logger = logging.getLogger()
        root_logger.addHandler(handler)
        root_logger.setLevel(logging.INFO)
        logging.getLogger("urllib3").setLevel(logging.ERROR)
        logging.getLogger("selenium.webdriver.remote.remote_connection").setLevel(logging.ERROR)

    # ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
    # UI 구성
    # ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
    def _build_ui(self):
        FONT = (FONT_UI, 10)

        # ── 메인 컨테이너 (수평 분할: 사이드바 | 콘텐츠) ──
        main_container = tk.Frame(self.root, bg=COLORS["content_bg"])
        main_container.pack(fill=BOTH, expand=YES)

        # ━━ 좌측 사이드바 ━━
        self.sidebar = tk.Frame(main_container, bg=COLORS["sidebar_bg"], width=200)
        self.sidebar.pack(side=LEFT, fill=Y)
        self.sidebar.pack_propagate(False)

        # 직원 카드
        brand_frame = tk.Frame(self.sidebar, bg=COLORS["sidebar_bg"])
        brand_frame.pack(fill=X, pady=(20, 24))

        # 아바타 (이미지 없을 때 이니셜 원형)
        self._build_employee_avatar(brand_frame)

        tk.Label(
            brand_frame, text=self.mode_config.get("employee_name", self.mode_config["brand"]), font=(FONT_UI, 13, "bold"),
            bg=COLORS["sidebar_bg"], fg=COLORS["text_primary"],
        ).pack(pady=(10, 2))
        tk.Label(
            brand_frame, text=self.mode_config.get("brand_subtitle", self.mode_config["subtitle"]), font=(FONT_UI, 8),
            bg=COLORS["sidebar_bg"], fg=COLORS["text_muted"],
        ).pack()
        # 근무 상태 표시
        self.employee_status_label = tk.Label(
            brand_frame, text="🟢  근무 중", font=(FONT_UI, 8),
            bg=COLORS["sidebar_bg"], fg="#5cb85c",
        )
        self.employee_status_label.pack(pady=(6, 0))

        # 네비게이션 메뉴 — 설정 + 작업 카테고리
        # 구조: [설정 항목들] / (구분선) / [작업 항목들]
        nav_sections = [
            # (section_label, [(key, label, desc), ...])
            (None, [
                ("settings", "직원 설정", "계정·API 키"),
            ]),
            ("작업 지시", self.mode_config["nav_items"]),
        ]

        nav_frame = tk.Frame(self.sidebar, bg=COLORS["sidebar_bg"])
        nav_frame.pack(fill=X, padx=12)

        def _add_nav_item(key, label, desc):
            # 가짜 그림자 wrapper (입체감) — 사이드바보다 살짝 진한 색
            shadow = tk.Frame(nav_frame, bg=COLORS["card_shadow"])
            shadow.pack(fill=X, pady=(0, 8), padx=1)

            # 실제 버튼 (흰색 카드 + 테두리)
            btn = tk.Frame(
                shadow,
                bg=COLORS["nav_btn_bg"],
                cursor="hand2",
                highlightbackground=COLORS["nav_btn_border"],
                highlightthickness=1,
            )
            btn.pack(fill=X, padx=0, pady=(0, 2))  # 아래쪽 2px만 shadow 노출

            # 왼쪽 accent bar (inactive 시 투명)
            accent_bar = tk.Frame(btn, bg=COLORS["nav_btn_bg"], width=4)
            accent_bar.pack(side=LEFT, fill=Y)

            inner = tk.Frame(btn, bg=COLORS["nav_btn_bg"])
            inner.pack(side=LEFT, fill=BOTH, expand=YES)

            lbl = tk.Label(
                inner, text=label, font=(FONT_UI, 10, "bold"),
                bg=COLORS["nav_btn_bg"], fg=COLORS["text_primary"],
                anchor="w", padx=12, pady=9,
            )
            lbl.pack(fill=X)

            sub = tk.Label(
                inner, text=desc, font=(FONT_UI, 8),
                bg=COLORS["nav_btn_bg"], fg=COLORS["text_muted"],
                anchor="w", padx=12,
            )
            sub.pack(fill=X, pady=(0, 7))

            widgets = (btn, accent_bar, inner, lbl, sub)
            for widget in widgets:
                widget.bind("<Button-1>", lambda e, k=key: self._show_panel(k))
                widget.bind("<Enter>", lambda e, k=key: self._nav_hover(k, True))
                widget.bind("<Leave>", lambda e, k=key: self._nav_hover(k, False))

            self.nav_buttons[key] = {
                "shadow": shadow, "btn": btn, "accent_bar": accent_bar,
                "inner": inner, "lbl": lbl, "sub": sub,
            }

        for section_idx, (section_label, items) in enumerate(nav_sections):
            # 섹션 사이 구분선 + 라벨
            if section_idx > 0:
                # 구분선
                tk.Frame(
                    nav_frame, bg=COLORS["card_border"], height=1,
                ).pack(fill=X, padx=12, pady=(14, 8))
                if section_label:
                    tk.Label(
                        nav_frame, text=section_label,
                        font=(FONT_UI, 8),
                        bg=COLORS["sidebar_bg"], fg=COLORS["text_muted"],
                        anchor="w", padx=12,
                    ).pack(fill=X, pady=(0, 4))
            for key, label, desc in items:
                _add_nav_item(key, label, desc)

        # 사이드바 하단 버전
        tk.Label(
            self.sidebar, text=aimax.APP_VERSION_LABEL, font=(FONT_UI, 8),
            bg=COLORS["sidebar_bg"], fg=COLORS["text_muted"],
        ).pack(side=BOTTOM, pady=10)

        # ━━ 우측 콘텐츠 영역 ━━
        right_area = tk.Frame(main_container, bg=COLORS["content_bg"])
        right_area.pack(side=LEFT, fill=BOTH, expand=YES)

        # 콘텐츠 패널 영역
        self.content_area = tk.Frame(right_area, bg=COLORS["content_bg"])
        self.content_area.pack(fill=BOTH, expand=YES)

        # 하단 로그 패널
        self._build_log_panel(right_area)

        # ── 패널 생성 ──
        # split_version에서는 APP_MODE에 맞는 기능 패널만 생성한다.
        self._build_settings_panel(FONT)
        if self.app_mode in ("all", "find"):
            self._build_find_keyword_panel(FONT)
            self._build_find_link_panel(FONT)
        if self.app_mode in ("all", "engage", "engage_write"):
            self._build_engage_panel(FONT)
        if self.app_mode in ("all", "write", "engage_write"):
            self._build_write_panel(FONT)
        # scraper, bulk, like, comment 패널은 lazy 빌드 (현재 노출 안 함)

    # ── 직원 아바타 ──
    def _build_employee_avatar(self, parent):
        """변형별 원형 PNG가 있으면 사용, 없으면 이니셜로 표시."""
        from paths import BUNDLE_DIR
        avatar_file = self.mode_config.get("avatar_file", "avatar_yeri_circle.png")
        fallback_files = [avatar_file, "avatar_yeri_circle.png", "avatar_yeri.png"]
        for file_name in dict.fromkeys(fallback_files):
            avatar_path = BUNDLE_DIR / "assets" / file_name
            if avatar_path.exists():
                try:
                    self._avatar_photo = tk.PhotoImage(file=str(avatar_path))
                    tk.Label(
                        parent, image=self._avatar_photo,
                        bg=COLORS["sidebar_bg"], borderwidth=0,
                    ).pack()
                    return
                except tk.TclError:
                    pass
        # 폴백: 이니셜 원형
        tk.Label(
            parent, text=self.mode_config.get("employee_initial", "예"), font=(FONT_UI, 22, "bold"),
            bg="#FF6B9D", fg="white", width=3, height=1,
        ).pack()

    # ── 네비게이션 ──
    def _apply_nav_style(self, key, state):
        """state: 'normal' | 'hover' | 'active'"""
        if key not in self.nav_buttons:
            return
        parts = self.nav_buttons[key]
        if state == "active":
            bg = COLORS["nav_btn_active_bg"]
            border = COLORS["nav_btn_active_border"]
            accent = COLORS["accent"]
            lbl_fg = COLORS["accent"]
            sub_fg = "#C4508A"
        elif state == "hover":
            bg = COLORS["nav_btn_hover_bg"]
            border = COLORS["nav_btn_border"]
            accent = COLORS["nav_btn_hover_bg"]
            lbl_fg = COLORS["text_primary"]
            sub_fg = COLORS["text_muted"]
        else:  # normal
            bg = COLORS["nav_btn_bg"]
            border = COLORS["nav_btn_border"]
            accent = COLORS["nav_btn_bg"]
            lbl_fg = COLORS["text_primary"]
            sub_fg = COLORS["text_muted"]

        parts["btn"].configure(bg=bg, highlightbackground=border)
        parts["accent_bar"].configure(bg=accent)
        parts["inner"].configure(bg=bg)
        parts["lbl"].configure(bg=bg, fg=lbl_fg)
        parts["sub"].configure(bg=bg, fg=sub_fg)

    def _nav_hover(self, key, entering):
        active_key = self._resolve_nav_key(self.current_panel) if self.current_panel else None
        if key == active_key:
            return  # 활성 버튼은 hover 효과 안 줌
        self._apply_nav_style(key, "hover" if entering else "normal")

    # 자식 패널 → 사이드바에서 활성화할 부모 nav 키 매핑
    PANEL_NAV_PARENT = {
        "find_link": "find_keyword",  # "찾아볼게요" 안의 특정 블로거 링크 탭 (B안 준비중)
        "neighbor": "find_keyword",   # 서이추는 찾아볼게요 체인 안으로 통합됨 (nav 표시용)
        "bulk": "write",               # "설득할게요" 안의 대량발행 (v1 보존)
    }

    def _resolve_nav_key(self, panel_key):
        return self.PANEL_NAV_PARENT.get(panel_key, panel_key)

    def _show_panel(self, key):
        if not self._mode_allows_panel(key):
            key = self.mode_config["default_panel"]

        # 이전 패널 숨기기
        if self.current_panel and self.current_panel in self.panels:
            self.panels[self.current_panel].pack_forget()

        # 이전 네비 버튼 비활성 스타일
        prev_nav = self._resolve_nav_key(self.current_panel) if self.current_panel else None
        if prev_nav and prev_nav in self.nav_buttons:
            self._apply_nav_style(prev_nav, "normal")

        # 새 패널 표시 (없으면 생성)
        self.current_panel = key
        if key not in self.panels:
            self._build_panel_lazy(key)
        if key in self.panels:
            self.panels[key].pack(fill=BOTH, expand=YES)

        # 활성 네비 버튼 스타일 (자식이면 부모 강조)
        nav_key = self._resolve_nav_key(key)
        if nav_key in self.nav_buttons:
            self._apply_nav_style(nav_key, "active")

    def _build_panel_lazy(self, key):
        """nav가 child key를 처음 호출할 때 lazy build"""
        builders = {
            "bulk": self._build_bulk_panel,
            "like": self._build_like_panel,
            "comment": self._build_comment_panel,
        }
        if key in builders:
            builders[key]((FONT_UI, 10))

    # ── 글로벌 마우스 휠 라우터 ──
    # 이전 구현은 패널별 <Enter>/<Leave> bind_all 토글이라 자식 위젯 위로 마우스가
    # 이동할 때 Leave가 터져 스크롤이 꺼지는 버그가 있었다. 단일 전역 핸들러로
    # 현재 활성 패널의 canvas에만 라우팅한다.
    _WHEEL_PASS_CLASSES = {"Text", "Listbox", "Scrollbar", "TCombobox", "TSpinbox", "TScrollbar"}

    def _on_mousewheel_global(self, event):
        try:
            target = event.widget
            if target is not None:
                cls = target.winfo_class() if hasattr(target, "winfo_class") else ""
                if cls in self._WHEEL_PASS_CLASSES:
                    return
            canvas = self.panel_canvases.get(self.current_panel)
            if canvas is None:
                return
            delta = self._wheel_units(event)
            if delta != 0:
                canvas.yview_scroll(delta, "units")
                return "break"
        except Exception:
            pass

    def _on_touchpad_scroll_global(self, event):
        try:
            target = event.widget
            if target is not None:
                cls = target.winfo_class() if hasattr(target, "winfo_class") else ""
                if cls in self._WHEEL_PASS_CLASSES:
                    return

            canvas = self.panel_canvases.get(self.current_panel)
            if canvas is None:
                return

            delta_x, delta_y = self._touchpad_deltas(event)
            if delta_y:
                self._canvas_scroll_pixels(canvas, -delta_y)
                return "break"
            if delta_x:
                self._canvas_scroll_pixels(canvas, -delta_x)
                return "break"
        except Exception:
            pass

    def _wheel_units(self, event):
        num = getattr(event, "num", None)
        if num == 4:
            self._wheel_remainder = 0.0
            return -1
        if num == 5:
            self._wheel_remainder = 0.0
            return 1

        delta = float(getattr(event, "delta", 0) or 0)
        if delta == 0:
            return 0

        if sys.platform == "darwin":
            # macOS trackpad(Tk 9)는 ±1 전후의 float을 보내지만,
            # 일부 외장 마우스는 Windows처럼 ±120을 보내므로 큰 값만 정규화한다.
            raw = -delta / 120 if abs(delta) >= 100 else -delta
            whole = int(raw)
            frac = raw - whole
            self._wheel_remainder += frac
            if abs(self._wheel_remainder) >= 1:
                extra = int(self._wheel_remainder)
                whole += extra
                self._wheel_remainder -= extra
            return whole

        self._wheel_remainder = 0.0
        units = int(-delta / 120)
        if units == 0:
            units = -1 if delta > 0 else 1
        return units

    def _touchpad_deltas(self, event):
        packed_delta = getattr(event, "delta", 0)
        try:
            delta_x, delta_y = self.root.tk.call("tk::PreciseScrollDeltas", packed_delta)
            return int(delta_x), int(delta_y)
        except Exception:
            return 0, int(packed_delta or 0)

    def _canvas_scroll_pixels(self, canvas, pixels):
        if pixels == 0:
            return
        try:
            region = canvas.cget("scrollregion")
            if region:
                x1, y1, x2, y2 = [float(v) for v in region.split()]
            else:
                bbox = canvas.bbox("all")
                if not bbox:
                    return
                x1, y1, x2, y2 = [float(v) for v in bbox]
            content_height = max(1.0, y2 - y1)
            first, last = canvas.yview()
            max_first = max(0.0, 1.0 - (last - first))
            next_first = min(max(first + (pixels / content_height), 0.0), max_first)
            canvas.yview_moveto(next_first)
        except Exception:
            fallback_units = 1 if pixels > 0 else -1
            canvas.yview_scroll(fallback_units, "units")

    # ── 단계 완료 팝업 ──
    # 단계 완료 시 tkinter 모달로 결과 요약 + 다음 단계 이동 제안
    STAGE_NEXT = {
        "find_keyword": ("engage", "고객과 친해질게요"),
        "find_link": ("engage", "고객과 친해질게요"),
        "engage": ("write", "고객을 설득할게요"),
        "write": (None, None),  # 마지막 단계
    }

    def _show_update_popup(self, version_info):
        """실행기 업데이트 안내 팝업. 버전별로 다시 보지 않기를 저장한다."""
        info = version_info if isinstance(version_info, dict) else {}
        if _is_update_popup_dismissed(info):
            return
        try:
            parent = getattr(self, "root", None)
            owns_root = parent is None
            popup = tk.Tk() if owns_root else tk.Toplevel(parent)
            required = bool(info.get("update_required"))
            current_version = str(info.get("current_version") or aimax.APP_VERSION)
            latest_version = str(info.get("latest_version") or "-")
            min_version = str(info.get("min_version") or "-")
            popup.title("실행기 필수 업데이트" if required else "실행기 업데이트")
            popup.configure(bg=COLORS["card_bg"])
            popup.resizable(False, False)
            if not owns_root:
                popup.transient(parent)
                popup.grab_set()

            w, h = 480, 300
            popup.update_idletasks()
            try:
                if parent is not None and parent.winfo_exists():
                    parent.update_idletasks()
                    rx = parent.winfo_rootx()
                    ry = parent.winfo_rooty()
                    rw = parent.winfo_width()
                    rh = parent.winfo_height()
                    x = rx + (rw - w) // 2
                    y = ry + (rh - h) // 2
                else:
                    x = (popup.winfo_screenwidth() - w) // 2
                    y = (popup.winfo_screenheight() - h) // 2
            except Exception:
                x = (popup.winfo_screenwidth() - w) // 2
                y = (popup.winfo_screenheight() - h) // 2
            popup.geometry(f"{w}x{h}+{max(0, x)}+{max(0, y)}")

            accent = "#C0392B" if required else "#1E5E8C"
            title = "필수 업데이트가 필요합니다" if required else "새 실행기 업데이트가 있습니다"
            body = (
                "현재 실행기 버전에서는 웹앱 작업이 제한될 수 있습니다.\n"
                "웹앱의 업데이트 탭에서 새 설치 파일을 내려받아 설치해주세요."
                if required
                else "더 안정적인 실행기 버전이 준비되어 있습니다.\n웹앱의 업데이트 탭에서 새 설치 파일을 내려받아 설치해주세요."
            )

            tk.Label(
                popup, text="업데이트", font=(FONT_UI, 10, "bold"),
                bg=COLORS["card_bg"], fg=accent,
            ).pack(pady=(22, 4))
            tk.Label(
                popup, text=title, font=(FONT_UI, 15, "bold"),
                bg=COLORS["card_bg"], fg=COLORS["text_primary"],
            ).pack(pady=(0, 10))
            tk.Label(
                popup,
                text=body,
                font=(FONT_UI, 10),
                bg=COLORS["card_bg"],
                fg=COLORS["text_secondary"],
                justify="center",
                wraplength=420,
            ).pack(pady=(0, 12))
            version_text = f"현재 {current_version}  /  최신 {latest_version}  /  최소 지원 {min_version}"
            tk.Label(
                popup,
                text=version_text,
                font=(FONT_UI, 9),
                bg=COLORS["card_bg"],
                fg=COLORS["text_muted"],
            ).pack(pady=(0, 12))

            dont_show_var = tk.BooleanVar(master=popup, value=False)
            tk.Checkbutton(
                popup,
                text="다시 보지 않기",
                variable=dont_show_var,
                font=(FONT_UI, 9),
                bg=COLORS["card_bg"],
                fg=COLORS["text_secondary"],
                activebackground=COLORS["card_bg"],
                anchor="w",
            ).pack()

            def _close():
                if dont_show_var.get():
                    dismiss_update_popup(info)
                popup.destroy()

            btn = tk.Button(
                popup,
                text="확인",
                command=_close,
                font=(FONT_UI, 10, "bold"),
                bg=accent,
                fg="white",
                activebackground=accent,
                activeforeground="white",
                relief="flat",
                padx=24,
                pady=8,
                cursor="hand2",
            )
            btn.pack(pady=(14, 18))
            popup.protocol("WM_DELETE_WINDOW", _close)
            try:
                popup.attributes("-topmost", True)
                popup.after(250, lambda: popup.attributes("-topmost", False))
            except Exception:
                pass
            popup.focus_force()
            if owns_root:
                popup.mainloop()
        except Exception as e:
            self._log(f"[업데이트 팝업 오류] {e}")
            self._log("[업데이트] 새 실행기 버전이 있습니다. 웹앱에서 업데이트 파일을 내려받아 설치해주세요.")

    def _show_stage_completion_popup(self, stage_key, title, body, next_stage=None):
        """작업 완료 팝업 — 대표님 호칭 + 다음 단계 이동 제안"""
        try:
            from tkinter import Toplevel
            popup = Toplevel(self.root)
            popup.title("작업 완료")
            popup.configure(bg=COLORS["card_bg"])
            popup.transient(self.root)
            popup.grab_set()

            # 크기 + 중앙 정렬
            w, h = 440, 260
            self.root.update_idletasks()
            rx = self.root.winfo_rootx()
            ry = self.root.winfo_rooty()
            rw = self.root.winfo_width()
            rh = self.root.winfo_height()
            x = rx + (rw - w) // 2
            y = ry + (rh - h) // 2
            popup.geometry(f"{w}x{h}+{x}+{y}")
            popup.resizable(False, False)

            # 체크마크 + 타이틀
            tk.Label(
                popup, text="✓", font=(FONT_UI, 36, "bold"),
                bg=COLORS["card_bg"], fg="#5cb85c",
            ).pack(pady=(18, 0))
            tk.Label(
                popup, text=title,
                font=(FONT_UI, 13, "bold"),
                bg=COLORS["card_bg"], fg=COLORS["text_primary"],
            ).pack(pady=(4, 2))
            tk.Label(
                popup, text="— 대표님 —",
                font=(FONT_UI, 9),
                bg=COLORS["card_bg"], fg=COLORS["text_muted"],
            ).pack()

            # 본문 (결과 요약)
            tk.Label(
                popup, text=body,
                font=(FONT_UI, 10),
                bg=COLORS["card_bg"], fg=COLORS["text_secondary"],
                justify="center", wraplength=400,
            ).pack(pady=(12, 10))

            # 버튼들
            btn_frame = tk.Frame(popup, bg=COLORS["card_bg"])
            btn_frame.pack(pady=(4, 14))

            next_panel, next_label = self.STAGE_NEXT.get(stage_key, (None, None))
            if next_panel:
                def _go_next():
                    popup.destroy()
                    self._show_panel(next_panel)
                ttk.Button(
                    btn_frame, text=f"네, {next_label}로 이동할게요",
                    bootstyle="primary",
                    command=_go_next,
                ).pack(side=LEFT, padx=4)
                ttk.Button(
                    btn_frame, text="나중에",
                    bootstyle="secondary-outline",
                    command=popup.destroy,
                ).pack(side=LEFT, padx=4)
            else:
                # 마지막 단계 — 종료 버튼만
                ttk.Button(
                    btn_frame, text=f"오늘도 수고했어요, {self.mode_config.get('employee_name', '예리씨')}!",
                    bootstyle="primary",
                    command=popup.destroy,
                ).pack()
        except Exception as e:
            # 팝업 실패 시 로그로만 알림
            self._log(f"[완료] {title} — {body}")
            self._log(f"[팝업 오류] {e}")

    # ── 패널 헤더 생성 유틸 ──
    def _make_panel(self, key, title, subtitle, tabs=None):
        """
        tabs: [(child_key, label), ...] 형태. 주어지면 헤더 아래 탭 바를 그린다.
              현재 패널 key가 활성으로 표시됨.
        """
        panel = tk.Frame(self.content_area, bg=COLORS["content_bg"])
        self.panels[key] = panel

        # 헤더
        header = tk.Frame(panel, bg=COLORS["content_bg"])
        header.pack(fill=X, padx=30, pady=(25, 5))
        tk.Label(
            header, text=title, font=(FONT_UI, 18, "bold"),
            bg=COLORS["content_bg"], fg=COLORS["text_primary"], anchor="w",
        ).pack(fill=X)
        tk.Label(
            header, text=subtitle, font=(FONT_UI, 10),
            bg=COLORS["content_bg"], fg=COLORS["text_muted"], anchor="w",
        ).pack(fill=X, pady=(2, 0))

        # 서브 탭 바 (옵션)
        if tabs:
            tab_bar = tk.Frame(panel, bg=COLORS["content_bg"])
            tab_bar.pack(fill=X, padx=30, pady=(12, 0))
            for child_key, label in tabs:
                is_active = child_key == key
                tab_bg = COLORS["sidebar_active"] if is_active else COLORS["card_bg"]
                tab_fg = "white" if is_active else COLORS["text_secondary"]
                tab_btn = tk.Label(
                    tab_bar, text=f"  {label}  ",
                    font=(FONT_UI, 10, "bold" if is_active else "normal"),
                    bg=tab_bg, fg=tab_fg, padx=14, pady=8, cursor="hand2",
                    borderwidth=1, relief="flat",
                )
                tab_btn.pack(side=LEFT, padx=(0, 6))
                tab_btn.bind("<Button-1>", lambda e, k=child_key: self._show_panel(k))

        # 구분선
        tk.Frame(panel, bg=COLORS["card_border"], height=1).pack(fill=X, padx=30, pady=(12, 0))

        # 스크롤 가능 콘텐츠 영역
        canvas = tk.Canvas(panel, bg=COLORS["content_bg"], highlightthickness=0)
        scrollbar = tk.Scrollbar(panel, orient=tk.VERTICAL, command=canvas.yview)
        content = tk.Frame(canvas, bg=COLORS["content_bg"])

        content.bind("<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
        # content 너비를 canvas 너비에 맞춤 (가로 스크롤 방지 + 좌우 정렬)
        content_window = canvas.create_window((0, 0), window=content, anchor="nw")
        canvas.bind("<Configure>", lambda e, w=content_window: canvas.itemconfigure(w, width=e.width))
        canvas.configure(yscrollcommand=scrollbar.set)

        canvas.pack(side=LEFT, fill=BOTH, expand=YES, padx=(30, 0), pady=15)
        scrollbar.pack(side=RIGHT, fill=Y, pady=15)

        # 패널 canvas 등록
        self.panel_canvases[key] = canvas

        return panel, content

    # ── 카드 프레임 유틸 ──
    def _make_card(self, parent):
        # 가짜 그림자 wrapper (입체감)
        shadow = tk.Frame(parent, bg=COLORS["card_shadow"])
        shadow.pack(fill=X, pady=(0, 14))

        card = tk.Frame(
            shadow, bg=COLORS["card_bg"],
            highlightbackground=COLORS["card_border"], highlightthickness=1,
        )
        card.pack(fill=X, padx=0, pady=(0, 2))  # 아래 2px shadow 노출

        inner = tk.Frame(card, bg=COLORS["card_bg"])
        inner.pack(fill=X, padx=20, pady=15)
        return inner

    # ── 입력 필드 유틸 ──
    def _make_field(self, parent, label_text, var, row, show="", width=40, font=None, secret=False):
        font = font or (FONT_UI, 10)
        tk.Label(
            parent, text=label_text, font=(FONT_UI, 9, "bold"),
            bg=COLORS["card_bg"], fg=COLORS["text_secondary"], anchor="w",
        ).grid(row=row, column=0, sticky=W, pady=(0, 8), padx=(0, 15))

        if secret:
            wrapper = tk.Frame(parent, bg=COLORS["card_bg"])
            wrapper.grid(row=row, column=1, sticky=EW, pady=(0, 8))
            parent.columnconfigure(1, weight=1)
            wrapper.columnconfigure(0, weight=1)

            entry = ttk.Entry(wrapper, textvariable=var, font=font, show="*")
            entry.grid(row=0, column=0, sticky=EW)

            _visible = [False]

            def _toggle(e=None, _entry=entry, _flag=_visible, _btn=None):
                _flag[0] = not _flag[0]
                _entry.config(show="" if _flag[0] else "*")
                if _btn:
                    _btn.config(text="🙈" if _flag[0] else "👁")

            eye_btn = tk.Label(
                wrapper, text="👁", font=(FONT_UI, 12),
                bg=COLORS["card_bg"], fg=COLORS["text_muted"],
                cursor="hand2", padx=6,
            )
            eye_btn.grid(row=0, column=1, sticky=W)
            eye_btn.bind("<Button-1>", lambda e, b=eye_btn: _toggle(_btn=b))
            return entry

        entry = ttk.Entry(parent, textvariable=var, font=font, width=width)
        if show:
            entry.config(show=show)
        entry.grid(row=row, column=1, sticky=EW, pady=(0, 8))
        return entry

    # ── 라벨 유틸 ──
    def _make_hint(self, parent, text, row, col=1):
        tk.Label(
            parent, text=text, font=(FONT_UI, 8),
            bg=COLORS["card_bg"], fg=COLORS["text_muted"], anchor="w",
        ).grid(row=row, column=col, sticky=W, pady=(0, 8))

    def _build_aimax_info_card(self, parent):
        record = aimax.load_consent_record()
        card = self._make_card(parent)
        card.columnconfigure(1, weight=1)

        tk.Label(
            card, text="AIMAX 라이선스 및 사업자 정보",
            font=(FONT_UI, 10, "bold"),
            bg=COLORS["card_bg"], fg=COLORS["text_primary"], anchor="w",
        ).grid(row=0, column=0, columnspan=2, sticky=W, pady=(0, 10))

        rows = [
            ("로컬 라이선스 ID", record.get("license_id", "-")),
            ("사용권", f"{record.get('license_unit', aimax.LICENSE_UNIT)} / {record.get('license_term', aimax.LICENSE_TERM)}"),
            ("약관 버전", record.get("terms_version", "-")),
            ("동의 시각", record.get("consented_at", "-")),
        ]
        for idx, (label, value) in enumerate(rows, start=1):
            tk.Label(
                card, text=label, font=(FONT_UI, 9, "bold"),
                bg=COLORS["card_bg"], fg=COLORS["text_secondary"], anchor="w",
            ).grid(row=idx, column=0, sticky=W, pady=(0, 6), padx=(0, 15))
            tk.Label(
                card, text=value, font=(FONT_UI, 9),
                bg=COLORS["card_bg"], fg=COLORS["text_primary"], anchor="w",
                wraplength=560, justify="left",
            ).grid(row=idx, column=1, sticky=W, pady=(0, 6))

        btn_row = tk.Frame(card, bg=COLORS["card_bg"])
        btn_row.grid(row=len(rows) + 1, column=0, columnspan=2, sticky=W, pady=(8, 0))
        ttk.Button(
            btn_row, text="약관·사업자 정보 보기",
            bootstyle="secondary-outline",
            command=self._show_aimax_info_dialog,
        ).pack(side=LEFT, padx=(0, 8))
        ttk.Button(
            btn_row, text="카카오 문의 채널 열기",
            bootstyle="info-outline",
            command=self._open_kakao_contact,
        ).pack(side=LEFT)

    def _show_aimax_info_dialog(self):
        dlg = tk.Toplevel(self.root)
        dlg.title("AIMAX 약관·사업자 정보")
        dlg.configure(bg=COLORS["content_bg"])
        dlg.transient(self.root)
        dlg.resizable(True, True)
        self._center_child(dlg, 820, 650)
        dlg.minsize(760, 580)

        tk.Label(
            dlg, text="AIMAX 라이선스 및 사업자 정보",
            font=(FONT_UI, 16, "bold"),
            bg=COLORS["content_bg"], fg=COLORS["text_primary"], anchor="w",
        ).pack(fill=X, padx=24, pady=(18, 8))

        notebook = ttk.Notebook(dlg)
        notebook.pack(fill=BOTH, expand=YES, padx=24, pady=(0, 12))
        tabs = [
            ("사업자 정보", aimax.business_info_text()),
            ("라이선스", aimax.license_info_text()),
            ("이용약관", aimax.TERMS_TEXT),
            ("면책 조항", aimax.DISCLAIMER_TEXT),
            ("개인정보 처리방침", aimax.PRIVACY_TEXT),
        ]
        for title, body in tabs:
            tab = tk.Frame(notebook, bg=COLORS["card_bg"])
            notebook.add(tab, text=title)
            self._make_readonly_text(tab, body)

        ttk.Button(
            dlg, text="닫기", bootstyle="primary",
            command=dlg.destroy,
        ).pack(pady=(0, 18))

    def _open_kakao_contact(self):
        try:
            import webbrowser
            webbrowser.open(aimax.CONTACT_URL)
            self._log(f"카카오 문의 채널 열기: {aimax.CONTACT_URL}")
        except Exception as e:
            self._log(f"[오류] 문의 채널 열기 실패: {e}")

    def _confirm_limit_override(self, item_name, configured_value, recommended_value):
        if configured_value != 0 and configured_value <= recommended_value:
            return True

        dlg = tk.Toplevel(self.root)
        dlg.title("권장 한도 초과 확인")
        dlg.configure(bg=COLORS["content_bg"])
        dlg.transient(self.root)
        dlg.resizable(False, False)
        self._center_child(dlg, 560, 330)

        confirmed = {"value": False}
        check_var = ttk.BooleanVar(value=False)

        tk.Frame(dlg, bg="#E6B800", height=6).pack(fill=X)
        tk.Label(
            dlg, text="권장 한도를 초과했습니다",
            font=(FONT_UI, 15, "bold"),
            bg=COLORS["content_bg"], fg=COLORS["text_primary"],
        ).pack(anchor=W, padx=24, pady=(22, 8))
        body = (
            f"{item_name} 설정값({configured_value if configured_value else '무제한'})이 "
            f"권장값 {recommended_value}회를 초과했습니다.\n\n"
            "권장 한도 초과 시 외부 플랫폼 정책에 따라 계정 제재, 검색 노출 제한, "
            "활동 제한 등이 발생할 수 있습니다.\n\n"
            "한도 변경 결과의 책임은 사용자에게 있습니다."
        )
        tk.Label(
            dlg, text=body, font=(FONT_UI, 10),
            bg=COLORS["content_bg"], fg=COLORS["text_secondary"],
            justify="left", wraplength=500,
        ).pack(anchor=W, padx=24, pady=(0, 12))

        proceed_btn = ttk.Button(dlg, text="확인 후 진행", bootstyle="warning", state=DISABLED)

        def _refresh():
            proceed_btn.configure(state=NORMAL if check_var.get() else DISABLED)

        ttk.Checkbutton(
            dlg, text="위 위험과 책임을 확인했습니다",
            variable=check_var, command=_refresh,
        ).pack(anchor=W, padx=24, pady=(0, 14))

        btn_row = tk.Frame(dlg, bg=COLORS["content_bg"])
        btn_row.pack(fill=X, padx=24, pady=(0, 20))

        def _cancel():
            confirmed["value"] = False
            try:
                dlg.grab_release()
            except tk.TclError:
                pass
            dlg.destroy()

        def _proceed():
            aimax.log_limit_override(
                item_name=item_name,
                configured_value=int(configured_value),
                recommended_value=int(recommended_value),
                app_mode=self.app_mode,
            )
            confirmed["value"] = True
            try:
                dlg.grab_release()
            except tk.TclError:
                pass
            dlg.destroy()

        proceed_btn.configure(command=_proceed)
        ttk.Button(btn_row, text="취소", bootstyle="secondary-outline", command=_cancel).pack(side=RIGHT, padx=(8, 0))
        proceed_btn.pack(side=RIGHT)
        dlg.protocol("WM_DELETE_WINDOW", _cancel)
        dlg.grab_set()
        self.root.wait_window(dlg)
        return confirmed["value"]

    # ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
    # 패널 빌드
    # ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━

    def _build_settings_panel(self, font):
        employee_name = self.mode_config.get("employee_name", "예리씨")
        panel, content = self._make_panel(
            "settings",
            "직원 설정",
            f"{employee_name}가 일할 네이버 계정과 이 PC용 AI 키를 등록해 주세요",
        )

        card = self._make_card(content)
        card.columnconfigure(1, weight=1)

        self._make_field(card, "네이버 ID", self.naver_id_var, 0, font=font)
        self._make_field(card, "비밀번호", self.naver_pw_var, 1, secret=True, font=font)
        self._make_field(card, "Gemini API Key", self.api_key_var, 2, secret=True, font=font)
        self._make_field(card, "Claude API Key", self.claude_key_var, 3, secret=True, font=font)
        self._make_field(card, "OpenAI API Key", self.openai_key_var, 4, secret=True, font=font)
        self._make_field(card, "Apify API Token", self.apify_key_var, 5, secret=True, font=font)

        tk.Label(
            card,
            text=(
                "이 AI/API 키는 이 PC의 로컬 실행기 작업용입니다. "
                "송이 자료조사용 Gemini/Apify/OpenAI/Claude 키는 웹 설정 탭의 AI/API 연결에서 관리하세요."
            ),
            font=(FONT_UI, 8),
            bg=COLORS["card_bg"], fg=COLORS["text_muted"],
            anchor="w", wraplength=720, justify="left",
        ).grid(row=6, column=0, columnspan=2, sticky=EW, pady=(2, 10))

        # AI 모델 선택
        tk.Label(
            card, text="글 생성 AI", font=(FONT_UI, 9, "bold"),
            bg=COLORS["card_bg"], fg=COLORS["text_secondary"], anchor="w",
        ).grid(row=7, column=0, sticky=W, pady=(0, 8), padx=(0, 15))
        _AI_MODELS = [
            ("gemini-2.5-pro",         "Gemini 2.5 Pro  (~44원/글)  ★ 기본"),
            ("gemini-2.5-flash",       "Gemini 2.5 Flash  (~11원/글)"),
            ("gpt-5.4-mini",           "GPT-5.4 mini  (~21원/글)"),
            ("gpt-5-mini",             "GPT-5 mini  (~9원/글)"),
            ("claude",                 "Claude Sonnet  (~70원/글)"),
        ]
        self.ai_model_var.set(_normalize_ai_model(self.ai_model_var.get()))

        model_frame = tk.Frame(card, bg=COLORS["card_bg"])
        model_frame.grid(row=7, column=1, sticky=EW, pady=(0, 8))
        model_frame.columnconfigure(0, weight=1)

        model_rows = {}

        def _select_model(model_id):
            self.ai_model_var.set(_normalize_ai_model(model_id))

        def _refresh_model_rows(*_):
            current = _normalize_ai_model(self.ai_model_var.get())
            if current != self.ai_model_var.get():
                self.ai_model_var.set(current)
                return
            for model_id, parts in model_rows.items():
                selected = model_id == current
                icon = parts["icon"]
                text = parts["text"]
                row = parts["row"]
                fg = COLORS["text_primary"] if selected else COLORS["text_secondary"]
                icon.configure(
                    text="●" if selected else "○",
                    fg=COLORS["accent"] if selected else COLORS["text_muted"],
                )
                text.configure(fg=fg)
                row.configure(bg=COLORS["card_bg"])

        for row_idx, (model_id, label) in enumerate(_AI_MODELS):
            row = tk.Frame(model_frame, bg=COLORS["card_bg"], cursor="hand2")
            row.grid(row=row_idx, column=0, sticky=W, pady=2)
            icon = tk.Label(
                row,
                text="○",
                font=(FONT_UI, 13, "bold"),
                bg=COLORS["card_bg"],
                fg=COLORS["text_muted"],
                width=2,
                cursor="hand2",
            )
            icon.pack(side=LEFT)
            text = tk.Label(
                row,
                text=label,
                font=(FONT_UI, 9),
                bg=COLORS["card_bg"],
                fg=COLORS["text_secondary"],
                cursor="hand2",
            )
            text.pack(side=LEFT)
            for widget in (row, icon, text):
                widget.bind("<Button-1>", lambda e, value=model_id: _select_model(value))
            model_rows[model_id] = {"row": row, "icon": icon, "text": text}

        self.ai_model_var.trace_add("write", _refresh_model_rows)
        _refresh_model_rows()

        # ── 웹앱 연결 카드 (웹에서 작업 지시 → 이 로컬 실행기가 처리) ──
        web_card = self._make_card(content)
        web_card.columnconfigure(1, weight=1)

        tk.Label(
            web_card, text="웹앱 연결",
            font=(FONT_UI, 10, "bold"),
            bg=COLORS["card_bg"], fg=COLORS["text_primary"], anchor="w",
        ).grid(row=0, column=0, columnspan=2, sticky=W, pady=(0, 4))
        tk.Label(
            web_card,
            text=(
                "웹앱에서 로그인한 계정과 이 실행기를 연결합니다. "
                "연결되면 웹앱에서 만든 작업을 이 PC의 네이버/브라우저 환경으로 처리할 수 있습니다."
            ),
            font=(FONT_UI, 9),
            bg=COLORS["card_bg"], fg=COLORS["text_secondary"],
            anchor="w", wraplength=720, justify="left",
        ).grid(row=1, column=0, columnspan=2, sticky=EW, pady=(0, 10))

        try:
            from web_agent.client import PASSWORD_INPUT_HINT
        except Exception:
            PASSWORD_INPUT_HINT = "비밀번호는 영문 입력 상태에서 입력해주세요. 한글로 입력된 값은 사용할 수 없습니다."

        self._make_field(web_card, "웹앱 이메일", self.web_email_var, 2, font=font)
        web_password_entry = self._make_field(web_card, "웹앱 비밀번호", self.web_password_var, 3, secret=True, font=font)
        web_password_entry.bind(
            "<FocusIn>",
            lambda _event: self._set_web_agent_status(
                PASSWORD_INPUT_HINT,
                COLORS["text_muted"],
            ),
        )

        tk.Label(
            web_card, text="연결 상태", font=(FONT_UI, 9, "bold"),
            bg=COLORS["card_bg"], fg=COLORS["text_secondary"], anchor="w",
        ).grid(row=4, column=0, sticky=W, pady=(0, 8), padx=(0, 15))
        self.web_status_label = tk.Label(
            web_card, textvariable=self.web_status_var, font=(FONT_UI, 9),
            bg=COLORS["card_bg"], fg=COLORS["text_muted"], anchor="w",
            justify="left", wraplength=560,
        )
        self.web_status_label.grid(row=4, column=1, sticky=EW, pady=(0, 8))

        web_btns = tk.Frame(web_card, bg=COLORS["card_bg"])
        web_btns.grid(row=5, column=1, sticky=W, pady=(2, 0))
        self.web_connect_btn = ttk.Button(
            web_btns, text="웹앱 로그인/연결", bootstyle="primary",
            command=self._connect_web_agent, width=18,
        )
        self.web_connect_btn.pack(side=LEFT, padx=(0, 8))
        self.web_disconnect_btn = ttk.Button(
            web_btns, text="연결 해제", bootstyle="secondary-outline",
            command=self._disconnect_web_agent, width=12,
        )
        self.web_disconnect_btn.pack(side=LEFT)

        # ── 내 블로그 프로필 카드 (직원이 멘트 생성할 때 참고) ──
        profile_card = self._make_card(content)
        profile_card.columnconfigure(0, weight=1)

        tk.Label(
            profile_card, text="✨ 내 블로그 프로필",
            font=(FONT_UI, 10, "bold"),
            bg=COLORS["card_bg"], fg=COLORS["text_primary"], anchor="w",
        ).grid(row=0, column=0, sticky=W, pady=(0, 4))
        tk.Label(
            profile_card,
            text=(
                f"{employee_name}가 서로이웃 신청 멘트를 쓸 때 참고합니다. "
                "대표님 블로그를 간단히 소개해 주세요."
            ),
            font=(FONT_UI, 9),
            bg=COLORS["card_bg"], fg=COLORS["text_secondary"],
            anchor="w", wraplength=720, justify="left",
        ).grid(row=1, column=0, sticky=W, pady=(0, 8))

        # 학습 중 안내 배지 (Phase 2 자동 분석 예고)
        learn_badge = tk.Frame(
            profile_card, bg="#E8F4FD",
            highlightbackground="#B3DDF7", highlightthickness=1,
        )
        learn_badge.grid(row=2, column=0, sticky=EW, pady=(0, 10))
        tk.Label(
            learn_badge,
            text=(
                f"🎓  {employee_name}가 블로그를 스스로 분석하는 기능은 다음 업데이트에 만나요!  "
                "지금은 대표님이 직접 적어주시면 그대로 배울게요."
            ),
            font=(FONT_UI, 8),
            bg="#E8F4FD", fg="#1E5E8C",
            anchor="w", padx=10, pady=7, justify="left",
        ).pack(fill=X)

        self.blog_profile_text = tk.Text(
            profile_card, height=7, font=font, wrap="word",
            bg="white", fg=COLORS["text_primary"],
            highlightbackground=COLORS["card_border"], highlightthickness=1,
            relief="flat", padx=10, pady=8,
        )
        self.blog_profile_text.grid(row=3, column=0, sticky=EW)

        _PROFILE_PLACEHOLDER = (
            "예시) 저는 경제적 자유를 주제로 쓰는 블로거입니다. "
            "재테크와 사이드잡에 관심이 많고, 실제 경험담 위주로 진솔하게 씁니다. "
            "친근하고 따뜻한 톤을 선호해요."
        )
        # 현재 네이버 ID로 저장된 프로필 로드
        _initial_profile = load_blog_profile(self.naver_id_var.get())
        if _initial_profile:
            self.blog_profile_text.insert("1.0", _initial_profile)
        else:
            # placeholder 동작
            self.blog_profile_text.insert("1.0", _PROFILE_PLACEHOLDER)
            self.blog_profile_text.configure(fg=COLORS["text_muted"])

            def _on_focus_in(_e):
                if self.blog_profile_text.get("1.0", "end-1c").strip() == _PROFILE_PLACEHOLDER:
                    self.blog_profile_text.delete("1.0", "end")
                    self.blog_profile_text.configure(fg=COLORS["text_primary"])

            def _on_focus_out(_e):
                if not self.blog_profile_text.get("1.0", "end-1c").strip():
                    self.blog_profile_text.insert("1.0", _PROFILE_PLACEHOLDER)
                    self.blog_profile_text.configure(fg=COLORS["text_muted"])

            self.blog_profile_text.bind("<FocusIn>", _on_focus_in)
            self.blog_profile_text.bind("<FocusOut>", _on_focus_out)

        # 네이버 ID가 바뀌면 자동으로 해당 계정 프로필 로드
        def _on_naver_id_change(*_):
            new_profile = load_blog_profile(self.naver_id_var.get())
            self.blog_profile_text.delete("1.0", "end")
            if new_profile:
                self.blog_profile_text.insert("1.0", new_profile)
                self.blog_profile_text.configure(fg=COLORS["text_primary"])
            else:
                self.blog_profile_text.insert("1.0", _PROFILE_PLACEHOLDER)
                self.blog_profile_text.configure(fg=COLORS["text_muted"])
        self.naver_id_var.trace_add("write", _on_naver_id_change)

        # ── API 비용 참고 카드 (라이트 모드) ──
        cost_card = tk.Frame(
            content, bg="#F0F7FB",
            highlightbackground="#B3DDF7", highlightthickness=1,
        )
        cost_card.pack(fill=X, pady=(4, 8))
        cost_inner = tk.Frame(cost_card, bg="#F0F7FB")
        cost_inner.pack(fill=X, padx=16, pady=10)

        tk.Label(
            cost_inner,
            text="💰  API 비용 참고  (블로그 글 1편 기준 · 참고용)",
            font=(FONT_UI, 8, "bold"),
            bg="#F0F7FB", fg="#1E5E8C", anchor="w",
        ).pack(fill=X)

        tk.Frame(cost_inner, bg="#B3DDF7", height=1).pack(fill=X, pady=(5, 6))

        cost_text = (
            "[ 글 생성 — 텍스트, 모델별 실측 ]\n"
            "  • Gemini 2.5 Pro            ~44원/글  ★ 기본\n"
            "  • Gemini 2.5 Flash          ~11원/글\n"
            "  • GPT-5.4 mini              ~21원/글\n"
            "  • GPT-5 mini                ~9원/글\n"
            "  • Claude Sonnet             ~70원/글\n"
            "\n"
            "[ 이미지 생성 — 선택 모델 기준, 1024px 정사각형 ]\n"
            "  • Gemini 이미지 1장        약 58원\n"
            "  • OpenAI 이미지 1장        약 62원 (medium 기준)\n"
            "\n"
            "환율 1,476원/$ 기준 · 완료 후 실제 토큰 사용량으로 원화 비용을 다시 계산합니다"
        )
        tk.Label(
            cost_inner,
            text=cost_text,
            font=(FONT_UI, 8),
            bg="#F0F7FB", fg="#4A7A95",
            justify=LEFT, anchor="w",
        ).pack(fill=X)

        self._build_aimax_info_card(content)

        # 저장 버튼
        btn_frame = tk.Frame(content, bg=COLORS["content_bg"])
        btn_frame.pack(fill=X, pady=(5, 0))

        def _open_api_key_guide():
            try:
                import webbrowser

                webbrowser.open(API_KEY_GUIDE_URL)
                self._log("API 키 발급 가이드를 열었습니다.")
            except Exception as error:
                self._log(f"[오류] API 키 발급 가이드를 열 수 없습니다: {error}")

        ttk.Button(
            btn_frame, text="설정 저장", bootstyle="info",
            command=self._save_credentials, width=15,
        ).pack(side=LEFT)
        ttk.Button(
            btn_frame, text="API 키 발급 가이드", bootstyle="secondary-outline",
            command=_open_api_key_guide,
        ).pack(side=LEFT, padx=(8, 0))

    def _build_write_panel(self, font):
        panel, content = self._make_panel(
            "write",
            "고객을 설득할게요",
            "예리가 글로 고객 마음을 움직여 판매로 이어가요",
        )

        self.write_keyword_var = ttk.StringVar()
        self.write_file_var = ttk.StringVar()
        self.write_mode_var = ttk.StringVar(value="publish")
        self.write_style_var = ttk.StringVar(value="info")
        self.write_wordcount_var = ttk.StringVar(value="1500")
        self.write_font_var = ttk.StringVar(value="기본값")
        self.write_category_var = ttk.StringVar()
        self.cta_link_var = ttk.StringVar()
        self.cta_text_var = ttk.StringVar()
        self.schedule_date_var = ttk.StringVar()
        self.schedule_hour_var = ttk.StringVar()
        self.schedule_interval_var = ttk.StringVar()

        # 콘텐츠 설정 카드
        card = self._make_card(content)
        card.columnconfigure(1, weight=1)

        tk.Label(
            card, text="글 형식", font=(FONT_UI, 9, "bold"),
            bg=COLORS["card_bg"], fg=COLORS["text_secondary"], anchor="w",
        ).grid(row=0, column=0, sticky=W, pady=(0, 8), padx=(0, 15))
        style_combo = ttk.Combobox(
            card, textvariable=self.write_style_var, font=font, width=30, state="readonly",
            values=["info - 정보성 글", "buy - 구매성 글", "ad - 광고성 글(파워컨텐츠)"]
        )
        style_combo.grid(row=0, column=1, sticky=W, pady=(0, 8))
        style_combo.current(0)

        # 글 분량
        tk.Label(
            card, text="글 분량", font=(FONT_UI, 9, "bold"),
            bg=COLORS["card_bg"], fg=COLORS["text_secondary"], anchor="w",
        ).grid(row=1, column=0, sticky=W, pady=(0, 8), padx=(0, 15))
        wordcount_combo = ttk.Combobox(
            card, textvariable=self.write_wordcount_var, font=font, width=30, state="readonly",
            values=["800자 - 짧게", "1500자 - 보통", "2500자 - 길게"]
        )
        wordcount_combo.grid(row=1, column=1, sticky=W, pady=(0, 8))
        wordcount_combo.current(1)

        self._make_field(card, "키워드 (AI 생성)", self.write_keyword_var, 2, font=font)
        self._make_hint(card, "쉼표로 여러 키워드 입력 가능 (예: 키워드1, 키워드2, 키워드3)", 2, col=2)
        self._make_field(card, "또는 MD 파일", self.write_file_var, 3, width=30, font=font)

        ttk.Button(
            card, text="찾아보기", bootstyle="outline",
            command=lambda: self._browse_file(self.write_file_var, [("Markdown", "*.md")])
        ).grid(row=3, column=2, padx=(8, 0), pady=(0, 8))

        self._make_field(card, "카테고리", self.write_category_var, 4, width=20, font=font)
        self._make_hint(card, "블로그 카테고리 이름 (비우면 기본)", 4, col=2)

        # 에디터 서식 카드
        card_fmt = self._make_card(content)
        card_fmt.columnconfigure(1, weight=1)

        tk.Label(
            card_fmt, text="에디터 서식", font=(FONT_UI, 10, "bold"),
            bg=COLORS["card_bg"], fg=COLORS["text_primary"], anchor="w",
        ).grid(row=0, column=0, columnspan=3, sticky=W, pady=(0, 10))

        tk.Label(
            card_fmt, text="글꼴", font=(FONT_UI, 9, "bold"),
            bg=COLORS["card_bg"], fg=COLORS["text_secondary"], anchor="w",
        ).grid(row=1, column=0, sticky=W, pady=(0, 8), padx=(0, 15))
        font_combo = ttk.Combobox(
            card_fmt, textvariable=self.write_font_var, font=font, width=30, state="readonly",
            values=["기본값", "나눔고딕", "나눔명조", "나눔스퀘어", "나눔바른고딕", "마루부리"]
        )
        font_combo.grid(row=1, column=1, sticky=W, pady=(0, 8))
        font_combo.current(0)

        # CTA 카드
        card2 = self._make_card(content)
        card2.columnconfigure(1, weight=1)

        tk.Label(
            card2, text="CTA (선택)", font=(FONT_UI, 10, "bold"),
            bg=COLORS["card_bg"], fg=COLORS["text_primary"], anchor="w",
        ).grid(row=0, column=0, columnspan=3, sticky=W, pady=(0, 10))

        self._make_field(card2, "CTA 링크", self.cta_link_var, 1, font=font)
        self._make_field(card2, "CTA 문구", self.cta_text_var, 2, font=font)
        self._make_hint(card2, "카톡 상담방, 제품 구매 페이지 등 (비우면 CTA 없음)", 3, col=0)

        # 발행 방식
        mode_frame = tk.Frame(content, bg=COLORS["content_bg"])
        mode_frame.pack(fill=X, pady=(0, 8))
        tk.Label(
            mode_frame, text="발행 방식", font=(FONT_UI, 9, "bold"),
            bg=COLORS["content_bg"], fg=COLORS["text_secondary"],
        ).pack(side=LEFT, padx=(0, 12))
        for text, val in [("즉시 발행", "publish"), ("예약 발행", "schedule"), ("임시 저장", "save")]:
            ttk.Radiobutton(mode_frame, text=text, variable=self.write_mode_var, value=val).pack(side=LEFT, padx=5)

        # 예약 일시
        schedule_frame = tk.Frame(content, bg=COLORS["content_bg"])
        schedule_frame.pack(fill=X, pady=(0, 10))
        tk.Label(
            schedule_frame, text="예약 일시", font=(FONT_UI, 9, "bold"),
            bg=COLORS["content_bg"], fg=COLORS["text_secondary"],
        ).pack(side=LEFT, padx=(0, 12))
        ttk.Entry(schedule_frame, textvariable=self.schedule_date_var, font=font, width=12).pack(side=LEFT)
        tk.Label(schedule_frame, text="(YYYY-MM-DD)", font=(FONT_UI, 8), bg=COLORS["content_bg"], fg=COLORS["text_muted"]).pack(side=LEFT, padx=(3, 10))
        ttk.Spinbox(schedule_frame, from_=0, to=23, textvariable=self.schedule_hour_var, font=font, width=4).pack(side=LEFT)
        tk.Label(schedule_frame, text="시", font=(FONT_UI, 8), bg=COLORS["content_bg"], fg=COLORS["text_muted"]).pack(side=LEFT, padx=(3, 10))
        tk.Label(schedule_frame, text="간격", font=(FONT_UI, 9, "bold"), bg=COLORS["content_bg"], fg=COLORS["text_secondary"]).pack(side=LEFT, padx=(0, 4))
        ttk.Spinbox(schedule_frame, from_=1, to=72, textvariable=self.schedule_interval_var, font=font, width=4).pack(side=LEFT)
        tk.Label(schedule_frame, text="시간 (다중 키워드 시)", font=(FONT_UI, 8), bg=COLORS["content_bg"], fg=COLORS["text_muted"]).pack(side=LEFT, padx=3)

        self._add_run_buttons(content, self._run_write)

    def _build_bulk_panel(self, font):
        panel, content = self._make_panel("bulk", "Bulk Post", "엑셀 파일로 여러 계정에 대량 발행합니다")

        self.bulk_file_var = ttk.StringVar()
        self.bulk_mode_var = ttk.StringVar(value="schedule")

        card = self._make_card(content)
        card.columnconfigure(1, weight=1)

        self._make_field(card, "엑셀 파일", self.bulk_file_var, 0, width=30, font=font)
        ttk.Button(
            card, text="찾아보기", bootstyle="outline",
            command=lambda: self._browse_file(self.bulk_file_var, [("Excel", "*.xlsx")])
        ).grid(row=0, column=2, padx=(8, 0), pady=(0, 8))

        # 발행 방식
        mode_frame = tk.Frame(content, bg=COLORS["content_bg"])
        mode_frame.pack(fill=X, pady=(0, 10))
        tk.Label(
            mode_frame, text="발행 방식", font=(FONT_UI, 9, "bold"),
            bg=COLORS["content_bg"], fg=COLORS["text_secondary"],
        ).pack(side=LEFT, padx=(0, 12))
        ttk.Radiobutton(mode_frame, text="예약 발행", variable=self.bulk_mode_var, value="schedule").pack(side=LEFT, padx=5)
        ttk.Radiobutton(mode_frame, text="즉시 발행", variable=self.bulk_mode_var, value="publish").pack(side=LEFT, padx=5)

        self._add_run_buttons(content, self._run_bulk)

    def _build_like_panel(self, font):
        panel, content = self._make_panel("like", "Like", "이웃 새글 또는 키워드 검색으로 자동 공감합니다")

        self.like_mode_var = ttk.StringVar(value="neighbor")
        self.like_keyword_var = ttk.StringVar()
        self.like_count_var = ttk.IntVar(value=20)

        card = self._make_card(content)
        card.columnconfigure(1, weight=1)

        # 대상 선택
        tk.Label(
            card, text="공감 대상", font=(FONT_UI, 9, "bold"),
            bg=COLORS["card_bg"], fg=COLORS["text_secondary"], anchor="w",
        ).grid(row=0, column=0, sticky=W, pady=(0, 8), padx=(0, 15))
        mode_frame = tk.Frame(card, bg=COLORS["card_bg"])
        mode_frame.grid(row=0, column=1, sticky=W, pady=(0, 8))
        ttk.Radiobutton(mode_frame, text="이웃 새글", variable=self.like_mode_var, value="neighbor").pack(side=LEFT, padx=(0, 15))
        ttk.Radiobutton(mode_frame, text="키워드 검색", variable=self.like_mode_var, value="search").pack(side=LEFT)

        self._make_field(card, "검색 키워드", self.like_keyword_var, 1, width=30, font=font)
        self._make_hint(card, "키워드 검색 모드일 때 사용", 1, col=2)

        tk.Label(
            card, text="공감 수", font=(FONT_UI, 9, "bold"),
            bg=COLORS["card_bg"], fg=COLORS["text_secondary"], anchor="w",
        ).grid(row=2, column=0, sticky=W, pady=(0, 8), padx=(0, 15))
        ttk.Spinbox(card, from_=1, to=100, textvariable=self.like_count_var, font=font, width=10).grid(row=2, column=1, sticky=W, pady=(0, 8))

        self._add_run_buttons(content, self._run_like)

    def _build_comment_panel(self, font):
        panel, content = self._make_panel("comment", "Comment", "이웃 새글 또는 키워드 검색으로 AI 댓글을 작성합니다")

        self.comment_mode_var = ttk.StringVar(value="neighbor")
        self.comment_keyword_var = ttk.StringVar()
        self.comment_count_var = ttk.IntVar(value=10)
        self.comment_tone_var = ttk.StringVar(value="friendly")
        self.comment_custom_var = ttk.StringVar()

        card = self._make_card(content)
        card.columnconfigure(1, weight=1)

        # 대상 선택
        tk.Label(
            card, text="댓글 대상", font=(FONT_UI, 9, "bold"),
            bg=COLORS["card_bg"], fg=COLORS["text_secondary"], anchor="w",
        ).grid(row=0, column=0, sticky=W, pady=(0, 8), padx=(0, 15))
        mode_frame = tk.Frame(card, bg=COLORS["card_bg"])
        mode_frame.grid(row=0, column=1, sticky=W, pady=(0, 8))
        ttk.Radiobutton(mode_frame, text="이웃 새글", variable=self.comment_mode_var, value="neighbor").pack(side=LEFT, padx=(0, 15))
        ttk.Radiobutton(mode_frame, text="키워드 검색", variable=self.comment_mode_var, value="search").pack(side=LEFT)

        self._make_field(card, "검색 키워드", self.comment_keyword_var, 1, width=30, font=font)
        self._make_hint(card, "키워드 검색 모드일 때 사용", 1, col=2)

        tk.Label(
            card, text="댓글 수", font=(FONT_UI, 9, "bold"),
            bg=COLORS["card_bg"], fg=COLORS["text_secondary"], anchor="w",
        ).grid(row=2, column=0, sticky=W, pady=(0, 8), padx=(0, 15))
        ttk.Spinbox(card, from_=1, to=50, textvariable=self.comment_count_var, font=font, width=10).grid(row=2, column=1, sticky=W, pady=(0, 8))

        # 톤
        tk.Label(
            card, text="댓글 톤", font=(FONT_UI, 9, "bold"),
            bg=COLORS["card_bg"], fg=COLORS["text_secondary"], anchor="w",
        ).grid(row=3, column=0, sticky=W, pady=(0, 8), padx=(0, 15))
        tone_combo = ttk.Combobox(
            card, textvariable=self.comment_tone_var, font=font, width=30, state="readonly",
            values=["friendly - 친근한 톤", "professional - 전문적/정중한 톤", "casual - 가벼운 말투", "enthusiastic - 열정적/감탄"]
        )
        tone_combo.grid(row=3, column=1, sticky=W, pady=(0, 8))
        tone_combo.current(0)

        self._make_field(card, "커스텀 지시", self.comment_custom_var, 4, font=font)
        self._make_hint(card, "직접 입력하면 위 톤 설정 대신 사용", 4, col=2)

        self._add_run_buttons(content, self._run_comment)

    def _build_engage_panel(self, font):
        panel, content = self._make_panel(
            "engage",
            "고객과 친해질게요",
            "예리가 이웃 글에 공감하고 댓글을 남기며 친밀도를 쌓아요",
        )

        self.engage_mode_var = ttk.StringVar(value="neighbor")
        self.engage_keyword_var = ttk.StringVar()
        self.engage_like_var = ttk.IntVar(value=20)
        self.engage_comment_var = ttk.IntVar(value=10)
        self.engage_tone_var = ttk.StringVar(value="friendly")
        # AI 변형 옵션 — 매 댓글마다 AI가 자연스럽게 변형 (밴 회피)
        self.engage_ai_variation_var = ttk.BooleanVar(value=True)

        card = self._make_card(content)
        card.columnconfigure(1, weight=1)

        tk.Label(
            card, text="공감 수", font=(FONT_UI, 9, "bold"),
            bg=COLORS["card_bg"], fg=COLORS["text_secondary"], anchor="w",
        ).grid(row=0, column=0, sticky=W, pady=(0, 8), padx=(0, 15))
        ttk.Spinbox(card, from_=1, to=100, textvariable=self.engage_like_var, font=font, width=10).grid(row=0, column=1, sticky=W, pady=(0, 8))

        tk.Label(
            card, text="댓글 수", font=(FONT_UI, 9, "bold"),
            bg=COLORS["card_bg"], fg=COLORS["text_secondary"], anchor="w",
        ).grid(row=1, column=0, sticky=W, pady=(0, 8), padx=(0, 15))
        ttk.Spinbox(card, from_=1, to=50, textvariable=self.engage_comment_var, font=font, width=10).grid(row=1, column=1, sticky=W, pady=(0, 8))

        # 톤
        tk.Label(
            card, text="댓글 톤", font=(FONT_UI, 9, "bold"),
            bg=COLORS["card_bg"], fg=COLORS["text_secondary"], anchor="w",
        ).grid(row=2, column=0, sticky=W, pady=(0, 8), padx=(0, 15))
        tone_combo = ttk.Combobox(
            card, textvariable=self.engage_tone_var, font=font, width=30, state="readonly",
            values=["friendly - 친근한 톤", "professional - 전문적/정중한 톤", "casual - 가벼운 말투", "enthusiastic - 열정적/감탄"]
        )
        tone_combo.grid(row=2, column=1, sticky=W, pady=(0, 8))
        tone_combo.current(0)

        # AI 변형 옵션 토글 — 댓글 다양화로 밴 회피
        tk.Label(
            card, text="AI 변형", font=(FONT_UI, 9, "bold"),
            bg=COLORS["card_bg"], fg=COLORS["text_secondary"], anchor="w",
        ).grid(row=3, column=0, sticky=W, pady=(0, 8), padx=(0, 15))
        ai_frame = tk.Frame(card, bg=COLORS["card_bg"])
        ai_frame.grid(row=3, column=1, columnspan=2, sticky=W, pady=(0, 8))
        ttk.Checkbutton(
            ai_frame, text="매 댓글마다 AI가 자연스럽게 변형 (권장)",
            variable=self.engage_ai_variation_var,
            bootstyle="round-toggle",
        ).pack(side=LEFT)
        self._make_hint(
            card,
            "ON일 때 같은 톤이라도 매번 다른 문장 구조로 작성 — 봇 탐지 회피",
            4, col=1,
        )

        # ── 차단 위험 고지 배너 (엔게이지 전용) ──
        warn_wrap = tk.Frame(
            content,
            bg="#FFF4CE",
            highlightbackground="#E6B800",
            highlightthickness=1,
        )
        warn_wrap.pack(fill=X, padx=0, pady=(8, 0))
        warn_inner = tk.Frame(warn_wrap, bg="#FFF4CE")
        warn_inner.pack(fill=X, padx=15, pady=12)
        tk.Label(
            warn_inner,
            text="⚠ 공감·댓글 자동화도 탐지 대상입니다",
            font=(FONT_UI, 10, "bold"),
            bg="#FFF4CE", fg="#8B4500", anchor="w",
        ).pack(fill=X)
        tk.Label(
            warn_inner,
            text=(
                "• 권장: 하루 공감 20~30회, 댓글 30회 이하\n"
                "• 같은 댓글 문구 반복은 봇 탐지 신호입니다 — AI 변형 ON 권장\n"
                "• 계정 정지·저품질 리스크 책임은 사용자에게 있습니다"
            ),
            font=(FONT_UI, 9),
            bg="#FFF4CE", fg="#5C3100", anchor="w", justify="left",
            wraplength=700,
        ).pack(fill=X, pady=(6, 0))

        self._add_run_buttons(content, self._run_engage)

    def _build_find_keyword_panel(self, font):
        panel, content = self._make_panel(
            "find_keyword",
            "고객을 찾아볼게요",
            "현주씨가 키워드로 잠재 고객을 찾아 서로이웃까지 신청합니다",
            tabs=[("find_keyword", "🔍  키워드 검색"), ("find_link", "🔗  특정 블로거 링크")],
        )

        self.neighbor_keywords_var = ttk.StringVar()
        self.neighbor_count_var = ttk.IntVar(value=10)
        # 속도 모드: safe / normal / fast
        self.neighbor_speed_var = ttk.StringVar(value="safe")
        # Cool-down: N명마다 60~180초 긴 휴식. 0이면 비활성화.
        self.neighbor_cooldown_var = ttk.IntVar(value=10)
        # 일일 상한 — 출시 기본값은 보수적으로 50명.
        self.neighbor_daily_limit_var = ttk.IntVar(value=aimax.RECOMMENDED_LIMITS["daily_neighbor_requests"])

        # ── 위험 고지 배너 ──
        warn_wrap = tk.Frame(
            content,
            bg="#FFF4CE",  # 연한 노란색
            highlightbackground="#E6B800",
            highlightthickness=1,
        )
        warn_wrap.pack(fill=X, padx=0, pady=(0, 12))
        warn_inner = tk.Frame(warn_wrap, bg="#FFF4CE")
        warn_inner.pack(fill=X, padx=15, pady=12)
        tk.Label(
            warn_inner,
            text="⚠ 현주씨가 일하다 막힐 수 있어요 — 꼭 읽어 주세요",
            font=(FONT_UI, 10, "bold"),
            bg="#FFF4CE", fg="#8B4500", anchor="w",
        ).pack(fill=X)
        warn_body = (
            "• 네이버는 2025년 7월부터 서로이웃·공감·댓글 자동화 탐지를 강화했습니다. "
            "현주씨가 일하는 모습도 탐지 대상이며, 계정 제한·저품질 블로그 지정 위험이 있어요.\n"
            "• AIMAX는 안전을 위해 일일 서로이웃 신청 50명을 기본 권장값으로 사용합니다.\n"
            "• 첫 1주일은 '안전 모드 + 일일 30명' 이하로 시켜보고 계정 반응을 살펴 주세요.\n"
            "• 같은 멘트 반복은 위험합니다. 인사 멘트는 최소 3개 이상 + {닉네임} 변수 활용을 권장합니다.\n"
            "• 현주씨가 일하다 발생한 계정 제재/블로그 손상 책임은 사용자 본인에게 있습니다."
        )
        tk.Label(
            warn_inner, text=warn_body,
            font=(FONT_UI, 9),
            bg="#FFF4CE", fg="#5C3100", anchor="w", justify="left",
            wraplength=700,
        ).pack(fill=X, pady=(6, 0))

        card = self._make_card(content)
        card.columnconfigure(1, weight=1)

        self._make_field(card, "검색 키워드", self.neighbor_keywords_var, 0, font=font)
        self._make_hint(card, "쉼표로 구분: 맛집,여행,카페", 0, col=2)

        tk.Label(
            card, text="키워드당 신청 수", font=(FONT_UI, 9, "bold"),
            bg=COLORS["card_bg"], fg=COLORS["text_secondary"], anchor="w",
        ).grid(row=1, column=0, sticky=W, pady=(0, 8), padx=(0, 15))
        ttk.Spinbox(card, from_=1, to=30, textvariable=self.neighbor_count_var, font=font, width=10).grid(row=1, column=1, sticky=W, pady=(0, 8))

        # ── 멘트 다중화: Text 위젯 (줄바꿈으로 여러 멘트) ──
        tk.Label(
            card, text="신청 멘트", font=(FONT_UI, 9, "bold"),
            bg=COLORS["card_bg"], fg=COLORS["text_secondary"], anchor="nw",
        ).grid(row=2, column=0, sticky="nw", pady=(0, 8), padx=(0, 15))
        msg_wrap = tk.Frame(card, bg=COLORS["card_bg"])
        msg_wrap.grid(row=2, column=1, columnspan=2, sticky=EW, pady=(0, 8))
        msg_wrap.columnconfigure(0, weight=1)

        # AI 생성 버튼 바 (멘트 Text 위)
        ai_bar = tk.Frame(msg_wrap, bg=COLORS["card_bg"])
        ai_bar.grid(row=0, column=0, sticky=EW, pady=(0, 6))
        self.ai_ment_btn = ttk.Button(
            ai_bar, text="✨ 진정성 있게 10개 생성", bootstyle="info-outline",
            command=self._generate_neighbor_messages_ai, width=24,
        )
        self.ai_ment_btn.pack(side=LEFT)
        self.ai_ment_status = tk.Label(
            ai_bar, text="", font=(FONT_UI, 8),
            bg=COLORS["card_bg"], fg=COLORS["text_muted"],
        )
        self.ai_ment_status.pack(side=LEFT, padx=(10, 0))

        self.neighbor_msg_text = tk.Text(
            msg_wrap, height=9, font=font, wrap="word",
            bg="white", fg=COLORS["text_primary"],
            highlightbackground=COLORS["card_border"], highlightthickness=1,
            relief="flat", padx=10, pady=8,
        )
        self.neighbor_msg_text.grid(row=1, column=0, sticky=EW)
        # 기본 멘트 풀 — 닉네임 치환 변수 포함
        default_msgs = (
            "{닉네임}님 안녕하세요! 좋은 글 잘 보고 갑니다 :) 서로이웃 신청드려요~\n"
            "{닉네임}님 글 잘 읽었습니다^^ 자주 소통하고 싶어서 서로이웃 신청드려요!\n"
            "{닉네임}님 반갑습니다! 블로그 잘 보고 갑니다. 서로이웃해요~\n"
            "안녕하세요 {닉네임}님! 좋은 포스팅 감사합니다. 서로이웃 신청드릴게요!\n"
            "{닉네임}님 글이 참 좋네요 :) 서로이웃으로 자주 찾아뵐게요!"
        )
        saved_msgs = load_neighbor_messages(self.naver_id_var.get())
        if saved_msgs:
            self.neighbor_msg_text.insert("1.0", "\n".join(saved_msgs))
            self.ai_ment_status.configure(text=f"저장된 멘트 {len(saved_msgs)}개 불러옴", fg=COLORS["text_muted"])
        else:
            self.neighbor_msg_text.insert("1.0", default_msgs)
        self._make_hint(
            msg_wrap,
            "한 줄에 하나씩. 네이버 제한은 400자이며, AI는 90~180자 안팎으로 생성합니다.",
            2, col=0,
        )

        def _reload_neighbor_messages_for_account(*_):
            if not hasattr(self, "neighbor_msg_text"):
                return
            account_msgs = load_neighbor_messages(self.naver_id_var.get())
            if not account_msgs:
                return
            self.neighbor_msg_text.delete("1.0", "end")
            self.neighbor_msg_text.insert("1.0", "\n".join(account_msgs))
            self.ai_ment_status.configure(
                text=f"저장된 멘트 {len(account_msgs)}개 불러옴", fg=COLORS["text_muted"],
            )

        self.naver_id_var.trace_add("write", _reload_neighbor_messages_for_account)

        # ── 속도 모드 라디오 ──
        tk.Label(
            card, text="속도 모드", font=(FONT_UI, 9, "bold"),
            bg=COLORS["card_bg"], fg=COLORS["text_secondary"], anchor="w",
        ).grid(row=3, column=0, sticky=W, pady=(0, 8), padx=(0, 15))
        speed_wrap = tk.Frame(card, bg=COLORS["card_bg"])
        speed_wrap.grid(row=3, column=1, columnspan=2, sticky=W, pady=(0, 8))
        for idx, (val, label) in enumerate([
            ("safe", "안전 (30~90초/건)"),
            ("normal", "보통 (4~15초/건)"),
            ("fast", "빠름 (3~8초/건)"),
        ]):
            ttk.Radiobutton(
                speed_wrap, text=label, value=val,
                variable=self.neighbor_speed_var,
            ).grid(row=0, column=idx, sticky=W, padx=(0, 12))

        # ── Cool-down ──
        tk.Label(
            card, text="Cool-down (명)", font=(FONT_UI, 9, "bold"),
            bg=COLORS["card_bg"], fg=COLORS["text_secondary"], anchor="w",
        ).grid(row=4, column=0, sticky=W, pady=(0, 8), padx=(0, 15))
        ttk.Spinbox(
            card, from_=0, to=50, textvariable=self.neighbor_cooldown_var,
            font=font, width=10,
        ).grid(row=4, column=1, sticky=W, pady=(0, 8))
        self._make_hint(card, "N명 신청마다 60~180초 휴식 (0=끔)", 4, col=2)

        # ── 일일 상한 ──
        tk.Label(
            card, text="일일 상한 (명)", font=(FONT_UI, 9, "bold"),
            bg=COLORS["card_bg"], fg=COLORS["text_secondary"], anchor="w",
        ).grid(row=5, column=0, sticky=W, pady=(0, 8), padx=(0, 15))
        ttk.Spinbox(
            card, from_=0, to=100, textvariable=self.neighbor_daily_limit_var,
            font=font, width=10,
        ).grid(row=5, column=1, sticky=W, pady=(0, 8))
        self._make_hint(card, "권장: 50 이하 (0=무제한, 초과 시 확인 필요)", 5, col=2)

        self._add_run_buttons(content, self._run_neighbor)

    def _build_find_link_panel(self, font):
        """특정 블로거 링크 기반 서이추."""
        panel, content = self._make_panel(
            "find_link",
            "고객을 찾아볼게요",
            "특정 블로거의 이웃 목록에서 잠재 고객을 찾아옵니다",
            tabs=[("find_keyword", "🔍  키워드 검색"), ("find_link", "🔗  특정 블로거 링크")],
        )

        self.link_blogger_url_var = ttk.StringVar()
        self.link_neighbor_count_var = ttk.IntVar(value=10)

        warn_wrap = tk.Frame(
            content,
            bg="#FFF4CE",
            highlightbackground="#E6B800",
            highlightthickness=1,
        )
        warn_wrap.pack(fill=X, padx=0, pady=(0, 12))
        warn_inner = tk.Frame(warn_wrap, bg="#FFF4CE")
        warn_inner.pack(fill=X, padx=15, pady=12)
        tk.Label(
            warn_inner,
            text="⚠ 특정 블로거의 공개 이웃 목록만 찾아옵니다",
            font=(FONT_UI, 10, "bold"),
            bg="#FFF4CE", fg="#8B4500", anchor="w",
        ).pack(fill=X)
        tk.Label(
            warn_inner,
            text=(
                "• 이웃 목록이 비공개이거나 네이버가 접근을 제한하면 수집하지 못할 수 있어요.\n"
                "• 수집 후 서로이웃 신청은 같은 일일 상한·속도 모드·Cool-down 규칙을 따릅니다."
            ),
            font=(FONT_UI, 9),
            bg="#FFF4CE", fg="#5C3100", anchor="w", justify="left",
            wraplength=700,
        ).pack(fill=X, pady=(6, 0))

        card = self._make_card(content)
        card.columnconfigure(1, weight=1)

        self._make_field(card, "블로거 링크/ID", self.link_blogger_url_var, 0, font=font)
        self._make_hint(card, "예: https://blog.naver.com/example 또는 example", 0, col=2)

        tk.Label(
            card, text="신청 수", font=(FONT_UI, 9, "bold"),
            bg=COLORS["card_bg"], fg=COLORS["text_secondary"], anchor="w",
        ).grid(row=1, column=0, sticky=W, pady=(0, 8), padx=(0, 15))
        ttk.Spinbox(
            card, from_=1, to=30, textvariable=self.link_neighbor_count_var,
            font=font, width=10,
        ).grid(row=1, column=1, sticky=W, pady=(0, 8))
        self._make_hint(card, "후보는 더 많이 모으고, 성공 신청 수는 이 값까지만 진행", 1, col=2)

        tk.Label(
            card, text="신청 멘트", font=(FONT_UI, 9, "bold"),
            bg=COLORS["card_bg"], fg=COLORS["text_secondary"], anchor="nw",
        ).grid(row=2, column=0, sticky="nw", pady=(0, 8), padx=(0, 15))
        msg_wrap = tk.Frame(card, bg=COLORS["card_bg"])
        msg_wrap.grid(row=2, column=1, columnspan=2, sticky=EW, pady=(0, 8))
        msg_wrap.columnconfigure(0, weight=1)

        self.link_neighbor_msg_text = tk.Text(
            msg_wrap, height=8, font=font, wrap="word",
            bg="white", fg=COLORS["text_primary"],
            highlightbackground=COLORS["card_border"], highlightthickness=1,
            relief="flat", padx=10, pady=8,
        )
        self.link_neighbor_msg_text.grid(row=0, column=0, sticky=EW)
        default_msgs = (
            "{닉네임}님 안녕하세요! 좋은 글 잘 보고 갑니다 :) 서로이웃 신청드려요~\n"
            "{닉네임}님 글 잘 읽었습니다^^ 자주 소통하고 싶어서 서로이웃 신청드려요!\n"
            "{닉네임}님 반갑습니다! 블로그 잘 보고 갑니다. 서로이웃해요~\n"
            "안녕하세요 {닉네임}님! 좋은 포스팅 감사합니다. 서로이웃 신청드릴게요!\n"
            "{닉네임}님 글이 참 좋네요 :) 서로이웃으로 자주 찾아뵐게요!"
        )
        saved_msgs = load_neighbor_messages(self.naver_id_var.get())
        self.link_neighbor_msg_text.insert("1.0", "\n".join(saved_msgs) if saved_msgs else default_msgs)
        self.link_neighbor_status = tk.Label(
            msg_wrap,
            text=f"저장된 멘트 {len(saved_msgs)}개 불러옴" if saved_msgs else "",
            font=(FONT_UI, 8),
            bg=COLORS["card_bg"], fg=COLORS["text_muted"],
        )
        self.link_neighbor_status.grid(row=1, column=0, sticky=W, pady=(5, 0))

        def _reload_link_neighbor_messages_for_account(*_):
            if not hasattr(self, "link_neighbor_msg_text"):
                return
            account_msgs = load_neighbor_messages(self.naver_id_var.get())
            if not account_msgs:
                return
            self.link_neighbor_msg_text.delete("1.0", "end")
            self.link_neighbor_msg_text.insert("1.0", "\n".join(account_msgs))
            self.link_neighbor_status.configure(
                text=f"저장된 멘트 {len(account_msgs)}개 불러옴", fg=COLORS["text_muted"],
            )

        self.naver_id_var.trace_add("write", _reload_link_neighbor_messages_for_account)

        tk.Label(
            card, text="속도 모드", font=(FONT_UI, 9, "bold"),
            bg=COLORS["card_bg"], fg=COLORS["text_secondary"], anchor="w",
        ).grid(row=3, column=0, sticky=W, pady=(0, 8), padx=(0, 15))
        speed_wrap = tk.Frame(card, bg=COLORS["card_bg"])
        speed_wrap.grid(row=3, column=1, columnspan=2, sticky=W, pady=(0, 8))
        for idx, (val, label) in enumerate([
            ("safe", "안전 (30~90초/건)"),
            ("normal", "보통 (4~15초/건)"),
            ("fast", "빠름 (3~8초/건)"),
        ]):
            ttk.Radiobutton(
                speed_wrap, text=label, value=val,
                variable=self.neighbor_speed_var,
            ).grid(row=0, column=idx, sticky=W, padx=(0, 12))

        tk.Label(
            card, text="Cool-down (명)", font=(FONT_UI, 9, "bold"),
            bg=COLORS["card_bg"], fg=COLORS["text_secondary"], anchor="w",
        ).grid(row=4, column=0, sticky=W, pady=(0, 8), padx=(0, 15))
        ttk.Spinbox(
            card, from_=0, to=50, textvariable=self.neighbor_cooldown_var,
            font=font, width=10,
        ).grid(row=4, column=1, sticky=W, pady=(0, 8))
        self._make_hint(card, "N명 신청마다 60~180초 휴식 (0=끔)", 4, col=2)

        tk.Label(
            card, text="일일 상한 (명)", font=(FONT_UI, 9, "bold"),
            bg=COLORS["card_bg"], fg=COLORS["text_secondary"], anchor="w",
        ).grid(row=5, column=0, sticky=W, pady=(0, 8), padx=(0, 15))
        ttk.Spinbox(
            card, from_=0, to=100, textvariable=self.neighbor_daily_limit_var,
            font=font, width=10,
        ).grid(row=5, column=1, sticky=W, pady=(0, 8))
        self._make_hint(card, "권장: 50 이하 (0=무제한, 초과 시 확인 필요)", 5, col=2)

        self._add_run_buttons(content, self._run_link_neighbor)

    def _build_scraper_panel(self, font):
        panel, content = self._make_panel(
            "scraper",
            "1. 고객을 찾아볼게요",
            "현주씨가 키워드로 잠재 고객 블로거를 찾아 모아옵니다",
        )

        self.scraper_keyword_var = ttk.StringVar()
        self.scraper_count_var = ttk.IntVar(value=50)

        input_card = self._make_card(content)
        input_card.columnconfigure(1, weight=1)
        self._make_field(input_card, "검색 키워드", self.scraper_keyword_var, 0, font=font)
        self._make_hint(input_card, "수집할 키워드 입력", 0, col=2)

        tk.Label(input_card, text="최대 수집 수", font=(FONT_UI, 9, "bold"),
                 bg=COLORS["card_bg"], fg=COLORS["text_secondary"], anchor="w").grid(
            row=1, column=0, sticky=W, pady=(0, 8), padx=(0, 15))
        ttk.Spinbox(input_card, from_=10, to=500, textvariable=self.scraper_count_var,
                    font=font, width=10).grid(row=1, column=1, sticky=W, pady=(0, 8))

        self._add_run_buttons(content, self._run_scraper)

        action_row = tk.Frame(content, bg=COLORS["content_bg"])
        action_row.pack(fill=X, pady=(0, 6))
        ttk.Button(
            action_row, text="최근 CSV 열기", bootstyle="secondary-outline",
            command=self._open_last_scraper_csv, width=18
        ).pack(side=LEFT, padx=(0, 8))
        ttk.Button(
            action_row, text="내보내기 폴더 열기", bootstyle="secondary-outline",
            command=self._open_exports_folder, width=18
        ).pack(side=LEFT)

    # ── 하단 로그 패널 ──
    def _build_log_panel(self, parent):
        log_container = tk.Frame(parent, bg=COLORS["terminal_bg"])
        log_container.pack(fill=X, side=BOTTOM)

        # 프로그레스 바
        self.progress_var = ttk.DoubleVar()
        ttk.Progressbar(log_container, variable=self.progress_var, maximum=100, bootstyle="info-striped").pack(fill=X)

        # 로그 헤더
        log_header = tk.Frame(log_container, bg=COLORS["terminal_bg"])
        log_header.pack(fill=X, padx=10, pady=(6, 0))
        tk.Label(
            log_header, text="Console", font=(FONT_MONO, 9, "bold"),
            bg=COLORS["terminal_bg"], fg=COLORS["text_muted"], anchor="w",
        ).pack(side=LEFT)
        tk.Label(
            log_header, text="Clear", font=(FONT_MONO, 8), cursor="hand2",
            bg=COLORS["terminal_bg"], fg=COLORS["text_muted"], anchor="e",
        ).pack(side=RIGHT)
        # Clear 클릭 이벤트
        log_header.winfo_children()[-1].bind("<Button-1>", lambda e: self._clear_log())

        # 로그 텍스트
        text_frame = tk.Frame(log_container, bg=COLORS["terminal_bg"])
        text_frame.pack(fill=X, padx=10, pady=(4, 8))

        scrollbar = tk.Scrollbar(text_frame)
        scrollbar.pack(side=RIGHT, fill=Y)

        self.log_text = tk.Text(
            text_frame, height=8, wrap=WORD, font=(FONT_MONO, 9),
            yscrollcommand=scrollbar.set, state=DISABLED,
            bg=COLORS["terminal_bg"], fg=COLORS["terminal_fg"],
            insertbackground=COLORS["terminal_fg"],
            borderwidth=0, highlightthickness=0,
            selectbackground="#264f78", selectforeground="white",
        )
        self.log_text.pack(side=LEFT, fill=BOTH, expand=YES)
        scrollbar.config(command=self.log_text.yview)

    # ── 공통: 실행/중지 버튼 ──
    def _add_run_buttons(self, parent, command):
        btn_frame = tk.Frame(parent, bg=COLORS["content_bg"])
        btn_frame.pack(fill=X, pady=(15, 5))

        run_btn = ttk.Button(
            btn_frame, text="  실행  ", bootstyle="success", command=command, width=14,
        )
        run_btn.pack(side=LEFT, padx=(0, 10))

        stop_btn = ttk.Button(
            btn_frame, text="  중지  ", bootstyle="danger-outline", command=self._stop_worker, width=14, state=DISABLED,
        )
        stop_btn.pack(side=LEFT)

    # ── 파일 찾아보기 ──
    def _browse_file(self, var, filetypes):
        path = filedialog.askopenfilename(filetypes=filetypes)
        if path:
            var.set(path)

    # ── 설정 저장 ──
    def _save_credentials(self):
        save_settings(
            self.naver_id_var.get(), self.naver_pw_var.get(), self.api_key_var.get(),
            self.ai_model_var.get(), self.claude_key_var.get(), self.openai_key_var.get(), self.apify_key_var.get()
        )
        # 블로그 프로필도 키체인에 저장 (네이버 계정별)
        profile = self._get_blog_profile_text()
        save_blog_profile(self.naver_id_var.get(), profile)
        self._log("설정이 저장되었습니다.")
        if profile:
            self._log(f"내 블로그 프로필 저장됨 (네이버 ID: {self.naver_id_var.get()})")
        self._send_immediate_web_agent_heartbeat("settings_saved")

    def _get_blog_profile_text(self):
        """프로필 Text 위젯에서 실제 내용 반환 (placeholder는 빈 문자열로)."""
        if not hasattr(self, "blog_profile_text"):
            return ""
        raw = self.blog_profile_text.get("1.0", "end-1c").strip()
        # placeholder 상태 체크 (fg가 muted면 placeholder 표시 중)
        try:
            current_fg = str(self.blog_profile_text.cget("fg"))
            if current_fg == COLORS["text_muted"]:
                return ""
        except Exception:
            pass
        return raw

    # ── 로그 ──
    def _log(self, msg):
        self.queue.put(("log", msg))

    def _write_crew_result(self, employee_id, status="done", inputs=None, outputs=None, summary=""):
        """Al Crew Creator가 읽을 수 있는 직원 실행 결과 JSON을 저장한다."""
        try:
            from datetime import datetime

            employee_id = (employee_id or "unknown").strip() or "unknown"
            now = datetime.now().astimezone()
            job_id = f"{now.strftime('%Y%m%d_%H%M%S')}_{employee_id}"
            payload = {
                "schemaVersion": 1,
                "jobId": job_id,
                "employeeId": employee_id,
                "sourceApp": aimax.APP_NAME,
                "appMode": self.app_mode,
                "status": status,
                "createdAt": now.isoformat(timespec="seconds"),
                "inputs": inputs or {},
                "outputs": outputs or {},
                "summary": summary or "",
            }

            os.makedirs(CREW_RESULTS_DIR, exist_ok=True)
            result_path = os.path.join(CREW_RESULTS_DIR, f"{job_id}.json")
            tmp_path = f"{result_path}.tmp"
            with open(tmp_path, "w", encoding="utf-8") as f:
                json.dump(payload, f, ensure_ascii=False, indent=2)
            os.replace(tmp_path, result_path)
            self._log(f"[Crew] 결과 저장: {result_path}")
            return result_path
        except Exception as e:
            self._log(f"[Crew] 결과 저장 실패: {e}")
            return None

    def _clear_log(self):
        self.log_text.config(state=NORMAL)
        self.log_text.delete("1.0", END)
        self.log_text.config(state=DISABLED)

    # ── Queue 폴링 ──
    def _poll_queue(self):
        from queue import Empty
        try:
            while True:
                msg_type, msg_data = self.queue.get_nowait()
                if msg_type == "log":
                    self.log_text.config(state=NORMAL)
                    self.log_text.insert(END, msg_data + "\n")
                    self.log_text.see(END)
                    self.log_text.config(state=DISABLED)
                elif msg_type == "progress":
                    self.progress_var.set(msg_data)
                elif msg_type == "done":
                    self._on_worker_done()
                elif msg_type == "popup":
                    # msg_data: (stage_key, title, body, next_stage)
                    self._show_stage_completion_popup(*msg_data)
                elif msg_type == "ai_ment_done":
                    self._apply_generated_messages(msg_data)
                elif msg_type == "ai_ment_error":
                    self.ai_ment_btn.configure(state="normal")
                    self.ai_ment_status.configure(
                        text=f"⚠ 오류: {msg_data[:40]}", fg="#C0392B"
                    )
                    self._log(f"[AI 멘트] 생성 오류: {msg_data}")
                elif msg_type == "web_agent_status":
                    self._set_web_agent_status(*msg_data)
                elif msg_type == "web_agent_update_popup":
                    self._show_update_popup(msg_data)
                elif msg_type == "web_agent_controls":
                    if hasattr(self, "web_connect_btn"):
                        self.web_connect_btn.configure(state=NORMAL)
                elif msg_type == "web_agent_clear_password":
                    self.web_password_var.set("")
                elif msg_type == "web_agent_command":
                    self._handle_web_agent_command(msg_data)
                elif msg_type == "remote_job":
                    self._start_remote_job(msg_data)
        except Empty:
            pass
        except Exception as e:
            self.log_text.config(state=NORMAL)
            self.log_text.insert(END, f"[내부 오류] {e}\n")
            self.log_text.see(END)
            self.log_text.config(state=DISABLED)
        self.root.after(100, self._poll_queue)

    # ── 워커 시작/중지 ──
    def _start_worker(self, target, **kwargs):
        if self.running:
            self._log("이미 실행 중입니다.")
            return

        self.stop_event.clear()
        self.running = True
        self.progress_var.set(0)
        self._set_buttons_running(True)

        self.worker_thread = threading.Thread(target=target, kwargs=kwargs, daemon=True)
        self.worker_thread.start()

    def _stop_worker(self):
        if self.running:
            self.stop_event.set()
            self.running = False
            self._log("중지 요청됨. 진행 중인 작업을 정리하고 있습니다...")
            if self.driver:
                try:
                    self.driver.quit()
                except Exception:
                    pass
                self.driver = None

    def _on_worker_done(self):
        self.running = False
        self._set_buttons_running(False)
        self.driver = None

    def _set_buttons_running(self, is_running):
        for widget in self.root.winfo_children():
            self._toggle_buttons_recursive(widget, is_running)

    def _toggle_buttons_recursive(self, widget, is_running):
        if isinstance(widget, ttk.Button):
            text = widget.cget("text").strip()
            if "실행" in text:
                widget.config(state=DISABLED if is_running else NORMAL)
            elif "중지" in text:
                widget.config(state=NORMAL if is_running else DISABLED)
        for child in widget.winfo_children():
            self._toggle_buttons_recursive(child, is_running)

    # ── 인증정보 검증 ──
    def _validate_credentials(self, need_api=True):
        if not self.naver_id_var.get():
            self._log("[오류] 네이버 ID를 입력해주세요.")
            return False
        if not self.naver_pw_var.get():
            self._log("[오류] 비밀번호를 입력해주세요.")
            return False
        if need_api:
            model = self.ai_model_var.get()
            if model == "claude" and not self.claude_key_var.get():
                self._log("[오류] Claude API Key를 입력해주세요.")
                return False
            if _is_openai_model(model) and not self.openai_key_var.get():
                self._log("[오류] OpenAI API Key를 입력해주세요.")
                return False
            if model != "claude" and not _is_openai_model(model) and not self.api_key_var.get():
                self._log("[오류] Gemini API Key를 입력해주세요.")
                return False
        return True

    # ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
    # 워커 함수들 (백그라운드 스레드에서 실행)
    # ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━

    def _run_write(self):
        keyword_str = self.write_keyword_var.get().strip()
        md_file = self.write_file_var.get().strip()
        mode = self.write_mode_var.get()
        style_id = self.write_style_var.get().split(" - ")[0].strip()
        word_count = _normalize_target_char_count(self.write_wordcount_var.get())
        font_name = self.write_font_var.get().strip()
        if font_name == "기본값":
            font_name = None
        category = self.write_category_var.get().strip() or None
        cta_link = self.cta_link_var.get().strip() or None
        cta_text = self.cta_text_var.get().strip() or None
        schedule_date = self.schedule_date_var.get().strip() or None
        schedule_hour = self.schedule_hour_var.get().strip() or None
        schedule_interval = self.schedule_interval_var.get().strip() or None

        # 예약 일시가 입력되어 있으면 자동으로 예약 모드로 전환
        if schedule_date and mode != "schedule":
            mode = "schedule"
            self.write_mode_var.set("schedule")
            self._log("[자동 전환] 예약 일시가 입력되어 예약 발행 모드로 전환합니다.")

        # 쉼표로 키워드 분리
        keywords = [k.strip() for k in keyword_str.split(",") if k.strip()] if keyword_str else []

        if not keywords and not md_file:
            self._log("[오류] 키워드 또는 MD 파일을 입력해주세요.")
            return
        if not self._validate_credentials(need_api=True):
            return
        post_count = len(keywords) if keywords else 1
        if not self._confirm_limit_override(
            "일일 자동 포스팅",
            post_count,
            aimax.RECOMMENDED_LIMITS["daily_auto_posts"],
        ):
            self._log("권장 한도 초과 확인이 취소되어 실행하지 않았습니다.")
            return

        self._start_worker(self._worker_write, keywords=keywords, md_file=md_file, mode=mode, style_id=style_id, category=category, cta_link=cta_link, cta_text=cta_text, schedule_date=schedule_date, schedule_hour=schedule_hour, schedule_interval=schedule_interval, word_count=word_count, font_name=font_name)

    def _worker_write(self, keywords, md_file, mode, style_id="info", category=None, cta_link=None, cta_text=None, schedule_date=None, schedule_hour=None, schedule_interval=None, word_count=1500, image_count=3, font_name=None, ai_model=None, seo_brief=None, artifact=None):
        success = 0
        total = 0
        last_error = None
        stage = "init"
        failed_keyword = ""
        usage_totals = {}
        image_totals = {"attempted": 0, "generated": 0, "inserted": 0, "providers": {"gemini": 0, "openai": 0}, "failures": []}
        post_results = []
        failed_posts = []
        try:
            import time
            from datetime import datetime, timedelta
            from browser.stealth_driver import create_stealth_driver
            from auth.naver_login import login
            from content.ai_text import generate_blog_content, measure_visible_char_count
            from content.markdown_parser import parse_markdown, parse_markdown_file
            from posting.editor import navigate_to_editor, input_title, input_content, set_font
            from posting.publisher import save_draft, publish_now, schedule_publish
            from utils.delays import random_delay

            nid = self.naver_id_var.get()
            npw = self.naver_pw_var.get()
            ai_model = _normalize_ai_model(ai_model or self.ai_model_var.get())
            image_count = _normalize_image_count(image_count)
            gemini_key = self.api_key_var.get()
            claude_key = self.claude_key_var.get()
            openai_key = self.openai_key_var.get()
            artifact_mode = isinstance(artifact, dict) and bool(str(artifact.get("content_markdown") or "").strip())
            if ai_model == "claude":
                text_api_key = claude_key
            elif _is_openai_model(ai_model):
                text_api_key = openai_key
            else:
                text_api_key = gemini_key
            if artifact_mode and openai_key:
                image_provider = "openai"
            else:
                image_provider = "openai" if _is_openai_model(ai_model) and openai_key else "gemini"
            image_api_key = openai_key if image_provider == "openai" else gemini_key
            fallback_image_api_key = gemini_key if image_provider == "openai" else openai_key

            # 포스트 목록 구성
            if artifact_mode:
                posts = [{
                    "type": "artifact",
                    "source": str(artifact.get("title") or (keywords[0] if keywords else "서버 생성 글")).strip(),
                    "artifact": artifact,
                }]
            elif md_file and not keywords:
                posts = [{"type": "md_file", "source": md_file}]
            else:
                posts = [{"type": "keyword", "source": kw} for kw in keywords]

            total = len(posts)
            is_multi = total > 1

            if is_multi:
                self._log(f"다중 키워드 모드: {total}개 글 작성 예정")

            usage_totals = {}
            image_totals = {"attempted": 0, "generated": 0, "inserted": 0, "providers": {"gemini": 0, "openai": 0}, "failures": []}
            post_results = []
            failed_posts = []

            def _record_failed(post, post_stage, error, title="", char_count=0, images=None):
                failed = {
                    "source": post.get("source", ""),
                    "keyword": post.get("source", ""),
                    "type": post.get("type", ""),
                    "title": title or "",
                    "status": "failed",
                    "stage": post_stage or "unknown",
                    "error": str(error)[:500],
                    "char_count": char_count or 0,
                    "target_char_count": word_count,
                }
                if images:
                    failed["images"] = images
                failed_posts.append(failed)

            def _reset_driver(reason=None):
                nonlocal stage
                if reason:
                    self._log(f"[복구] {reason}")
                if self.driver:
                    try:
                        self.driver.quit()
                    except Exception:
                        pass
                    self.driver = None
                stage = "browser_start"
                self._log("스텔스 브라우저 시작...")
                self.driver = create_stealth_driver()
                stage = "naver_login"
                self._log("네이버 로그인 중...")
                login(self.driver, nid, npw)

            # 예약 기본 시간 파싱
            base_date = None
            base_hour = None
            interval_hours = int(schedule_interval) if schedule_interval else 1
            if schedule_date:
                try:
                    base_date = datetime.strptime(schedule_date, "%Y-%m-%d")
                except ValueError:
                    self._log(f"[경고] 날짜 형식 오류 ({schedule_date}), 기본값 사용")
            if schedule_hour:
                try:
                    base_hour = int(schedule_hour)
                except ValueError:
                    self._log(f"[경고] 시간 형식 오류 ({schedule_hour}), 기본값 사용")

            # 1. 브라우저 시작 + 로그인 (1회, 창/DevTools 조기 종료는 1회 재시도)
            self.queue.put(("progress", 5))
            for login_attempt in range(2):
                try:
                    _reset_driver("네이버 로그인 창 재시도" if login_attempt else None)
                    break
                except Exception as login_error:
                    if login_attempt == 0 and _is_browser_session_error(login_error):
                        self._log("[복구] 네이버 로그인 창이 닫혀 브라우저를 새로 열고 1회 재시도합니다.")
                        continue
                    raise
            self.queue.put(("progress", 15))

            if not self.running:
                return _build_write_result(
                    False, success, total, mode, ai_model, usage_totals, image_totals,
                    post_results, failed_posts, error="글쓰기 작업이 중지되었습니다.", stage=stage,
                    failed_keyword=failed_keyword,
                )

            # 2. 키워드별 순차 처리
            success = 0
            for i, post in enumerate(posts):
                if not self.running:
                    self._log("사용자에 의해 중지됨")
                    break

                pct = 15 + int((i / total) * 80)
                self.queue.put(("progress", pct))

                if is_multi:
                    self._log(f"── [{i+1}/{total}] 시작 ──")

                try:
                    post_stage = "content_generation"
                    title = ""
                    visible_char_count = 0
                    image_attempted = 0
                    image_generated = 0
                    image_inserted = 0
                    image_failures = []
                    post_usage = {}
                    # 콘텐츠 생성/로드
                    if post["type"] == "artifact":
                        post_stage = "server_artifact_parse"
                        stage = post_stage
                        post_artifact = post.get("artifact") or {}
                        content = str(post_artifact.get("content_markdown") or "").strip()
                        if not content:
                            raise RuntimeError("서버 생성 글 artifact가 비어 있습니다.")
                        post_usage = post_artifact.get("usage") if isinstance(post_artifact.get("usage"), dict) else {}
                        _merge_usage_totals(usage_totals, post_usage)
                        visible_char_count = measure_visible_char_count(content)
                        parsed_title, content_list = parse_markdown(content)
                        title = str(post_artifact.get("title") or parsed_title or post["source"]).strip()
                        self._log(f"서버 생성 글 사용: {title} (모델: {post_artifact.get('text_model') or ai_model}, 글자 수: {visible_char_count}자)")
                    elif post["type"] == "keyword":
                        keyword = post["source"]
                        cta_info = f", CTA: {cta_link}" if cta_link else ""
                        self._log(f"AI 글 생성 중: {keyword} (스타일: {style_id}, 모델: {ai_model}, 분량: {word_count}자, 이미지: {image_count}장{cta_info})")
                        generated = generate_blog_content(keyword, text_api_key, style_id=style_id, model=ai_model, cta_link=cta_link, cta_text=cta_text, word_count=word_count, image_count=image_count, seo_brief=seo_brief, return_usage=True)
                        if isinstance(generated, tuple):
                            content, post_usage = generated
                        else:
                            content, post_usage = generated, {}
                        _merge_usage_totals(usage_totals, post_usage)
                        if not content:
                            last_error = f"글 생성 실패: {keyword}"
                            stage = post_stage
                            failed_keyword = keyword
                            self._log(f"[오류] {last_error}")
                            _record_failed(post, post_stage, last_error)
                            continue
                        visible_char_count = measure_visible_char_count(content)
                        self._log(f"생성 글자 수 확인: {visible_char_count}자 (요청 {word_count}자)")
                        title, content_list = parse_markdown(content)
                    else:
                        post_usage = {}
                        self._log(f"마크다운 파일 로드: {post['source']}")
                        with open(post["source"], "r", encoding="utf-8") as f:
                            md_content = f.read()
                        visible_char_count = measure_visible_char_count(md_content)
                        title, content_list = parse_markdown(md_content)
                    content_list, repaired_image_prompts = _repair_empty_image_prompts(content_list, title, post.get("source"))
                    if repaired_image_prompts:
                        self._log(f"빈 이미지 프롬프트 {repaired_image_prompts}개를 제목 기반 기본 프롬프트로 보정했습니다.")
                    content_list = _limit_image_blocks(content_list, image_count)
                    image_block_count = sum(1 for item in content_list if item and item[0] == "image")
                    if image_count > 0 and image_block_count < image_count:
                        raise RuntimeError(
                            f"이미지 프롬프트 생성 실패: 요청 {image_count}장 중 {image_block_count}장만 준비되었습니다. "
                            "글이 이미지 없이 저장되지 않도록 작업을 중단했습니다."
                        )

                    if not self.running:
                        break

                    for publish_attempt in range(2):
                        try:
                            # 글쓰기 화면 진입
                            post_stage = "smart_editor_open"
                            stage = post_stage
                            self._log("글쓰기 화면 진입...")
                            navigate_to_editor(self.driver, nid, npw)

                            post_stage = "smart_editor_title"
                            stage = post_stage
                            self._log(f"제목 입력: {title}")
                            input_title(self.driver, title)

                            # 글꼴 설정 (제목 입력 후 본문 영역에서 툴바 활성화됨)
                            if font_name:
                                post_stage = "smart_editor_font"
                                stage = post_stage
                                self._log(f"글꼴 설정: {font_name}")
                                set_font(self.driver, font_name)

                            post_stage = "smart_editor_input"
                            stage = post_stage
                            self._log("본문 입력 중...")
                            input_stats = input_content(
                                self.driver,
                                content_list,
                                image_api_key,
                                image_provider=image_provider,
                                fallback_api_key=fallback_image_api_key,
                            ) or {}
                            image_attempted = _usage_number(input_stats.get("image_attempted")) or image_block_count
                            image_generated = _usage_number(input_stats.get("image_generated"))
                            image_inserted = _usage_number(input_stats.get("image_inserted"))
                            image_totals["attempted"] += image_attempted
                            image_totals["generated"] += image_generated
                            image_totals["inserted"] += image_inserted
                            for provider, count in (input_stats.get("image_providers") or {}).items():
                                if provider in image_totals["providers"]:
                                    image_totals["providers"][provider] += _usage_number(count)
                            image_failures = _safe_image_failures(input_stats.get("image_failures"))
                            if image_failures:
                                image_totals["failures"].extend(image_failures)

                            if image_count > 0 and image_inserted < image_block_count:
                                post_stage = _image_failure_stage(image_failures)
                                stage = post_stage
                                raise RuntimeError(_image_failure_message(image_count, image_inserted, image_failures))

                            # 발행
                            if mode == "save":
                                post_stage = "smart_editor_save"
                                stage = post_stage
                                self._log("임시 저장 중...")
                                save_draft(self.driver)
                            elif mode == "schedule":
                                post_stage = "smart_editor_schedule"
                                stage = post_stage
                                target_date = base_date
                                hour = base_hour
                                if is_multi:
                                    if target_date is None:
                                        target_date = datetime.now() + timedelta(days=1)
                                    if hour is None:
                                        hour = 9
                                    base_dt = target_date.replace(hour=hour, minute=0, second=0)
                                    target_dt = base_dt + timedelta(hours=interval_hours * i)
                                    target_date = target_dt
                                    hour = target_dt.hour
                                    self._log(f"예약 시간: {target_dt.strftime('%Y-%m-%d %H:%M')}")
                                self._log("예약 발행 설정 중...")
                                schedule_publish(self.driver, target_date=target_date, hour=hour, category=category)
                            else:
                                post_stage = "smart_editor_publish"
                                stage = post_stage
                                self._log("즉시 발행 중...")
                                publish_now(self.driver, category=category)
                            break
                        except Exception as publish_error:
                            recoverable_stage = post_stage in {"smart_editor_open", "smart_editor_title", "smart_editor_font"}
                            if publish_attempt == 0 and recoverable_stage and _is_browser_session_error(publish_error):
                                self._log("[복구] 브라우저 세션이 끊겨 재로그인 후 같은 글로 1회 재시도합니다.")
                                _reset_driver("브라우저 세션 재연결")
                                continue
                            raise

                    success += 1
                    post_results.append({
                        "source": post["source"],
                        "keyword": post["source"],
                        "type": post["type"],
                        "title": title,
                        "status": "done",
                        "stage": "completed",
                        "usage": post_usage or {},
                        "images": {
                            "attempted": image_attempted,
                            "generated": image_generated,
                            "inserted": image_inserted,
                            "failure_count": 0,
                            "failures": [],
                        },
                        "char_count": visible_char_count,
                        "target_char_count": word_count,
                    })
                    if is_multi:
                        self._log(f"── [{i+1}/{total}] 완료 ──")

                except Exception as e:
                    last_error = f"{post['source']} 처리 실패: {e}"
                    stage = post_stage if "post_stage" in locals() else stage
                    failed_keyword = post["source"]
                    self._log(f"[오류] {last_error}")
                    _record_failed(
                        post,
                        stage,
                        e,
                        title=title,
                        char_count=visible_char_count,
                        images={
                            "attempted": image_attempted,
                            "generated": image_generated,
                            "inserted": image_inserted,
                            "failure_count": len(image_failures),
                            "failures": image_failures,
                        },
                    )

                # 다음 글까지 대기 (마지막 제외)
                if self.running and i < total - 1:
                    delay = random_delay(10, 3.0, 5.0, 20.0)
                    self._log(f"다음 글까지 {delay:.0f}초 대기...")
                    time.sleep(delay)

            # 3. 완료
            self.queue.put(("progress", 100))
            action_label = {"save": "임시저장", "schedule": "예약 발행", "publish": "발행"}.get(mode, "발행")
            if success != total:
                error_msg = last_error or f"글쓰기 성공 건수가 부족합니다: {success}/{total}"
                self._log(f"[오류] 글쓰기 작업 실패: {success}/{total} 글 성공")
                body = (
                    f"총 {success}/{total}개 글만 {action_label}됐어요.\n"
                    f"앱 로그를 확인해주세요.\n\n"
                    f"마지막 오류: {error_msg}"
                )
                self._write_crew_result(
                    "blog_write",
                    inputs={
                        "keywords": keywords,
                        "mdFile": md_file,
                        "mode": mode,
                        "styleId": style_id,
                        "category": category,
                        "scheduleDate": schedule_date,
                        "scheduleHour": schedule_hour,
                        "scheduleInterval": schedule_interval,
                        "wordCount": word_count,
                        "fontName": font_name,
                    },
                    outputs={
                        "successCount": success,
                        "totalCount": total,
                        "postSources": [post["source"] for post in posts],
                        "mode": mode,
                    },
                    summary=f"글쓰기 작업 실패: {success}/{total}개 성공",
                )
                self.queue.put((
                    "popup",
                    ("write", "고객 설득하기 실패", body, None),
                ))
                return _build_write_result(
                    False, success, total, mode, ai_model, usage_totals, image_totals,
                    post_results, failed_posts, error=error_msg, stage=stage,
                    failed_keyword=failed_keyword,
                )

            if is_multi:
                self._log(f"다중 글쓰기 완료: {success}/{total} 글 {action_label} 성공")
                body = (
                    f"총 {success}/{total}개 글을 {action_label}했어요 📝\n"
                    f"이웃들의 반응을 지켜보자구요!\n\n"
                    f"오늘도 정말 수고하셨어요."
                )
            else:
                self._log("완료!")
                body = (
                    f"글 한 편을 정성껏 써서 {action_label}했어요 📝\n"
                    f"이웃들의 반응을 지켜보자구요!\n\n"
                    f"오늘도 정말 수고하셨어요."
                )
            self._write_crew_result(
                "blog_write",
                inputs={
                    "keywords": keywords,
                    "mdFile": md_file,
                    "mode": mode,
                    "styleId": style_id,
                    "category": category,
                    "scheduleDate": schedule_date,
                    "scheduleHour": schedule_hour,
                    "scheduleInterval": schedule_interval,
                    "wordCount": word_count,
                    "fontName": font_name,
                },
                outputs={
                    "successCount": success,
                    "totalCount": total,
                    "postSources": [post["source"] for post in posts],
                    "mode": mode,
                },
                summary=f"글쓰기 작업 완료: {success}/{total}개 성공",
            )
            self.queue.put((
                "popup",
                ("write", "고객 설득하기 완료", body, None),
            ))
            return _build_write_result(True, success, total, mode, ai_model, usage_totals, image_totals, post_results, failed_posts)

        except Exception as e:
            self._log(f"[오류] {e}")
            traceback.print_exc()
            return _build_write_result(
                False, success, total, mode, ai_model or _DEFAULT_AI_MODEL, usage_totals,
                image_totals, post_results, failed_posts, error=str(e), stage=stage,
                failed_keyword=failed_keyword,
            )
        finally:
            if self.driver:
                try:
                    self.driver.quit()
                except Exception:
                    pass
            self.queue.put(("done", None))

    def _run_bulk(self):
        excel_path = self.bulk_file_var.get().strip()
        if not excel_path:
            self._log("[오류] 엑셀 파일을 선택해주세요.")
            return
        if not self._validate_credentials(need_api=True):
            return
        mode = self.bulk_mode_var.get()
        self._start_worker(self._worker_bulk, excel_path=excel_path, schedule=(mode == "schedule"))

    def _worker_bulk(self, excel_path, schedule):
        try:
            from browser.stealth_driver import create_stealth_driver
            from auth.naver_login import login
            from bulk.excel_loader import load_bulk_rows
            from content.ai_text import generate_blog_content
            from content.markdown_parser import parse_markdown
            from posting.editor import navigate_to_editor, input_title, input_content
            from posting.publisher import publish_now, schedule_publish
            from utils.delays import random_delay
            import time

            ai_model = self.ai_model_var.get()
            if ai_model == "claude":
                text_api_key = self.claude_key_var.get()
            elif _is_openai_model(ai_model):
                text_api_key = self.openai_key_var.get()
            else:
                text_api_key = self.api_key_var.get()
            image_provider = "openai" if _is_openai_model(ai_model) and self.openai_key_var.get() else "gemini"
            image_api_key = self.openai_key_var.get() if image_provider == "openai" else self.api_key_var.get()
            fallback_image_api_key = self.api_key_var.get() if image_provider == "openai" else self.openai_key_var.get()

            self._log(f"엑셀 파일 로드: {excel_path}")
            rows = load_bulk_rows(excel_path)
            total = len(rows)
            success = 0

            for index, row in enumerate(rows):
                if not self.running:
                    break

                account_id = row["account_id"]
                account_pw = row["account_pw"]
                keyword = row["keyword"]
                pct = int((index / total) * 100)
                self.queue.put(("progress", pct))

                self._log(f"[{index + 1}/{total}] {keyword} (계정: {account_id})")

                try:
                    content = generate_blog_content(keyword, text_api_key, model=ai_model)
                    if not content:
                        self._log(f"  글 생성 실패: {keyword}")
                        continue

                    title, content_list = parse_markdown(content)
                    content_list, repaired_image_prompts = _repair_empty_image_prompts(content_list, title, keyword)
                    if repaired_image_prompts:
                        self._log(f"  빈 이미지 프롬프트 {repaired_image_prompts}개를 제목 기반 기본 프롬프트로 보정했습니다.")
                    self.driver = create_stealth_driver()
                    login(self.driver, account_id, account_pw)
                    navigate_to_editor(self.driver, account_id, account_pw)
                    input_title(self.driver, title)
                    input_content(
                        self.driver,
                        content_list,
                        image_api_key,
                        image_provider=image_provider,
                        fallback_api_key=fallback_image_api_key,
                    )

                    if schedule:
                        schedule_publish(self.driver)
                    else:
                        publish_now(self.driver)

                    success += 1
                    self._log(f"  성공: {keyword}")
                except Exception as e:
                    self._log(f"  실패: {keyword} - {e}")
                finally:
                    if self.driver:
                        try:
                            self.driver.quit()
                        except Exception:
                            pass
                        self.driver = None

                if self.running and index < total - 1:
                    delay = random_delay(10, 3.0, 5.0, 20.0)
                    self._log(f"  다음 계정까지 {delay:.0f}초 대기...")
                    time.sleep(delay)

            self.queue.put(("progress", 100))
            self._log(f"대량 발행 완료: 성공 {success}/{total}")

        except Exception as e:
            self._log(f"[오류] {e}")
            traceback.print_exc()
        finally:
            self.queue.put(("done", None))

    def _run_like(self):
        if not self._validate_credentials(need_api=False):
            return
        mode = self.like_mode_var.get()
        keyword = self.like_keyword_var.get().strip() or None
        if mode == "search" and not keyword:
            self._log("[오류] 키워드 검색 모드에서는 키워드를 입력해주세요.")
            return
        self._start_worker(self._worker_like, count=self.like_count_var.get(), mode=mode, keyword=keyword)

    def _worker_like(self, count, mode="neighbor", keyword=None):
        try:
            from browser.stealth_driver import create_stealth_driver
            from auth.naver_login import login
            from engagement.auto_like import auto_like

            nid = self.naver_id_var.get()
            npw = self.naver_pw_var.get()

            self.queue.put(("progress", 10))
            self._log("스텔스 브라우저 시작...")
            self.driver = create_stealth_driver()
            if self.stop_event.is_set():
                self._log("사용자에 의해 중지됨")
                return

            self.queue.put(("progress", 30))
            self._log("네이버 로그인 중...")
            login(self.driver, nid, npw)
            if self.stop_event.is_set():
                self._log("사용자에 의해 중지됨")
                return

            self.queue.put(("progress", 50))
            mode_desc = f"키워드 '{keyword}'" if mode == "search" else "이웃 새글"
            self._log(f"자동 공감 시작 - {mode_desc} (최대 {count}개)...")
            liked = auto_like(
                self.driver, max_posts=count, mode=mode, keyword=keyword,
                stop_event=self.stop_event
            )

            if self.stop_event.is_set():
                self._log("사용자에 의해 중지됨")
                return

            self.queue.put(("progress", 100))
            self._log(f"자동 공감 완료: {liked}개 성공")

        except Exception as e:
            if self.stop_event.is_set():
                self._log("사용자에 의해 중지됨")
            else:
                self._log(f"[오류] {e}")
                traceback.print_exc()
        finally:
            if self.driver:
                try:
                    self.driver.quit()
                except Exception:
                    pass
            self.queue.put(("done", None))

    def _run_comment(self):
        if not self._validate_credentials(need_api=True):
            return
        mode = self.comment_mode_var.get()
        keyword = self.comment_keyword_var.get().strip() or None
        tone = self.comment_tone_var.get().split(" - ")[0].strip()
        custom = self.comment_custom_var.get().strip() or None
        if mode == "search" and not keyword:
            self._log("[오류] 키워드 검색 모드에서는 키워드를 입력해주세요.")
            return
        comment_count = int(self.comment_count_var.get() or 0)
        if not self._confirm_limit_override(
            "일일 댓글",
            comment_count,
            aimax.RECOMMENDED_LIMITS["daily_comments"],
        ):
            self._log("권장 한도 초과 확인이 취소되어 실행하지 않았습니다.")
            return
        self._start_worker(
            self._worker_comment, count=comment_count,
            mode=mode, keyword=keyword, tone=tone, custom_instruction=custom,
            ai_model=self.ai_model_var.get()
        )

    def _worker_comment(self, count, mode="neighbor", keyword=None, tone="friendly", custom_instruction=None, ai_model=_DEFAULT_AI_MODEL):
        try:
            from browser.stealth_driver import create_stealth_driver
            from auth.naver_login import login
            from engagement.auto_comment import auto_comment

            nid = self.naver_id_var.get()
            npw = self.naver_pw_var.get()
            if ai_model == "claude":
                api_key = self.claude_key_var.get()
            elif _is_openai_model(ai_model):
                api_key = self.openai_key_var.get()
            else:
                api_key = self.api_key_var.get()

            self.queue.put(("progress", 10))
            self._log("스텔스 브라우저 시작...")
            self.driver = create_stealth_driver()
            if self.stop_event.is_set():
                self._log("사용자에 의해 중지됨")
                return

            self.queue.put(("progress", 30))
            self._log("네이버 로그인 중...")
            login(self.driver, nid, npw)
            if self.stop_event.is_set():
                self._log("사용자에 의해 중지됨")
                return

            self.queue.put(("progress", 50))
            mode_desc = f"키워드 '{keyword}'" if mode == "search" else "이웃 새글"
            self._log(f"자동 댓글 시작 - {mode_desc} (최대 {count}개)...")
            commented = auto_comment(
                self.driver, api_key, max_posts=count,
                mode=mode, keyword=keyword, tone=tone,
                custom_instruction=custom_instruction, naver_id=nid,
                ai_model=ai_model, stop_event=self.stop_event
            )

            if self.stop_event.is_set():
                self._log("사용자에 의해 중지됨")
                return

            self.queue.put(("progress", 100))
            self._log(f"자동 댓글 완료: {commented}개 성공")

        except Exception as e:
            if self.stop_event.is_set():
                self._log("사용자에 의해 중지됨")
            else:
                self._log(f"[오류] {e}")
                traceback.print_exc()
        finally:
            if self.driver:
                try:
                    self.driver.quit()
                except Exception:
                    pass
            self.queue.put(("done", None))

    def _run_engage(self):
        if not self._validate_credentials(need_api=True):
            return

        mode = "neighbor"
        keyword = None
        self.engage_mode_var.set(mode)
        self.engage_keyword_var.set("")

        tone = self.engage_tone_var.get().split(" - ")[0].strip()
        comment_count = int(self.engage_comment_var.get() or 0)
        if not self._confirm_limit_override(
            "일일 댓글",
            comment_count,
            aimax.RECOMMENDED_LIMITS["daily_comments"],
        ):
            self._log("권장 한도 초과 확인이 취소되어 실행하지 않았습니다.")
            return
        self._start_worker(
            self._worker_engage,
            like_count=self.engage_like_var.get(),
            comment_count=comment_count,
            mode=mode,
            keyword=keyword,
            tone=tone,
            ai_model=self.ai_model_var.get(),
        )

    def _worker_engage(self, like_count, comment_count, mode="neighbor", keyword=None, tone="friendly", ai_model=_DEFAULT_AI_MODEL):
        try:
            from browser.stealth_driver import create_stealth_driver
            from auth.naver_login import login
            from engagement.auto_like import auto_like
            from engagement.auto_comment import auto_comment

            nid = self.naver_id_var.get()
            npw = self.naver_pw_var.get()
            if ai_model == "claude":
                api_key = self.claude_key_var.get()
            elif _is_openai_model(ai_model):
                api_key = self.openai_key_var.get()
            else:
                api_key = self.api_key_var.get()

            self.queue.put(("progress", 10))
            self._log("스텔스 브라우저 시작...")
            self.driver = create_stealth_driver()
            if self.stop_event.is_set():
                self._log("사용자에 의해 중지됨")
                return

            self.queue.put(("progress", 20))
            self._log("네이버 로그인 중...")
            login(self.driver, nid, npw)
            if self.stop_event.is_set():
                self._log("사용자에 의해 중지됨")
                return

            mode_desc = f"키워드 '{keyword}'" if mode == "search" else "이웃 새글"

            self.queue.put(("progress", 35))
            self._log(f"자동 공감 시작 - {mode_desc} (최대 {like_count}개)...")
            liked = auto_like(
                self.driver, max_posts=like_count, mode=mode, keyword=keyword,
                stop_event=self.stop_event
            )
            if self.stop_event.is_set():
                self._log("사용자에 의해 중지됨")
                return
            self._log(f"자동 공감 완료: {liked}개 성공")

            commented = 0
            self.queue.put(("progress", 65))
            if not self.stop_event.is_set():
                self._log(f"자동 댓글 시작 - {mode_desc} (최대 {comment_count}개)...")
                commented = auto_comment(
                    self.driver, api_key, max_posts=comment_count,
                    mode=mode, keyword=keyword, tone=tone, naver_id=nid,
                    ai_model=ai_model, stop_event=self.stop_event
                )
                if self.stop_event.is_set():
                    self._log("사용자에 의해 중지됨")
                    return
                self._log(f"자동 댓글 완료: {commented}개 성공")

            self.queue.put(("progress", 100))
            self._log("공감+댓글 전체 완료!")

            # 완료 팝업
            body = (
                f"오늘 공감 {liked}개, 댓글 {commented}개를 남겼어요 💬\n"
                f"이웃들과 친밀도가 높아졌을 거예요.\n\n"
                f"이제 글로 마음을 움직일 시간이에요!"
            )
            self._write_crew_result(
                "blog_engage",
                inputs={
                    "likeCount": like_count,
                    "commentCount": comment_count,
                    "mode": mode,
                    "keyword": keyword,
                    "tone": tone,
                    "aiModel": ai_model,
                },
                outputs={
                    "likedCount": liked,
                    "commentedCount": commented,
                    "mode": mode,
                    "engagementSummary": f"공감 {liked}개, 댓글 {commented}개",
                },
                summary=f"관계 형성 작업 완료: 공감 {liked}개, 댓글 {commented}개",
            )
            self.queue.put((
                "popup",
                ("engage", "고객과 친해지기 완료", body, "write"),
            ))

        except Exception as e:
            if self.stop_event.is_set():
                self._log("사용자에 의해 중지됨")
            else:
                self._log(f"[오류] {e}")
                traceback.print_exc()
        finally:
            if self.driver:
                try:
                    self.driver.quit()
                except Exception:
                    pass
            self.queue.put(("done", None))

    def _generate_neighbor_messages_ai(self):
        """내 블로그 프로필 기반으로 서이추 멘트 10개 배치 생성.

        1. 프로필 검증 → 없으면 안내
        2. API 키 검증
        3. 백그라운드 스레드에서 생성 (UI 프리즈 방지)
        4. 결과를 Text 영역에 삽입
        """
        employee_name = self.mode_config.get("employee_name", "예리씨")
        profile = self._get_blog_profile_text()
        if not profile:
            self._log("[AI 멘트] 먼저 [직원 설정]에서 '내 블로그 프로필'을 입력해 주세요.")
            try:
                from tkinter import messagebox
                messagebox.showinfo(
                    "프로필이 없어요",
                    f"{employee_name}가 대표님 블로그를 아직 몰라서 멘트를 쓸 수 없어요.\n\n"
                    "[직원 설정] 탭에서 '내 블로그 프로필'을 먼저 입력해 주세요!",
                )
            except Exception:
                pass
            return

        model = self.ai_model_var.get() or _DEFAULT_AI_MODEL
        gemini_key = self.api_key_var.get().strip()
        claude_key = self.claude_key_var.get().strip()
        openai_key = self.openai_key_var.get().strip()

        if model == "claude" and not claude_key:
            self._log("[AI 멘트] Claude API Key가 없습니다.")
            return
        if _is_openai_model(model) and not openai_key:
            self._log("[AI 멘트] OpenAI API Key가 없습니다.")
            return
        if model != "claude" and not _is_openai_model(model) and not gemini_key:
            self._log("[AI 멘트] Gemini API Key가 없습니다.")
            return

        # UI 잠금 + 로딩 메시지
        self.ai_ment_btn.configure(state="disabled")
        self.ai_ment_status.configure(text=f"{employee_name}가 멘트를 고민하고 있어요…", fg="#FF6B9D")
        self._log("[AI 멘트] 프로필 기반 멘트 10개 생성 중...")

        def _worker():
            try:
                from content.neighbor_message_ai import generate_neighbor_messages
                messages = generate_neighbor_messages(
                    profile_text=profile,
                    api_key=gemini_key,
                    model=model,
                    count=10,
                    claude_key=claude_key,
                    openai_key=openai_key,
                )
                self.queue.put(("ai_ment_done", messages))
            except Exception as e:
                self.queue.put(("ai_ment_error", str(e)))

        import threading
        threading.Thread(target=_worker, daemon=True).start()

    def _apply_generated_messages(self, messages):
        """AI 생성 결과를 멘트 Text 영역에 반영"""
        self.ai_ment_btn.configure(state="normal")
        if not messages:
            self.ai_ment_status.configure(
                text="⚠ 생성 실패 — 프로필이나 API 키를 확인해 주세요", fg="#C0392B",
            )
            self._log("[AI 멘트] 생성 실패 또는 응답이 비어있습니다.")
            return
        self.neighbor_msg_text.delete("1.0", "end")
        self.neighbor_msg_text.insert("1.0", "\n".join(messages))
        if hasattr(self, "link_neighbor_msg_text"):
            self.link_neighbor_msg_text.delete("1.0", "end")
            self.link_neighbor_msg_text.insert("1.0", "\n".join(messages))
            self.link_neighbor_status.configure(
                text=f"✓ {len(messages)}개 생성·저장 완료 — 편집 가능", fg="#27AE60",
            )
        save_neighbor_messages(self.naver_id_var.get(), messages)
        self.ai_ment_status.configure(
            text=f"✓ {len(messages)}개 생성·저장 완료 — 편집 가능", fg="#27AE60",
        )
        self._log(f"[AI 멘트] {len(messages)}개 생성 후 저장 완료 (네이버 ID: {self.naver_id_var.get()})")

    def _collect_neighbor_messages(self, text_widget, status_label=None):
        """서이추 멘트 Text 위젯에서 줄 단위 멘트 풀을 읽고 검증한다."""
        raw_msgs = text_widget.get("1.0", END)
        messages = [m.strip() for m in raw_msgs.splitlines() if m.strip()]
        if not messages:
            self._log("[오류] 멘트를 최소 1개 이상 입력해주세요.")
            return None
        too_long = [m for m in messages if len(m) > 400]
        if too_long:
            self._log(f"[오류] 서로이웃 신청 멘트는 400자를 넘을 수 없습니다. ({len(too_long)}개 초과)")
            try:
                from tkinter import messagebox
                messagebox.showwarning(
                    "멘트가 너무 길어요",
                    "네이버 서로이웃 신청 메시지는 400자를 넘으면 제출되지 않습니다.\n\n"
                    f"400자를 넘는 멘트 {len(too_long)}개를 줄인 뒤 다시 실행해 주세요.",
                )
            except Exception:
                pass
            return None
        save_neighbor_messages(self.naver_id_var.get(), messages)
        if status_label is not None:
            status_label.configure(text=f"멘트 {len(messages)}개 저장됨", fg=COLORS["text_muted"])
        return messages

    def _run_neighbor(self):
        keywords_str = self.neighbor_keywords_var.get().strip()
        if not keywords_str:
            self._log("[오류] 검색 키워드를 입력해주세요.")
            return
        if not self._validate_credentials(need_api=False):
            return
        keywords = [k.strip() for k in keywords_str.split(",") if k.strip()]
        max_per = self.neighbor_count_var.get()

        messages = self._collect_neighbor_messages(self.neighbor_msg_text, self.ai_ment_status)
        if messages is None:
            return

        speed_mode = self.neighbor_speed_var.get() or "normal"
        cooldown_every = max(0, int(self.neighbor_cooldown_var.get() or 0))
        daily_limit = max(0, int(self.neighbor_daily_limit_var.get() or 0))
        if not self._confirm_limit_override(
            "서로이웃 일일 상한",
            daily_limit,
            aimax.RECOMMENDED_LIMITS["daily_neighbor_requests"],
        ):
            self._log("권장 한도 초과 확인이 취소되어 실행하지 않았습니다.")
            return

        self._start_worker(
            self._worker_neighbor,
            keywords=keywords, max_per_keyword=max_per,
            messages=messages, speed_mode=speed_mode,
            cooldown_every=cooldown_every, daily_limit=daily_limit,
        )

    def _run_link_neighbor(self):
        blogger_url = self.link_blogger_url_var.get().strip()
        if not blogger_url:
            self._log("[오류] 블로거 링크 또는 ID를 입력해주세요.")
            return
        try:
            from scraper.follower_scraper import extract_blogger_id_from_url
            if not extract_blogger_id_from_url(blogger_url):
                self._log("[오류] 네이버 블로그 링크 또는 블로그 ID 형식을 확인해주세요.")
                return
        except Exception as e:
            self._log(f"[오류] 블로거 링크 확인 실패: {e}")
            return
        if not self._validate_credentials(need_api=False):
            return

        messages = self._collect_neighbor_messages(
            self.link_neighbor_msg_text, self.link_neighbor_status,
        )
        if messages is None:
            return

        max_requests = max(1, int(self.link_neighbor_count_var.get() or 1))
        speed_mode = self.neighbor_speed_var.get() or "normal"
        cooldown_every = max(0, int(self.neighbor_cooldown_var.get() or 0))
        daily_limit = max(0, int(self.neighbor_daily_limit_var.get() or 0))
        if not self._confirm_limit_override(
            "서로이웃 일일 상한",
            daily_limit,
            aimax.RECOMMENDED_LIMITS["daily_neighbor_requests"],
        ):
            self._log("권장 한도 초과 확인이 취소되어 실행하지 않았습니다.")
            return

        self._start_worker(
            self._worker_link_neighbor,
            blogger_url=blogger_url, max_requests=max_requests,
            messages=messages, speed_mode=speed_mode,
            cooldown_every=cooldown_every, daily_limit=daily_limit,
        )

    def _worker_neighbor(self, keywords, max_per_keyword, messages=None,
                        speed_mode="normal", cooldown_every=10, daily_limit=50):
        try:
            from browser.stealth_driver import create_stealth_driver
            from auth.naver_login import login
            from engagement.auto_neighbor import auto_neighbor

            nid = self.naver_id_var.get()
            npw = self.naver_pw_var.get()

            self.queue.put(("progress", 10))
            self._log("스텔스 브라우저 시작...")
            self.driver = create_stealth_driver()
            if self.stop_event.is_set():
                self._log("사용자에 의해 중지됨")
                return

            self.queue.put(("progress", 25))
            self._log("네이버 로그인 중...")
            login(self.driver, nid, npw)
            if self.stop_event.is_set():
                self._log("사용자에 의해 중지됨")
                return

            self.queue.put(("progress", 40))
            speed_label = {"safe": "안전", "normal": "보통", "fast": "빠름"}.get(speed_mode, speed_mode)
            self._log(
                f"서로이웃 추가 시작 (키워드: {', '.join(keywords)}, "
                f"속도={speed_label}, 멘트 {len(messages or [])}종, "
                f"Cool-down={cooldown_every}명, 일일 상한={daily_limit or '무제한'})"
            )
            total = auto_neighbor(
                self.driver, keywords, max_per_keyword=max_per_keyword,
                messages=messages, speed_mode=speed_mode,
                cooldown_every=cooldown_every, daily_limit=daily_limit,
                naver_id=nid, naver_pw=npw, stop_event=self.stop_event,
            )

            if self.stop_event.is_set():
                self._log("사용자에 의해 중지됨")
                return

            self.queue.put(("progress", 100))
            self._log(f"서로이웃 추가 완료: {total}명 신청")
            self._write_crew_result(
                "blog_find",
                inputs={
                    "source": "keyword",
                    "keywords": keywords,
                    "maxPerKeyword": max_per_keyword,
                    "speedMode": speed_mode,
                    "cooldownEvery": cooldown_every,
                    "dailyLimit": daily_limit,
                },
                outputs={
                    "source": "keyword",
                    "blogIds": [],
                    "sentCount": total,
                    "successCount": total,
                    "keywords": keywords,
                },
                summary=f"키워드 기반 고객 찾기 완료: {total}명 신청",
            )

            # 완료 팝업
            try:
                from engagement.neighbor_quota import get_today_count
                today = get_today_count(nid)
                daily_txt = f" / 일일 상한 {daily_limit}명" if daily_limit else ""
                body = (
                    f"이번에 {total}명에게 서로이웃을 신청했어요 ✨\n"
                    f"(오늘 누적 {today}명{daily_txt})\n\n"
                    f"이제 댓글과 공감으로 친해질 시간이에요!"
                )
            except Exception:
                body = f"이번에 {total}명에게 서로이웃을 신청했어요 ✨"
            self.queue.put((
                "popup",
                ("find_keyword", "고객을 찾아오기 완료", body, "engage"),
            ))

        except Exception as e:
            if self.stop_event.is_set():
                self._log("사용자에 의해 중지됨")
            else:
                self._log(f"[오류] {e}")
                traceback.print_exc()
        finally:
            if self.driver:
                try:
                    self.driver.quit()
                except Exception:
                    pass
            self.queue.put(("done", None))

    def _worker_link_neighbor(self, blogger_url, max_requests, messages=None,
                              speed_mode="normal", cooldown_every=10, daily_limit=50):
        try:
            from browser.stealth_driver import create_stealth_driver
            from auth.naver_login import login
            from scraper.follower_scraper import (
                extract_blogger_id_from_url,
                scrape_follower_ids,
            )
            from engagement.auto_neighbor import auto_neighbor_to_blog_ids

            nid = self.naver_id_var.get()
            npw = self.naver_pw_var.get()
            owner_id = extract_blogger_id_from_url(blogger_url) or blogger_url

            self.queue.put(("progress", 10))
            self._log("스텔스 브라우저 시작...")
            self.driver = create_stealth_driver()
            if self.stop_event.is_set():
                self._log("사용자에 의해 중지됨")
                return

            self.queue.put(("progress", 25))
            self._log("네이버 로그인 중...")
            login(self.driver, nid, npw)
            if self.stop_event.is_set():
                self._log("사용자에 의해 중지됨")
                return

            self.queue.put(("progress", 40))
            scrape_limit = max(max_requests * 3, max_requests + 10)
            self._log(f"'{owner_id}' 블로거의 이웃 목록 수집 중... (후보 최대 {scrape_limit}명)")
            blog_ids = scrape_follower_ids(
                self.driver, blogger_url, max_count=scrape_limit,
                stop_event=self.stop_event,
            )
            if self.stop_event.is_set():
                self._log("사용자에 의해 중지됨")
                return
            if not blog_ids:
                self.queue.put(("progress", 100))
                self._log("이웃 목록이 비공개이거나 접근 제한됨")
                self._write_crew_result(
                    "blog_find",
                    status="empty",
                    inputs={
                        "source": "blogger_link",
                        "bloggerUrl": blogger_url,
                        "ownerId": owner_id,
                        "maxRequests": max_requests,
                    },
                    outputs={
                        "source": "blogger_link",
                        "ownerId": owner_id,
                        "blogIds": [],
                        "candidateCount": 0,
                        "sentCount": 0,
                        "reason": "이웃 목록이 비공개이거나 접근 제한됨",
                    },
                    summary=f"{owner_id} 이웃 목록 접근 제한 또는 후보 없음",
                )
                return

            self.queue.put(("progress", 60))
            speed_label = {"safe": "안전", "normal": "보통", "fast": "빠름"}.get(speed_mode, speed_mode)
            self._log(
                f"서로이웃 추가 시작 (기준 블로거: {owner_id}, 후보 {len(blog_ids)}명, "
                f"최대 신청 {max_requests}명, 속도={speed_label}, 멘트 {len(messages or [])}종, "
                f"Cool-down={cooldown_every}명, 일일 상한={daily_limit or '무제한'})"
            )
            total = auto_neighbor_to_blog_ids(
                self.driver, blog_ids, max_requests=max_requests,
                messages=messages, speed_mode=speed_mode,
                cooldown_every=cooldown_every, daily_limit=daily_limit,
                naver_id=nid, naver_pw=npw, stop_event=self.stop_event,
                source_label=f"{owner_id} 이웃 목록",
            )

            if self.stop_event.is_set():
                self._log("사용자에 의해 중지됨")
                return

            self.queue.put(("progress", 100))
            self._log(f"특정 블로거 링크 작업 완료: 후보 {len(blog_ids)}명 중 {total}명 신청")
            self._write_crew_result(
                "blog_find",
                inputs={
                    "source": "blogger_link",
                    "bloggerUrl": blogger_url,
                    "ownerId": owner_id,
                    "maxRequests": max_requests,
                    "speedMode": speed_mode,
                    "cooldownEvery": cooldown_every,
                    "dailyLimit": daily_limit,
                },
                outputs={
                    "source": "blogger_link",
                    "ownerId": owner_id,
                    "blogIds": blog_ids,
                    "candidateCount": len(blog_ids),
                    "sentCount": total,
                    "successCount": total,
                },
                summary=f"특정 블로거 링크 기반 고객 찾기 완료: 후보 {len(blog_ids)}명 중 {total}명 신청",
            )

            try:
                from engagement.neighbor_quota import get_today_count
                today = get_today_count(nid)
                daily_txt = f" / 일일 상한 {daily_limit}명" if daily_limit else ""
                body = (
                    f"'{owner_id}' 이웃 목록에서 후보 {len(blog_ids)}명을 찾고,\n"
                    f"이번에 {total}명에게 서로이웃을 신청했어요.\n"
                    f"(오늘 누적 {today}명{daily_txt})"
                )
            except Exception:
                body = f"'{owner_id}' 이웃 목록에서 {total}명에게 서로이웃을 신청했어요."
            self.queue.put((
                "popup",
                ("find_link", "고객을 찾아오기 완료", body, "engage"),
            ))

        except Exception as e:
            if self.stop_event.is_set():
                self._log("사용자에 의해 중지됨")
            else:
                self._log(f"[오류] {e}")
                traceback.print_exc()
        finally:
            if self.driver:
                try:
                    self.driver.quit()
                except Exception:
                    pass
            self.queue.put(("done", None))

    def _run_scraper(self):
        keyword = self.scraper_keyword_var.get().strip()
        if not keyword:
            self._log("[오류] 검색 키워드를 입력해주세요.")
            return
        self._start_worker(self._worker_scraper_blog, keyword=keyword,
                           max_results=self.scraper_count_var.get())

    def _worker_scraper_blog(self, keyword, max_results):
        try:
            from browser.stealth_driver import create_stealth_driver
            from auth.naver_login import login
            from scraper.blog_scraper import scrape_blogger_ids, save_to_csv

            nid = self.naver_id_var.get()
            npw = self.naver_pw_var.get()

            self.queue.put(("progress", 10))
            self._log("스텔스 브라우저 시작...")
            self.driver = create_stealth_driver()

            self.queue.put(("progress", 25))
            self._log("네이버 로그인 중...")
            login(self.driver, nid, npw)

            self.queue.put(("progress", 40))
            self._log(f"블로거 ID 수집 중: '{keyword}'")
            ids = scrape_blogger_ids(self.driver, keyword, max_results=max_results)

            self.queue.put(("progress", 90))
            filepath = save_to_csv(keyword, ids, include_email=True)
            self.last_scraper_csv_path = filepath
            self._log(f"수집 완료: {len(ids)}명 → {filepath}")
            if self._open_local_path(filepath, quiet=True):
                self._log("CSV 파일을 기본 프로그램으로 열었습니다.")
            self.queue.put(("progress", 100))

        except Exception as e:
            self._log(f"[오류] {e}")
            traceback.print_exc()
        finally:
            if self.driver:
                try:
                    self.driver.quit()
                except Exception:
                    pass
            self.queue.put(("done", None))

    def _open_local_path(self, path, quiet=False):
        try:
            if not path or not os.path.exists(path):
                if not quiet:
                    self._log(f"[오류] 경로를 찾을 수 없습니다: {path}")
                return False

            if sys.platform.startswith("win"):
                os.startfile(path)
            elif sys.platform == "darwin":
                subprocess.Popen(["open", path])
            else:
                subprocess.Popen(["xdg-open", path])
            return True
        except Exception as e:
            if not quiet:
                self._log(f"[오류] 파일 열기 실패: {e}")
            return False

    def _open_last_scraper_csv(self):
        if self.last_scraper_csv_path and os.path.exists(self.last_scraper_csv_path):
            if self._open_local_path(self.last_scraper_csv_path):
                self._log(f"CSV 열기: {self.last_scraper_csv_path}")
            return

        self._log("최근 CSV가 없어 exports 폴더를 엽니다.")
        self._open_exports_folder()

    def _open_exports_folder(self):
        os.makedirs(EXPORTS_DIR, exist_ok=True)
        if self._open_local_path(EXPORTS_DIR):
            self._log(f"내보내기 폴더 열기: {EXPORTS_DIR}")

    # ── 종료 ──
    def _on_close(self):
        self.running = False
        try:
            self.web_agent_stop_event.set()
        except Exception:
            pass
        driver = getattr(self, "driver", None)
        if driver:
            try:
                driver.quit()
            except Exception:
                pass
        self.root.destroy()

    # ── 실행 ──
    def run(self):
        self.root.mainloop()


class HeadlessNaverBlogAgent(HeadlessAgentMixin, NaverBlogApp):
    def __init__(self, app_mode=None):
        self.app_mode = self._normalize_app_mode(app_mode or os.environ.get("APP_MODE", "all"))
        self.mode_config = self.MODE_CONFIG[self.app_mode]
        self._init_headless_agent(
            settings_loader=load_settings,
            settings_saver=save_settings,
            local_settings_saver=save_local_security_settings,
            settings_recoverer=recover_missing_settings_secrets,
            default_ai_model=_DEFAULT_AI_MODEL,
            normalizer=_normalize_ai_model,
            api_key_guide_url=API_KEY_GUIDE_URL,
        )


def _parse_runtime_args(argv=None):
    parser = argparse.ArgumentParser(description="AIMAX split Local Agent")
    parser.add_argument("--agent", action="store_true", help="화면 없는 Local Agent 모드로 실행")
    parser.add_argument("--legacy-ui", action="store_true", help="기존 Tkinter UI를 실행")
    parser.add_argument("--connect", action="store_true", help=argparse.SUPPRESS)
    parser.add_argument("--status", action="store_true", help=argparse.SUPPRESS)
    parser.add_argument("--open-settings", dest="open_settings", action="store_true", help=argparse.SUPPRESS)
    parser.add_argument("--open_settings", dest="open_settings", action="store_true", help=argparse.SUPPRESS)
    parser.add_argument("--diagnostics-probe", metavar="PATH", help=argparse.SUPPRESS)
    parser.add_argument("--repair-local-state", metavar="PATH", help=argparse.SUPPRESS)
    parser.add_argument("--repair-local-state-dry-run", action="store_true", help=argparse.SUPPRESS)
    parser.add_argument("--repair-local-state-stale-seconds", type=int, default=3600, help=argparse.SUPPRESS)
    args, _unknown = parser.parse_known_args(argv)
    return args


def _run_local_state_repair(args):
    output_path = getattr(args, "repair_local_state", None)
    if not output_path:
        return False
    from local_agent.state_repair import quarantine_local_state_conflicts

    result = quarantine_local_state_conflicts(
        dry_run=bool(getattr(args, "repair_local_state_dry_run", False)),
        stale_seconds=max(0, int(getattr(args, "repair_local_state_stale_seconds", 3600) or 3600)),
    )
    output_path = os.path.abspath(output_path)
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)
    return True


def _run_diagnostics_probe(args):
    output_path = getattr(args, "diagnostics_probe", None)
    if not output_path:
        return False
    from diagnostics.error_reporter import build_error_report

    report = build_error_report(
        work_context="runtime diagnostics probe",
        visible_error="diagnostics probe",
        user_note="",
        console_log="diagnostics probe",
        source="diagnostics_probe",
    )
    try:
        from content.ai_text import generate_blog_content, measure_visible_char_count

        report["ai_text_import_smoke"] = {
            "ok": True,
            "module_file": getattr(sys.modules.get("content.ai_text"), "__file__", None),
            "has_generate_blog_content": callable(generate_blog_content),
            "has_measure_visible_char_count": callable(measure_visible_char_count),
            "sample_visible_char_count": measure_visible_char_count("AIMAX 테스트 123"),
        }
    except Exception as exc:
        report["ai_text_import_smoke"] = {
            "ok": False,
            "error": f"{type(exc).__name__}: {exc}",
        }
    try:
        import tempfile
        from pathlib import Path

        from openpyxl import Workbook
        from bulk.excel_loader import load_bulk_rows

        with tempfile.TemporaryDirectory(prefix="aimax-excel-probe-") as temp_dir:
            excel_path = Path(temp_dir) / "bulk.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["아이디", "비밀번호", "핵심키워드"])
            sheet.append(["probe_id", "probe_pw", "AIMAX 진단"])
            workbook.save(excel_path)
            workbook.close()
            rows = load_bulk_rows(excel_path)
        report["excel_loader_import_smoke"] = {
            "ok": rows == [{"account_id": "probe_id", "account_pw": "probe_pw", "keyword": "AIMAX 진단"}],
            "row_count": len(rows),
        }
    except Exception as exc:
        report["excel_loader_import_smoke"] = {
            "ok": False,
            "error": f"{type(exc).__name__}: {exc}",
        }
    output_path = os.path.abspath(output_path)
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(report, f, ensure_ascii=False, indent=2)
    return True


if __name__ == "__main__":
    import multiprocessing

    multiprocessing.freeze_support()
    args = _parse_runtime_args()
    if _run_local_state_repair(args):
        sys.exit(0)
    if _run_diagnostics_probe(args):
        sys.exit(0)
    app = HeadlessNaverBlogAgent() if agent_mode_requested(args) else NaverBlogApp()
    if _EARLY_AGENT_LOCK is not None and agent_mode_requested(args):
        app._single_instance_lock = _EARLY_AGENT_LOCK
    app.run()
